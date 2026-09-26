#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
claim_guard.py — Preliminary Claim Linter (Bộ rà soát sơ bộ ngôn từ tiếp thị & over-claim)
Tuân thủ Luật R5: .agents/rules/R5-legal-claim-compliance.md
Căn cứ: Luật Quảng cáo 2012, NĐ 181/2013/NĐ-CP, NĐ 38/2021/NĐ-CP, TT 06/2011/TT-BYT.

QUAN TRỌNG:
Công cụ này là PRELIMINARY CLAIM LINTER (Bộ lọc từ khóa sơ bộ), KHÔNG PHẢI là Legal Correctness Verifier
(Bộ thẩm định tính đúng đắn pháp lý toàn diện) và không thay thế ý kiến tư vấn pháp lý chuyên môn.

Sử dụng:
    python3 scripts/claim_guard.py --input <đường_dẫn_file> [--strict] [--json]
"""

import sys
import os
import re
import json
import argparse
from pathlib import Path

# Màu sắc hiển thị Terminal
GREEN = "\033[92m"
RED = "\033[91m"
YELLOW = "\033[93m"
BLUE = "\033[94m"
BOLD = "\033[1m"
RESET = "\033[0m"

# Danh mục các mẫu cụm từ vi phạm pháp lý over-claim
CLAIM_RULES = [
    {
        "category": "KHẲNG ĐỊNH TUYỆT ĐỐI / SUPERLATIVES",
        "severity": "CAO (Vi phạm Luật Quảng cáo Đ.8 K.11)",
        "patterns": [
            (r"\ban toàn tuyệt đối\b", "Dịu nhẹ và lành tính / Được kiểm nghiệm dịu nhẹ"),
            (r"\btuyệt đối an toàn\b", "Lành tính cho làn da / An tâm khi sử dụng"),
            (r"\b100%\s*an toàn\b", "Được chứng minh dịu nhẹ"),
            (r"\bhoàn toàn không có tác dụng phụ\b", "Hạn chế tối đa nguy cơ kích ứng"),
            (r"\bkhông (gây )?kích ứng với mọi loại da\b", "Phù hợp cho cả làn da nhạy cảm"),
            (r"\bvĩnh viễn\b", "Lâu dài / Bền màu theo chu trình chăm sóc"),
            (r"\btrọn đời\b", "Dài hạn / Bền bỉ"),
            (r"\bkhông bao giờ (mọc lại|tái phát)\b", "Hạn chế xuất hiện trở lại"),
            (r"\bđỉnh cao nhất\b", "Nổi bật / Tiêu biểu"),
            (r"\bhoàn hảo nhất\b", "Tối ưu / Toàn diện"),
        ]
    },
    {
        "category": "NHẦM LẪN DƯỢC PHẨM / ĐIỀU TRỊ Y TẾ (MỸ PHẨM & TPCN)",
        "severity": "CAO (Vi phạm TT 06/2011/TT-BYT Phụ lục 03-MP)",
        "patterns": [
            (r"\bđặc trị\b", "Hỗ trợ cải thiện / Chăm sóc chuyên sâu"),
            (r"\btrị dứt điểm\b", "Hỗ trợ che phủ / Cải thiện rõ rệt"),
            (r"\bchữa khỏi\b", "Hỗ trợ phục hồi / Chăm sóc"),
            (r"\btiêu diệt tận gốc\b", "Làm sạch sâu / Loại bỏ nhẹ nhàng"),
            (r"\bkhỏi hẳn\b", "Cải thiện tích cực"),
            (r"\bchữa lành\b", "Phục hồi / Nuôi dưỡng"),
            (r"\bthay thế thuốc\b", "Hỗ trợ chăm sóc sức khỏe"),
            (r"\btái tạo tế bào gốc\b", "Hỗ trợ tái tạo bề mặt da/tóc"),
            (r"\bthuốc phủ bạc\b", "Dầu gội phủ màu / Sản phẩm phủ màu tóc bạc"),
            (r"\bthuốc nhuộm tóc bạc\b", "Dầu gội / Kem ủ phủ màu tóc bạc"),
        ]
    },
    {
        "category": "CAM KẾT TÀI CHÍNH & PHÁP LÝ PHI THỰC TẾ",
        "severity": "TRUNG BÌNH (Rủi ro tranh chấp & lừa dối khách hàng)",
        "patterns": [
            (r"\bcam kết sinh lời 100%\b", "Mục tiêu tối ưu hóa tỷ suất sinh lời"),
            (r"\bchắc chắn có lãi\b", "Kỳ vọng tăng trưởng lợi nhuận"),
            (r"\bkhông có (bất kỳ )?rủi ro nào\b", "Kiểm soát rủi ro ở mức tối thiểu"),
            (r"\bchắc chắn trúng thầu\b", "Tối đa hóa năng lực cạnh tranh hồ sơ thầu"),
            (r"\b(đảm bảo|chắc chắn) thắng kiện\b", "Bảo vệ tối đa quyền lợi hợp pháp"),
        ]
    },
    {
        "category": "BẢO MẬT & CÔNG NGHỆ CỰC ĐOAN",
        "severity": "TRUNG BÌNH (Tuyên bố phi kỹ thuật)",
        "patterns": [
            (r"\bbảo mật (dữ liệu )?tuyệt đối\b", "Bảo mật dữ liệu nhiều tầng theo tiêu chuẩn cao"),
            (r"\bkhông thể bị hack\b", "Khả năng chống chịu tấn công mạng kiên cố"),
            (r"\bmiễn nhiễm (hoàn toàn|100%)\b", "Tăng cường sức đề kháng trước mã độc"),
        ]
    }
]

# Các cụm từ cần kiểm tra nguồn chứng minh (Substantiation Check)
CONDITIONAL_CLAIMS = [
    (r"\bsố (1|một)\b", "Cần kèm trích dẫn nguồn công bố hợp pháp (Theo số liệu của... giai đoạn...)"),
    (r"\bduy nhất\b", "Cần tài liệu chứng minh tính độc quyền hoặc duy nhất"),
    (r"\btốt nhất\b", "Cần có chứng nhận giải thưởng hoặc bình chọn của bên thứ ba độc lập"),
    (r"\bhàng đầu\b", "Cần có căn cứ đo lường thị phần hoặc đánh giá tổ chức uy tín"),
]

BINARY_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".pptx", ".zip", ".tar", ".gz", ".png", ".jpg", ".jpeg", ".mp4", ".mp3"}

def extract_text_safely(file_path: Path) -> Tuple[Optional[str], Optional[str]]:
    """
    Trích xuất text từ tệp an toàn. 
    Không bao giờ silently read binary formats as raw text.
    """
    ext = file_path.suffix.lower()

    if ext in {".txt", ".md", ".json", ".html", ".htm", ".js", ".ts", ".jsx", ".tsx", ".csv", ".typ"}:
        try:
            return file_path.read_text(encoding="utf-8", errors="replace"), None
        except Exception as e:
            return None, f"Lỗi đọc text file: {e}"

    if ext == ".docx":
        try:
            import docx
            doc = docx.Document(str(file_path))
            full_text = []
            for p in doc.paragraphs:
                full_text.append(p.text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        full_text.append(cell.text)
            return "\n".join(full_text), None
        except ImportError:
            return None, "File .docx yêu cầu thư viện 'python-docx'. Chạy: pip3 install python-docx"
        except Exception as e:
            return None, f"Lỗi bóc tách text từ file .docx: {e}"

    if ext == ".pdf":
        try:
            import fitz
            doc = fitz.open(str(file_path))
            full_text = []
            for page in doc:
                full_text.append(page.get_text())
            doc.close()
            return "\n".join(full_text), None
        except ImportError:
            return None, "File .pdf yêu cầu thư viện 'pymupdf'. Chạy: pip3 install pymupdf"
        except Exception as e:
            return None, f"Lỗi bóc tách text từ file .pdf: {e}"

    if ext in BINARY_EXTENSIONS:
        return None, f"Định dạng nhị phân '{ext}' không được hỗ trợ đọc trực tiếp. Vui lòng trích xuất text trước khi rà soát."

    # Fallback cho các extension lạ: thử đọc text với kiểm tra byte null
    try:
        raw_bytes = file_path.read_bytes()
        if b"\x00" in raw_bytes[:1024]:
            return None, f"File chứa ký tự nhị phân (binary content). Vui lòng cung cấp file văn bản thuần."
        return raw_bytes.decode("utf-8", errors="replace"), None
    except Exception as e:
        return None, f"Không thể đọc file: {e}"


def check_file(file_path, quiet=False, strict=False):
    p = Path(file_path)
    if not p.exists():
        if not quiet:
            print(f"{RED}Lỗi: File không tồn tại: {file_path}{RESET}")
        return {
            "status": "FAIL",
            "failures": [f"File không tồn tại: {file_path}"],
            "warnings": [],
            "artifacts": []
        }

    text, err = extract_text_safely(p)
    if err:
        if not quiet:
            print(f"{RED}❌ {err}{RESET}")
        return {
            "status": "BLOCKED",
            "failures": [err],
            "warnings": [],
            "artifacts": [str(p)]
        }

    lines = text.split("\n")
    violations = []
    warnings = []

    for idx, line in enumerate(lines, 1):
        line_clean = line.strip()
        line_lower = line_clean.lower()

        # 1. Quét vi phạm cấm tuyệt đối
        for cat in CLAIM_RULES:
            for pattern, suggestion in cat["patterns"]:
                match = re.search(pattern, line_lower)
                if match:
                    violations.append({
                        "line": idx,
                        "matched": match.group(0),
                        "snippet": line_clean[:120],
                        "category": cat["category"],
                        "severity": cat["severity"],
                        "suggestion": suggestion
                    })

        # 2. Quét các từ so sánh có điều kiện (Conditional Claims)
        for pattern, note in CONDITIONAL_CLAIMS:
            match = re.search(pattern, line_lower)
            if match:
                # Kiểm tra xem trong dòng hoặc đoạn có nguồn trích dẫn không
                has_source = any(w in line_lower for w in ["theo số liệu", "công bố", "khảo sát", "nguồn:", "báo cáo", "pyuru", "năm 20"])
                if not has_source:
                    warnings.append({
                        "line": idx,
                        "matched": match.group(0),
                        "snippet": line_clean[:120],
                        "note": note
                    })

    # Quyết định trạng thái
    if violations:
        status = "FAIL"
    elif warnings:
        status = "FAIL" if strict else "REQUIRES_EVIDENCE"
    else:
        status = "PASS"

    # In báo cáo
    if not quiet:
        print(f"\n{BOLD}=== PRELIMINARY CLAIM LINTER (RULE R5): {p.name} ==={RESET}")
        print(f"{BLUE}ℹ️  Lưu ý: Đây là bộ lọc từ khóa sơ bộ, không xác nhận tính đúng đắn pháp lý toàn diện.{RESET}\n")

    if violations:
        print(f"{RED}{BOLD}❌ PHÁT HIỆN {len(violations)} VI PHẠM OVER-CLAIM (CẤM TUYỆT ĐỐI THEO LUẬT R5):{RESET}")
        for v in violations:
            print(f"  {RED}• Dòng {v['line']}:{RESET} Phát hiện cụm từ cấm {YELLOW}'{v['matched']}'{RESET}")
            print(f"    - Phân loại: {v['category']} [{v['severity']}]")
            print(f"    - Trích đoạn: \"{v['snippet']}\"")
            print(f"    - {GREEN}Gợi ý thay thế hợp chuẩn:{RESET} \"{v['suggestion']}\"")
            print()

    if warnings and not quiet:
        print(f"{YELLOW}{BOLD}⚠️ CẢNH BÁO {len(warnings)} TUYÊN BỐ CẦN MINH CHỨNG (REQUIRES EVIDENCE):{RESET}")
        print(f"{YELLOW}   (Không được mặc định coi là PASS nếu chưa có số liệu chứng minh công bố){RESET}")
        for w in warnings:
            print(f"  {YELLOW}• Dòng {w['line']}:{RESET} Sử dụng từ so sánh '{w['matched']}' nhưng thiếu ghi chú nguồn:")
            print(f"    - Yêu cầu: {w['note']}")
            print(f"    - Trích đoạn: \"{w['snippet']}\"")
            print()

    if status == "PASS" and not quiet:
        print(f"{GREEN}✅ LINTER_PASS: Không phát hiện từ khóa cấm trong từ điển linter.{RESET}")
        print(f"{GREEN}   (Lưu ý: Không thay thế thẩm định pháp lý nội dung chuyên sâu).{RESET}\n")
    elif status == "REQUIRES_EVIDENCE" and not quiet:
        print(f"{YELLOW}⚠️  LINTER_REQUIRES_EVIDENCE: Không có từ cấm, nhưng có tuyên bố cần bằng chứng nguồn.{RESET}\n")

    return {
        "status": status,
        "violations_count": len(violations),
        "warnings_count": len(warnings),
        "failures": [f"Dòng {v['line']}: {v['matched']} ({v['category']})" for v in violations],
        "warnings": [f"Dòng {w['line']}: {w['matched']} ({w['note']})" for w in warnings],
        "artifacts": [str(p)]
    }

def main():
    parser = argparse.ArgumentParser(description="Preliminary Claim Linter (Rule R5)")
    parser.add_argument("--input", "-i", required=True, help="Đường dẫn file cần kiểm tra")
    parser.add_argument("--quiet", "-q", action="store_true", help="Chỉ in kết quả ngắn gọn")
    parser.add_argument("--strict", action="store_true", help="Bắt buộc fail cả khi có cảnh báo cần minh chứng")
    parser.add_argument("--json", action="store_true", help="Xuất kết quả JSON")
    args = parser.parse_args()

    result = check_file(args.input, args.quiet, args.strict)

    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))

    if result["status"] == "FAIL":
        sys.exit(1)
    elif result["status"] == "BLOCKED":
        sys.exit(2)
    elif result["status"] == "REQUIRES_EVIDENCE" and args.strict:
        sys.exit(1)
    else:
        sys.exit(0)

if __name__ == "__main__":
    main()

