#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
doc_ingest.py — Bộ chuyển đổi tài liệu DÙNG CHUNG của AIWF: mọi định dạng → Markdown + manifest.

Agent KHÔNG đọc trực tiếp tệp nhị phân (DOC/DOCX/XLS/XLSX/PPT/PPTX/PDF/EPUB/ODT/RTF/ảnh...).
Chạy công cụ này trước, rồi đọc <out>/source.md + <out>/manifest.json.

CLI (hợp đồng cố định — các skill gọi đúng như sau):
    .venv/bin/python scripts/doc_ingest.py <input> [<input2> ...] --out <dir> [--json] [--no-ocr] [--max-chars N]

Đầu ra cho mỗi tệp (nhiều tệp → mỗi tệp một thư mục con <dir>/<slug-tên-tệp>/, kèm <dir>/index.json):
    source.md      Markdown: tiêu đề, bảng dạng pipe table, ảnh ![](media/..),
                   PPTX: "## Slide N" + ghi chú diễn giả; XLSX: mỗi sheet một mục (bảng giá trị + bảng công thức);
                   EPUB: các chương theo thứ tự spine; PDF: marker "<!-- page N -->" cho từng trang.
    media/         Ảnh thật trích từ tài liệu.
    ocr_pages/     PNG các trang/ảnh KHÔNG có lớp chữ (cần OCR bằng thị giác) + JSON Apple Vision.
    manifest.json  {input, detected_format, converter_chain, pages, sheets, slides, chapters, chars, tables,
                    images, warnings[], pages_needing_ocr[], ocr_pages[], legacy_encoding, status, exit_code, message}

Mã thoát:
    0  ok           — đọc đầy đủ.
    3  partial      — có trang cần OCR bằng thị giác: đọc ảnh trong ocr_pages/ (bản nháp OCR trong source.md CHƯA kiểm chứng).
    2  unreadable   — PDF có mật khẩu, EPUB có DRM, tệp Office mã hoá, tệp rỗng/hỏng/không hỗ trợ (manifest.message nói cần hỏi người dùng gì).
    4  missing_dependency — thiếu thư viện/công cụ (manifest.message nói chính xác cần cài gì).
    Nhiều tệp: trả về mã "nặng" nhất (4 > 2 > 3 > 0).

Python API:
    from doc_ingest import ingest
    manifest = ingest("file.docx", "out_dir", ocr=True)       # không ném lỗi với tệp xấu; xem manifest["status"]

Nếu python3 hệ thống thiếu markitdown, script tự chạy lại bằng <repo>/.venv/bin/python.
Không gửi dữ liệu ra dịch vụ ngoài (không tải ảnh từ URL; OCR nháp chạy cục bộ: Apple Vision / tesseract).
"""
from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import io
import json
import os
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
import time
import traceback
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from urllib.parse import unquote

__version__ = "1.0.0"
__all__ = ["ingest", "detect_format", "detect_legacy", "tcvn3_to_unicode", "decode_bytes",
           "diacritic_damage", "choose_least_damaged", "find_soffice", "main"]

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
MAC_OCR_SWIFT = REPO / ".agents" / "skills" / "boc-tach-pdf" / "scripts" / "mac_ocr.swift"
# mac_ocr.swift tự khớp "vi-VN" → "vi-VT" (mã Apple). Hai lượt: Việt và Nhật (Vision không đọc cả hai trong một lượt).
VISION_PASSES = "vi-VN,en-US;ja-JP,en-US"

EXIT_OK, EXIT_UNREADABLE, EXIT_PARTIAL, EXIT_MISSING_DEP = 0, 2, 3, 4
STATUS_BY_CODE = {0: "ok", 3: "partial", 2: "unreadable", 4: "missing_dependency"}
SEVERITY = {0: 0, 3: 1, 2: 2, 4: 3}

REQUIRED_MODULES = ("markitdown", "mammoth", "bs4", "pymupdf", "pdfplumber", "pdfminer", "openpyxl", "pptx", "PIL")
PIP_NAMES = {"pymupdf": "pymupdf", "pptx": "python-pptx", "PIL": "Pillow", "bs4": "beautifulsoup4",
             "pdfminer": "pdfminer.six", "markitdown": "markitdown[docx,pdf,pptx,xlsx,xls,outlook]"}
REEXEC_ENV = "AIWF_DOC_INGEST_REEXEC"
BIG_TEXT_CHARS = 300_000

IMAGE_FORMATS = ("png", "jpeg", "tiff", "gif", "bmp", "webp", "heic")

MSG = {
    "missing_python": (
        "Thiếu thư viện Python: {mods}. Cài đặt: trong thư mục ai-workforce chạy `bash scripts/auto-setup.sh` "
        "(macOS/Linux) hoặc `setup.bat` (Windows); hoặc thủ công: `python3 -m venv .venv && "
        ".venv/bin/pip install -r requirements.txt` (Windows: `.venv\\Scripts\\pip install -r requirements.txt`). "
        "Sau đó chạy lại: `.venv/bin/python scripts/doc_ingest.py ...`."),
    "missing_soffice": (
        "Cần LibreOffice để đọc định dạng {fmt} (Word/Excel/PowerPoint đời cũ, OpenDocument, RTF). Cài: macOS "
        "`brew install --cask libreoffice`; Windows `winget install -e --id TheDocumentFoundation.LibreOffice`; "
        "Linux `sudo apt install libreoffice`. Đã cài ở chỗ khác thì đặt biến môi trường SOFFICE=<đường dẫn soffice>. "
        "Hoặc hỏi người dùng: mở tệp và Lưu thành {modern} rồi gửi lại."),
    "soffice_failed": (
        "LibreOffice không chuyển được tệp {name} sang {modern} (tệp có thể có mật khẩu, bị hỏng hoặc quá lớn). "
        "Hãy hỏi người dùng: mở tệp trong Word/Excel/PowerPoint, bỏ mật khẩu nếu có, Lưu thành {modern} rồi gửi lại."),
    "pdf_password": (
        "PDF được bảo vệ bằng mật khẩu nên không đọc được. Hãy hỏi người dùng: gửi bản PDF không đặt mật khẩu "
        "(mở tệp bằng mật khẩu → In → Lưu thành PDF) hoặc tự gỡ mật khẩu rồi gửi lại. KHÔNG đoán mật khẩu."),
    "office_password": (
        "Tệp Office được đặt mật khẩu (nội dung bị mã hoá) nên không đọc được. Hãy hỏi người dùng: mở tệp bằng mật khẩu, "
        "gỡ mật khẩu (File → Info → Protect → Encrypt with Password → để trống), lưu lại rồi gửi lại."),
    "epub_drm": (
        "EPUB có DRM (nội dung bị mã hoá: META-INF/encryption.xml hoặc rights.xml) nên không đọc được. Hãy hỏi người dùng "
        "bản không DRM (EPUB/PDF gốc mà họ có quyền sử dụng)."),
    "empty": "Tệp rỗng (0 byte). Hãy hỏi người dùng gửi lại tệp (có thể tải xuống chưa xong).",
    "not_found": "Không tìm thấy tệp: {path}. Hãy hỏi người dùng đường dẫn chính xác.",
    "is_dir": "Đường dẫn là thư mục, không phải tệp: {path}. Hãy truyền từng tệp (có thể nhiều tệp cùng lúc).",
    "unsupported": (
        "Không nhận dạng/không hỗ trợ định dạng tệp ({detail}). Hãy hỏi người dùng tệp gốc ở dạng DOCX, XLSX, PPTX, PDF, "
        "EPUB, HTML, CSV, TXT hoặc ảnh JPG/PNG."),
    "corrupt": (
        "Tệp bị hỏng hoặc không mở được ({detail}). Hãy hỏi người dùng gửi lại bản gốc (hoặc xuất lại sang DOCX/PDF)."),
    "no_content": (
        "Không trích được nội dung nào từ tệp ({detail}). Hãy hỏi người dùng xác nhận tệp có nội dung hoặc gửi bản khác."),
    "heic": ("Cần thư viện pillow-heif để đọc ảnh HEIC trên hệ điều hành này: `.venv/bin/pip install pillow-heif`. "
             "Hoặc hỏi người dùng xuất ảnh sang JPG/PNG."),
}
MODERN = {"docx": "DOCX (Word)", "xlsx": "XLSX (Excel)", "pptx": "PPTX (PowerPoint)", "png": "PNG"}
OCR_REASON_VI = {
    "no_text_layer": "trang scan, không có lớp chữ",
    "garbage_text_layer": "lớp chữ hỏng/mã hoá font, không đọc được",
    "mostly_image": "trang gần như toàn ảnh, rất ít chữ",
    "image_input": "tệp ảnh",
}
ENGINE_VI = {"apple_vision": "Apple Vision (vi + ja)", "tesseract": "tesseract"}


# ════════════════════════════════════════════════════════════════════════════
# Lỗi
# ════════════════════════════════════════════════════════════════════════════
class IngestError(Exception):
    code = EXIT_UNREADABLE

    def __init__(self, message: str, detail: str | None = None):
        super().__init__(message)
        self.message = message
        self.detail = detail


class Unreadable(IngestError):
    code = EXIT_UNREADABLE


class MissingDependency(IngestError):
    code = EXIT_MISSING_DEP


# ════════════════════════════════════════════════════════════════════════════
# Môi trường: kiểm tra thư viện, tự chạy lại bằng .venv
# ════════════════════════════════════════════════════════════════════════════
def _missing_modules(mods=REQUIRED_MODULES) -> list[str]:
    import importlib.util
    out = []
    for m in mods:
        try:
            if importlib.util.find_spec(m) is None:
                out.append(m)
        except (ImportError, ValueError):
            out.append(m)
    return out


def _venv_python() -> Path | None:
    for rel in ((".venv", "bin", "python"), (".venv", "Scripts", "python.exe")):
        p = REPO.joinpath(*rel)
        if p.exists():
            return p
    return None


def _running_in_repo_venv() -> bool:
    try:
        return Path(sys.prefix).resolve() == (REPO / ".venv").resolve()
    except OSError:
        return False


def _missing_python_message(missing: list[str]) -> str:
    mods = ", ".join(f"{m} (pip: {PIP_NAMES.get(m, m)})" for m in missing)
    return MSG["missing_python"].format(mods=mods)


def find_soffice() -> str | None:
    """LibreOffice: $SOFFICE → PATH → vị trí cài đặt mặc định (macOS/Windows)."""
    cands = [os.environ.get("SOFFICE"), shutil.which("soffice"), shutil.which("libreoffice"),
             "/Applications/LibreOffice.app/Contents/MacOS/soffice",
             str(Path.home() / "Applications/LibreOffice.app/Contents/MacOS/soffice"),
             r"C:\Program Files\LibreOffice\program\soffice.exe",
             r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]
    for c in cands:
        if c and Path(c).is_file():
            return c
    return None


def _uid() -> str:
    try:
        return str(os.getuid())
    except AttributeError:
        return os.environ.get("USERNAME", "user")


# ════════════════════════════════════════════════════════════════════════════
# Nhận dạng định dạng theo magic bytes (không tin phần mở rộng)
# ════════════════════════════════════════════════════════════════════════════
def detect_format(path) -> tuple[str, dict]:
    """Trả về (format, info). info có thể chứa 'encrypted': True, 'detail'."""
    p = Path(path)
    size = p.stat().st_size
    if size == 0:
        return "empty", {}
    with open(p, "rb") as f:
        head = f.read(8192)
    if b"%PDF-" in head[:1024]:
        return "pdf", {}
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return _ole_kind(p)
    if head.startswith(b"{\\rtf"):
        return "rtf", {}
    img = _sniff_image(head)
    if img:
        return img, {}
    if head.startswith(b"PK\x03\x04") or head.startswith(b"PK\x05\x06") or zipfile.is_zipfile(p):
        return _zip_kind(p)
    # Văn bản
    is_bom = head.startswith((b"\xef\xbb\xbf", b"\xff\xfe", b"\xfe\xff"))
    if not is_bom and b"\x00" in head and not _looks_utf16(head):
        return "binary", {"detail": "tệp nhị phân không rõ loại"}
    text_head, _, _ = decode_bytes(head)
    low = text_head.lstrip("\ufeff \t\r\n").lower()
    ext = p.suffix.lower()
    if low.startswith(("<!doctype html", "<html")) or "<html" in low[:600] or ext in (".html", ".htm", ".xhtml"):
        return "html", {}
    if ext in (".csv", ".tsv"):
        return "csv", {}
    if ext in (".md", ".markdown"):
        return "markdown", {}
    return "text", {}


def _sniff_image(head: bytes) -> str | None:
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if head[:3] == b"\xff\xd8\xff":
        return "jpeg"
    if head[:4] in (b"II*\x00", b"MM\x00*"):
        return "tiff"
    if head[:6] in (b"GIF87a", b"GIF89a"):
        return "gif"
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return "webp"
    if head[4:8] == b"ftyp" and head[8:12] in (b"heic", b"heix", b"hevc", b"heim", b"heis", b"mif1", b"msf1", b"avif"):
        return "heic"
    if head[:2] == b"BM" and len(head) >= 26 and head[6:10] == b"\x00\x00\x00\x00":
        return "bmp"
    return None


def _img_ext(data: bytes, fallback: str = "") -> str:
    kind = _sniff_image(data[:64])
    if kind:
        return {"jpeg": "jpg"}.get(kind, kind)
    if data[:4] == b"\x01\x00\x00\x00" and data[40:44] == b" EMF":
        return "emf"
    if data[:4] in (b"\xd7\xcd\xc6\x9a", b"\x01\x00\x09\x00"):
        return "wmf"
    if b"<svg" in data[:2048]:
        return "svg"
    return fallback


def _looks_utf16(head: bytes) -> bool:
    s = head[:4096]
    if len(s) < 8:
        return False
    zeros_even, zeros_odd = s[0::2].count(0), s[1::2].count(0)
    half = len(s) // 2
    return max(zeros_even, zeros_odd) > half * 0.3 and min(zeros_even, zeros_odd) < half * 0.05


def _zip_kind(p: Path) -> tuple[str, dict]:
    try:
        with zipfile.ZipFile(p) as z:
            names = set(z.namelist())
            mt = z.read("mimetype")[:120].decode("ascii", "ignore").strip() if "mimetype" in names else ""
    except (zipfile.BadZipFile, OSError, KeyError) as e:
        return "corrupt", {"detail": f"tệp nén hỏng: {e}"}
    if mt == "application/epub+zip" or ("META-INF/container.xml" in names and any(n.endswith(".opf") for n in names)):
        return "epub", {}
    if mt.startswith("application/vnd.oasis.opendocument."):
        sub = mt.rsplit(".", 1)[-1]
        if sub in ("text", "text-master", "text-template"):
            return "odt", {}
        if sub in ("spreadsheet", "spreadsheet-template"):
            return "ods", {}
        if sub in ("presentation", "presentation-template"):
            return "odp", {}
        return "unsupported", {"detail": f"OpenDocument loại {sub}"}
    if "word/document.xml" in names:
        return "docx", {}
    if "ppt/presentation.xml" in names:
        return "pptx", {}
    if "xl/workbook.xml" in names:
        return "xlsx", {}
    if "xl/workbook.bin" in names:
        return "xlsb", {}
    return "zip", {"detail": "tệp ZIP thường (không phải tài liệu) — hãy giải nén và gửi từng tệp"}


def _ole_kind(p: Path) -> tuple[str, dict]:
    try:
        import olefile
    except ImportError:
        ext = p.suffix.lower().lstrip(".")
        return ({"doc": "doc", "dot": "doc", "xls": "xls", "xlt": "xls", "ppt": "ppt", "pps": "ppt", "msg": "msg"}
                .get(ext, "ole"), {"detail": "không có olefile — đoán theo phần mở rộng"})
    try:
        with olefile.OleFileIO(str(p)) as ole:
            streams = {"/".join(s) for s in ole.listdir()}
            if "EncryptionInfo" in streams and "EncryptedPackage" in streams:
                return "ooxml_encrypted", {"encrypted": True}
            if "WordDocument" in streams:
                fib = ole.openstream("WordDocument").read(12)
                enc = len(fib) >= 12 and bool(int.from_bytes(fib[10:12], "little") & 0x0100)
                return "doc", {"encrypted": enc}
            for wb in ("Workbook", "Book"):
                if wb in streams:
                    return "xls", {"encrypted": _xls_has_filepass(ole.openstream(wb).read(8192))}
            if "PowerPoint Document" in streams:
                return "ppt", {"encrypted": "EncryptedSummary" in streams}
            if any(s.startswith("__substg1.0_") for s in streams):
                return "msg", {}
    except Exception as e:  # noqa: BLE001
        return "corrupt", {"detail": f"tệp OLE hỏng: {e}"}
    return "ole", {"detail": "tệp OLE2 không phải Word/Excel/PowerPoint/Outlook"}


def _xls_has_filepass(data: bytes) -> bool:
    i = 0
    for _ in range(200):
        if i + 4 > len(data):
            break
        rtype = int.from_bytes(data[i:i + 2], "little")
        size = int.from_bytes(data[i + 2:i + 4], "little")
        if rtype == 0x002F:
            return True
        if rtype == 0x0A and i > 0:  # EOF của globals
            break
        i += 4 + size
    return False


# ════════════════════════════════════════════════════════════════════════════
# Mã hoá văn bản: UTF-8 → UTF-16 BOM → TCVN3/VNI → CP1258; phát hiện + chuyển TCVN3
# ════════════════════════════════════════════════════════════════════════════
# TCVN3 (ABC, font .VnTime/.VnArial…): byte → chữ Unicode. Bản .VnTimeH (chữ hoa) dùng cùng mã chữ thường.
_TCVN3 = {
    0xB8: "á", 0xB5: "à", 0xB6: "ả", 0xB7: "ã", 0xB9: "ạ",
    0xA8: "ă", 0xBE: "ắ", 0xBB: "ằ", 0xBC: "ẳ", 0xBD: "ẵ", 0xC6: "ặ",
    0xA9: "â", 0xCA: "ấ", 0xC7: "ầ", 0xC8: "ẩ", 0xC9: "ẫ", 0xCB: "ậ",
    0xAE: "đ",
    0xD0: "é", 0xCC: "è", 0xCE: "ẻ", 0xCF: "ẽ", 0xD1: "ẹ",
    0xAA: "ê", 0xD5: "ế", 0xD2: "ề", 0xD3: "ể", 0xD4: "ễ", 0xD6: "ệ",
    0xDD: "í", 0xD7: "ì", 0xD8: "ỉ", 0xDC: "ĩ", 0xDE: "ị",
    0xE3: "ó", 0xDF: "ò", 0xE1: "ỏ", 0xE2: "õ", 0xE4: "ọ",
    0xAB: "ô", 0xE8: "ố", 0xE5: "ồ", 0xE6: "ổ", 0xE7: "ỗ", 0xE9: "ộ",
    0xAC: "ơ", 0xED: "ớ", 0xEA: "ờ", 0xEB: "ở", 0xEC: "ỡ", 0xEE: "ợ",
    0xF3: "ú", 0xEF: "ù", 0xF1: "ủ", 0xF2: "ũ", 0xF4: "ụ",
    0xAD: "ư", 0xF8: "ứ", 0xF5: "ừ", 0xF6: "ử", 0xF7: "ữ", 0xF9: "ự",
    0xFD: "ý", 0xFA: "ỳ", 0xFB: "ỷ", 0xFC: "ỹ", 0xFE: "ỵ",
    0xA1: "Ă", 0xA2: "Â", 0xA3: "Ê", 0xA4: "Ô", 0xA5: "Ơ", 0xA6: "Ư", 0xA7: "Đ",
}
_TCVN3_TRANS = {k: v for k, v in _TCVN3.items()}
TCVN3_CHARS = frozenset(chr(k) for k in _TCVN3)
# Ký tự 0xA1–0xBE: trong văn bản phương Tây là ký hiệu hiếm (©®µ¸·…), trong TCVN3 là chữ Việt rất thường gặp.
# Bỏ « » (ngoặc kép tiếng Pháp) để tránh nhận nhầm.
TCVN3_STRONG = frozenset(chr(k) for k in _TCVN3 if 0xA1 <= k <= 0xBE and k not in (0xAB, 0xBB))
# Chữ chỉ có trong tiếng Việt Unicode (không nằm trong Latin-1) → bằng chứng văn bản đã là Unicode.
VI_UNICODE_RE = re.compile("[\u1EA0-\u1EF9\u0102\u0103\u0110\u0111\u01A0\u01A1\u01AF\u01B0\u0128\u0129\u0168\u0169]")
# VNI-Windows: nguyên âm + byte dấu (vd "Vieät", "cuûa", "ñöôïc").
VNI_RE = re.compile("[aeouyAEOUYôöÔÖ][ùøûõïâáàåãäêéèúüë]")


def detect_legacy(text: str, hint: str | None = None) -> str | None:
    """Trả về 'tcvn3', 'vni' hoặc None. hint: 'tcvn3'/'vni' nếu biết font (.VnTime / VNI-Times)."""
    if not text:
        return None
    sample = text if len(text) <= 400_000 else text[:200_000] + "\n" + text[-200_000:]
    toks = [t for t in re.findall(r"\S+", sample) if any(c.isalpha() for c in t)]
    if not toks:
        return None
    uni_vi = len(VI_UNICODE_RE.findall(sample))
    strong = sum(1 for c in sample if c in TCVN3_STRONG)
    tc_tok = sum(1 for t in toks if any(c in TCVN3_CHARS for c in t))
    ratio = tc_tok / len(toks)
    if uni_vi * 4 <= max(strong, tc_tok // 2):
        if hint == "tcvn3" and ratio >= 0.1 and tc_tok >= 1:
            return "tcvn3"
        if strong >= 3 and ratio >= 0.2:
            return "tcvn3"
        if strong >= 2 and ratio >= 0.35:
            return "tcvn3"
    vni = len(VNI_RE.findall(sample))
    if uni_vi * 4 <= vni and ((vni >= 5 and vni / len(toks) >= 0.1) or (hint == "vni" and vni >= 1)):
        return "vni"
    return None


def _tcvn3_token(t: str, line_upper: bool) -> str:
    if not any(c in TCVN3_CHARS for c in t):
        return t
    u = t.translate(_TCVN3_TRANS)
    ascii_letters = [c for c in t if c.isascii() and c.isalpha()]
    # Font .VnTimeH (chữ HOA) dùng chung mã với chữ thường: viết hoa khi cả dòng là chữ HOA,
    # hoặc token có ≥2 chữ Latin đều HOA ("CéNG" → "CỘNG"); "Tù" (1 chữ hoa đầu từ) giữ "Tự".
    if line_upper or (len(ascii_letters) >= 2 and all(c.isupper() for c in ascii_letters)):
        u = u.upper()
    return u


def tcvn3_to_unicode(text: str) -> tuple[str, int]:
    """Chuyển TCVN3 → Unicode theo từng dòng (bỏ qua dòng đã có chữ Việt Unicode). Trả (văn bản, số dòng đã chuyển)."""
    out, changed = [], 0
    for line in text.split("\n"):
        if VI_UNICODE_RE.search(line) or not any(c in TCVN3_CHARS for c in line):
            out.append(line)
            continue
        alpha = [t for t in re.findall(r"\S+", line) if any(c.isalpha() for c in t)]
        tc = sum(1 for t in alpha if any(c in TCVN3_CHARS for c in t))
        if not (any(c in TCVN3_STRONG for c in line) or (alpha and tc / len(alpha) >= 0.3)):
            out.append(line)
            continue
        letters = [c for c in line if c.isascii() and c.isalpha()]
        line_upper = len(letters) >= 3 and all(c.isupper() for c in letters)
        out.append(re.sub(r"\S+", lambda m: _tcvn3_token(m.group(0), line_upper), line))
        changed += 1
    return "\n".join(out), changed


def decode_bytes(b: bytes) -> tuple[str, str, str | None]:
    """UTF-8 (BOM) → UTF-16 (BOM/không BOM) → UTF-8 → TCVN3 → VNI → CP1258. Trả (text, encoding, legacy)."""
    if b.startswith(b"\xef\xbb\xbf"):
        return b[3:].decode("utf-8", "replace"), "utf-8-sig", None
    if b.startswith((b"\xff\xfe", b"\xfe\xff")):
        return b.decode("utf-16", "replace"), "utf-16", None
    if _looks_utf16(b[:4096]):
        s = b[:4096]
        enc = "utf-16-le" if s[1::2].count(0) > s[0::2].count(0) else "utf-16-be"
        try:
            return b.decode(enc), enc, None
        except UnicodeDecodeError:
            pass
    try:
        return b.decode("utf-8"), "utf-8", None
    except UnicodeDecodeError as e:
        # Tệp UTF-8 bị cắt cụt giữa một ký tự ở cuối
        if e.start >= len(b) - 3:
            try:
                return b[:e.start].decode("utf-8"), "utf-8", None
            except UnicodeDecodeError:
                pass
    latin = b.decode("latin-1")
    if detect_legacy(latin) == "tcvn3":
        return latin, "tcvn3", "tcvn3"
    w1252 = b.decode("cp1252", "replace")
    if detect_legacy(w1252) == "vni":
        return w1252, "vni-windows", "vni"
    return b.decode("cp1258", "replace"), "cp1258", "cp1258"


# ════════════════════════════════════════════════════════════════════════════
# Kiểm tra dấu tiếng Việt bị "rơi" (PDF dùng font dự phòng → dấu tách khỏi chữ)
# ════════════════════════════════════════════════════════════════════════════
def _has_vowel(tok: str) -> bool:
    base = unicodedata.normalize("NFD", tok.lower())
    return any(c in "aeiouy" for c in base)


def diacritic_damage(text: str) -> float:
    """Tỷ lệ "mảnh âm tiết" do rơi dấu: khi nguyên âm có dấu bị tách đi nơi khác, phần còn lại của chữ
    là cụm phụ âm không nguyên âm ('Ngh đ nh' ← 'Nghị định', 'th c' ← 'thức').
    Văn bản tiếng Việt bình thường < 0.02; văn bản rơi dấu thường > 0.08."""
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)|<!--.*?-->", " ", text, flags=re.S)
    toks = re.findall(r"[^\W\d_]+", text)
    vi = sum(1 for c in text if unicodedata.normalize("NFD", c) != c or c in "đĐ")
    if len(toks) < 20 or vi < 5:
        return 0.0
    bad = 0
    for t in toks:
        # viết tắt HOA (TP, NĐ, BHXH) và chữ đơn ('đ' tiền tệ, 'a)' mục liệt kê) không tính
        if not _has_vowel(t) and 2 <= len(t) <= 4 and not t.isupper():
            bad += 1
    return bad / len(toks)


def choose_least_damaged(primary: tuple[str, str], candidates: list[tuple[str, str]]) -> tuple[str, str, float]:
    """primary/candidates: (tên, văn bản). Giữ primary trừ khi nó vỡ dấu rõ rệt và có ứng viên tốt hơn nhiều."""
    p_score = diacritic_damage(primary[1])
    best = (primary[0], primary[1], p_score)
    if p_score < 0.06:
        return best
    for name, txt in candidates:
        if not txt or len(txt.strip()) < len(primary[1].strip()) * 0.5:
            continue
        s = diacritic_damage(txt)
        if s < best[2] * 0.5:
            best = (name, txt, s)
    return best


# ════════════════════════════════════════════════════════════════════════════
# Tiện ích Markdown
# ════════════════════════════════════════════════════════════════════════════
def _cell(x) -> str:
    s = "" if x is None else str(x)
    s = s.replace("\r\n", "\n").replace("\r", "\n").strip().replace("\n", "<br>")
    return s.replace("|", "\\|")


def md_table(rows, escape: bool = True) -> str:
    rows = [[(_cell(c) if escape else ("" if c is None else str(c))) for c in r] for r in rows if r is not None]
    rows = [r for r in rows if any(x.strip() for x in r)]
    if not rows:
        return ""
    n = max(len(r) for r in rows)
    rows = [r + [""] * (n - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "|".join(["---"] * n) + "|"]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def _space_tables(text: str) -> str:
    """Thêm dòng trống trước/sau khối bảng pipe để Markdown hiển thị đúng."""
    lines, out = text.split("\n"), []
    for i, ln in enumerate(lines):
        is_t = ln.lstrip().startswith("|")
        prev_t = i > 0 and lines[i - 1].lstrip().startswith("|")
        if is_t and not prev_t and out and out[-1].strip():
            out.append("")
        if not is_t and prev_t and ln.strip():
            out.append("")
        out.append(ln)
    return "\n".join(out)


TABLE_SEP_RE = re.compile(r"(?m)^\s*\|(?:\s*:?-{3,}:?\s*\|)+\s*$")


def _count_tables(md: str) -> int:
    return len(TABLE_SEP_RE.findall(md))


def slugify(name: str) -> str:
    s = name.replace("đ", "d").replace("Đ", "D")
    s = "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))
    s = re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-").lower()
    return s[:80] or "file"


def _oneline(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def _safe_xml(data: bytes):
    try:
        from defusedxml import ElementTree as DET  # type: ignore
        return DET.fromstring(data)
    except ImportError:
        return ET.fromstring(data)


# ════════════════════════════════════════════════════════════════════════════
# Ngữ cảnh chuyển đổi
# ════════════════════════════════════════════════════════════════════════════
class Ctx:
    def __init__(self, src: Path, out: Path, man: dict, ocr: bool):
        self.src, self.out, self.man, self.ocr = src, out, man, ocr
        self.tmp = Path(tempfile.mkdtemp(prefix="aiwf_ingest_"))
        self.font_hint: str | None = None
        self._media_by_hash: dict[str, str] = {}
        self._names: set[str] = set()
        self.media_limit = 2000
        self.soffice_runs = 0

    def warn(self, msg: str):
        if msg not in self.man["warnings"]:
            self.man["warnings"].append(msg)

    def chain(self, step: str):
        self.man["converter_chain"].append(step)

    def save_media(self, data: bytes, ext: str, stem: str) -> str | None:
        if not data:
            return None
        h = hashlib.sha1(data).hexdigest()
        if h in self._media_by_hash:
            return self._media_by_hash[h]
        if len(self._media_by_hash) >= self.media_limit:
            self.warn(f"Quá {self.media_limit} ảnh: chỉ trích {self.media_limit} ảnh đầu tiên vào media/.")
            return None
        ext = (ext or "").lower().lstrip(".")
        ext = {"jpeg": "jpg", "x-emf": "emf", "x-wmf": "wmf", "svg+xml": "svg", "tif": "tiff", "jpx": "jp2"}.get(ext, ext)
        if not re.fullmatch(r"[a-z0-9]{2,5}", ext or ""):
            ext = _img_ext(data, "bin")
        media = self.out / "media"
        media.mkdir(parents=True, exist_ok=True)
        base = re.sub(r"[^A-Za-z0-9_.-]+", "_", stem).strip("_.") or "image"
        name, k = f"{base}.{ext}", 1
        while name in self._names:
            k += 1
            name = f"{base}_{k}.{ext}"
        (media / name).write_bytes(data)
        self._names.add(name)
        rel = f"media/{name}"
        self._media_by_hash[h] = rel
        return rel

    @property
    def image_count(self) -> int:
        return len(self._media_by_hash)

    def media_paths(self) -> list[str]:
        return list(self._media_by_hash.values())

    def cleanup(self):
        shutil.rmtree(self.tmp, ignore_errors=True)


# ════════════════════════════════════════════════════════════════════════════
# LibreOffice headless (profile riêng, không đụng LibreOffice người dùng đang mở)
# ════════════════════════════════════════════════════════════════════════════
SOFFICE_FILTERS = {"docx": "docx:MS Word 2007 XML", "xlsx": "xlsx:Calc MS Excel 2007 XML",
                   "pptx": "pptx:Impress MS PowerPoint 2007 XML", "png": "png"}


def soffice_convert(ctx: Ctx, src: Path, target: str, fmt_label: str, input_ext: str | None = None) -> Path:
    exe = find_soffice()
    if not exe:
        raise MissingDependency(MSG["missing_soffice"].format(fmt=fmt_label.upper(), modern=MODERN.get(target, target)))
    ctx.soffice_runs += 1
    work = ctx.tmp / f"lo{ctx.soffice_runs}"
    (work / "in").mkdir(parents=True)
    outdir = work / "out"
    inp = work / "in" / ("input" + (input_ext or src.suffix.lower() or ".bin"))
    shutil.copyfile(src, inp)
    profiles = [Path(tempfile.gettempdir()) / f"aiwf_lo_profile_{_uid()}", work / "profile"]
    last = ""
    for prof in profiles:
        cmd = [exe, f"-env:UserInstallation={prof.as_uri()}", "--headless", "--invisible", "--nologo",
               "--norestore", "--nolockcheck", "--nodefault", "--convert-to", SOFFICE_FILTERS.get(target, target),
               "--outdir", str(outdir), str(inp)]
        try:
            r = subprocess.run(cmd, capture_output=True, text=True, timeout=240, stdin=subprocess.DEVNULL)
            last = (r.stderr or r.stdout or "").strip()[-400:]
        except subprocess.TimeoutExpired:
            last = "LibreOffice quá thời gian (240s)"
            continue
        except OSError as e:
            last = str(e)
            continue
        res = outdir / f"input.{target}"
        if res.exists() and res.stat().st_size > 0:
            return res
    raise Unreadable(MSG["soffice_failed"].format(name=ctx.src.name, modern=MODERN.get(target, target)), detail=last)


# ════════════════════════════════════════════════════════════════════════════
# OCR nháp: Apple Vision (mac_ocr.swift của boc-tach-pdf) → tesseract
# ════════════════════════════════════════════════════════════════════════════
CJK_RE = re.compile(r"[\u3040-\u30ff\u3400-\u4dbf\u4e00-\u9fff\uff66-\uff9f]")


def _mac_ocr_cmd(ctx: Ctx) -> list[str] | None:
    if sys.platform != "darwin" or not MAC_OCR_SWIFT.exists():
        return None
    src = MAC_OCR_SWIFT
    prebuilt = src.with_suffix("")  # boc-tach-pdf có thể đã biên dịch sẵn (gitignored)
    try:
        if prebuilt.is_file() and os.access(prebuilt, os.X_OK) and prebuilt.stat().st_mtime >= src.stat().st_mtime:
            return [str(prebuilt)]
    except OSError:
        pass
    digest = hashlib.sha1(src.read_bytes()).hexdigest()[:12]
    cache_dir = Path(tempfile.gettempdir()) / f"aiwf_doc_ingest_{_uid()}"
    cached = cache_dir / f"mac_ocr_{digest}"
    if cached.is_file() and os.access(cached, os.X_OK):
        return [str(cached)]
    swiftc = shutil.which("swiftc")
    if swiftc:
        cache_dir.mkdir(parents=True, exist_ok=True)
        tmp_bin = cache_dir / f".mac_ocr_{digest}_{os.getpid()}"
        try:
            r = subprocess.run([swiftc, "-O", str(src), "-o", str(tmp_bin)], capture_output=True, text=True, timeout=600)
            if r.returncode == 0 and tmp_bin.exists():
                os.replace(tmp_bin, cached)
                return [str(cached)]
        except (subprocess.TimeoutExpired, OSError):
            pass
        finally:
            tmp_bin.unlink(missing_ok=True)
    swift = shutil.which("swift")
    return [swift, str(src)] if swift else None


def _bbox(quad) -> tuple[float, float, float, float]:
    xs = [p[0] for p in quad]
    ys = [p[1] for p in quad]
    return min(xs), min(ys), max(xs), max(ys)


def _overlap(a, b) -> float:
    ix = max(0.0, min(a[2], b[2]) - max(a[0], b[0]))
    iy = max(0.0, min(a[3], b[3]) - max(a[1], b[1]))
    inter = ix * iy
    small = min((a[2] - a[0]) * (a[3] - a[1]), (b[2] - b[0]) * (b[3] - b[1])) or 1.0
    return inter / small


def _layout_lines(lines: list[list]) -> str:
    """lines: [[bbox, text]] → văn bản theo hàng (các ô cùng hàng nối bằng ' | ')."""
    if not lines:
        return ""
    heights = sorted(b[3] - b[1] for b, _ in lines)
    tol = max(2.0, heights[len(heights) // 2] * 0.5)
    lines = sorted(lines, key=lambda l: ((l[0][1] + l[0][3]) / 2, l[0][0]))
    rows: list[list] = []
    for b, t in lines:
        cy = (b[1] + b[3]) / 2
        if rows and abs(rows[-1][0] - cy) <= tol:
            rows[-1][1].append((b[0], t))
        else:
            rows.append([cy, [(b[0], t)]])
    return "\n".join(" | ".join(t for _, t in sorted(cells)) for _, cells in rows)


def merge_vision_passes(d: dict) -> str:
    passes = d.get("passes") or [d]

    def obs(i):
        if i is None or i >= len(passes):
            return []
        return [o for o in (passes[i].get("observations") or []) if (o.get("text") or "").strip() and o.get("quad")]

    langs = [[str(x) for x in (p.get("languages") or [])] for p in passes]
    vi_i = next((i for i, ls in enumerate(langs) if any(x.startswith("vi") for x in ls)), 0)
    ja_i = next((i for i, ls in enumerate(langs) if any(x.startswith("ja") for x in ls)), None)
    lines = [[_bbox(o["quad"]), o["text"].strip()] for o in obs(vi_i)]
    if ja_i is not None and ja_i != vi_i:
        for o in obs(ja_i):
            if not CJK_RE.search(o["text"]):
                continue
            jb = _bbox(o["quad"])
            hits = [ln for ln in lines if ln[1] and _overlap(ln[0], jb) > 0.5]
            for ln in hits:
                if not CJK_RE.search(ln[1]):
                    ln[1] = None
            lines.append([jb, o["text"].strip()])
        lines = [ln for ln in lines if ln[1]]
    return _layout_lines(lines)


def _vision_drafts(ctx: Ctx, pngs: list[Path]) -> dict:
    cmd = _mac_ocr_cmd(ctx)
    if not cmd:
        return {}
    res = {}
    outdir = pngs[0].parent
    for i in range(0, len(pngs), 20):
        chunk = pngs[i:i + 20]
        try:
            subprocess.run(cmd + [str(outdir), *map(str, chunk), "--passes", VISION_PASSES],
                           capture_output=True, text=True, timeout=180 + 30 * len(chunk))
        except (subprocess.TimeoutExpired, OSError) as e:
            ctx.warn(f"Apple Vision lỗi/quá thời gian ({type(e).__name__}); các trang còn lại không có bản nháp OCR.")
            break
        for png in chunk:
            j = outdir / f"{png.name}.vision.json"
            if j.exists():
                try:
                    res[png.name] = (merge_vision_passes(json.loads(j.read_text("utf-8"))), "apple_vision")
                except (ValueError, KeyError, TypeError):
                    pass
    return res


def _tesseract_drafts(ctx: Ctx, pngs: list[Path]) -> dict:
    exe = shutil.which("tesseract")
    if not exe:
        return {}
    try:
        avail = subprocess.run([exe, "--list-langs"], capture_output=True, text=True, timeout=30).stdout.split()
    except (subprocess.TimeoutExpired, OSError):
        return {}
    langs = [x for x in ("vie", "jpn", "eng") if x in avail] or ["eng"]
    if "vie" not in langs:
        ctx.warn("tesseract thiếu gói tiếng Việt (vie): bản nháp OCR sẽ mất dấu. Cài: brew install tesseract-lang.")
    res = {}
    for png in pngs:
        try:
            r = subprocess.run([exe, str(png), "stdout", "-l", "+".join(langs)], capture_output=True, text=True, timeout=240)
            res[png.name] = (r.stdout.strip(), "tesseract")
        except (subprocess.TimeoutExpired, OSError):
            continue
    return res


def ocr_drafts(ctx: Ctx, pngs: list[Path]) -> dict:
    if not ctx.ocr or not pngs:
        return {}
    res = _vision_drafts(ctx, pngs)
    rest = [p for p in pngs if p.name not in res]
    if rest:
        res.update(_tesseract_drafts(ctx, rest))
    return res


def _ocr_blocks(ctx: Ctx, items: list[tuple[int, Path, str, str]]) -> dict[int, str]:
    """items: (số trang, png, lý do, chữ một phần) → {trang: khối markdown}. Ghi manifest."""
    drafts = ocr_drafts(ctx, [png for _, png, _, _ in items])
    blocks = {}
    for page_no, png, reason, extra in items:
        rel = f"ocr_pages/{png.name}"
        text, engine = drafts.get(png.name, ("", None))
        ctx.man["pages_needing_ocr"].append(page_no)
        ctx.man["ocr_pages"].append({"page": page_no, "image": rel, "reason": reason,
                                     "draft_engine": engine, "draft_chars": len(text)})
        lines = [f"> [!WARNING] Trang {page_no} cần OCR bằng thị giác ({OCR_REASON_VI.get(reason, reason)}). "
                 f"Đọc ảnh `{rel}` để lấy nội dung chính xác."]
        if engine and text:
            lines.append(f"> Bản nháp OCR {ENGINE_VI.get(engine, engine)} bên dưới CHƯA KIỂM CHỨNG — chỉ để đối chiếu.")
        else:
            lines.append("> Không có bản nháp OCR (không có Apple Vision/tesseract, ảnh không có chữ, hoặc đã tắt bằng --no-ocr).")
        if extra.strip():
            lines += ["", "<!-- text-layer (một phần) -->", extra.strip()]
        if text:
            lines += ["", f"<!-- ocr-draft:start page={page_no} engine={engine} -->", text, "<!-- ocr-draft:end -->"]
        blocks[page_no] = "\n".join(lines)
    if items:
        ctx.chain("ocr-draft:" + ("+".join(sorted({e for _, e in drafts.values()})) or "none"))
    return blocks


# ════════════════════════════════════════════════════════════════════════════
# HTML → Markdown (dùng HtmlConverter của markitdown) + ảnh + bảng
# ════════════════════════════════════════════════════════════════════════════
def _html_to_md(html_text: str) -> str:
    from markitdown.converters import HtmlConverter
    return HtmlConverter().convert_string(html_text).markdown or ""


def _prep_tables(soup):
    """Hàng đầu làm tiêu đề bảng (tránh hàng tiêu đề rỗng) và thoát ký tự '|' trong ô."""
    for table in soup.find_all("table"):
        first = table.find("tr")
        if first is not None and table.find("th") is None:
            for td in first.find_all("td", recursive=False):
                td.name = "th"
        for cell in table.find_all(["td", "th"]):
            for s in list(cell.find_all(string=True)):
                if "|" in s:
                    s.replace_with(s.replace("|", "\\|"))


def _data_uri(src: str):
    m = re.match(r"data:([\w/+.-]+)?(;base64)?,(.*)", src, re.S)
    if not m:
        return None, None
    try:
        data = base64.b64decode(m.group(3)) if m.group(2) else unquote(m.group(3)).encode()
    except (ValueError, TypeError):
        return None, None
    ext = (m.group(1) or "").split("/")[-1]
    return data, _img_ext(data, ext)


def _local_image(base: Path, src: str):
    """Ảnh cục bộ cạnh tệp HTML/MD. KHÔNG tải URL từ mạng. Chỉ nhận tệp thật sự là ảnh."""
    if src.startswith("data:"):
        return _data_uri(src)
    if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", src) and not src.lower().startswith("file:"):
        return None, None
    path = unquote(src.split("#", 1)[0].split("?", 1)[0])
    if path.lower().startswith("file://"):
        path = path[7:]
    if not path:
        return None, None
    fp = Path(path) if os.path.isabs(path) else (base / path)
    try:
        if not fp.is_file() or fp.stat().st_size > 50 * 1024 * 1024:
            return None, None
        data = fp.read_bytes()
    except OSError:
        return None, None
    ext = _img_ext(data, "")
    return (data, ext) if ext else (None, None)


def _rewrite_html_images(ctx: Ctx, soup, resolver, stem: str) -> int:
    n = 0
    for tag in soup.find_all(["img", "image"]):
        attr = "src" if tag.name == "img" else next((a for a in ("xlink:href", "href") if tag.get(a)), None)
        if not attr or not tag.get(attr):
            continue
        data, ext = resolver(tag.get(attr))
        if not data:
            continue
        rel = ctx.save_media(data, ext, stem)
        if not rel:
            continue
        n += 1
        if tag.name == "img":
            tag["src"] = rel
        else:
            tag.replace_with(soup.new_tag("img", src=rel, alt=tag.get("alt", "")))
    return n


def _decode_html(raw: bytes) -> str:
    m = re.search(rb"""charset\s*=\s*["']?([A-Za-z0-9_\-]+)""", raw[:4096], re.I)
    if m:
        enc = m.group(1).decode("ascii", "ignore")
        try:
            if enc.lower().replace("_", "-") not in ("utf-8", "utf8"):
                return raw.decode(enc, "replace")
        except LookupError:
            pass
    return decode_bytes(raw)[0]


# ════════════════════════════════════════════════════════════════════════════
# Biểu đồ / SmartArt (OOXML dùng chung cho DOCX, XLSX, PPTX)
# ════════════════════════════════════════════════════════════════════════════
C_NS = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
DGM_NS = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def chart_xml_to_md(blob: bytes) -> str:
    try:
        root = _safe_xml(blob)
    except ET.ParseError:
        return ""
    title_el = root.find(f".//{C_NS}title")
    title = _oneline(" ".join(t.text or "" for t in title_el.iter(f"{A_NS}t"))) if title_el is not None else ""
    plot_area = root.find(f".//{C_NS}plotArea")
    blocks = []
    for plot in (list(plot_area) if plot_area is not None else []):
        tag = plot.tag.split("}")[-1]
        if not tag.endswith("Chart"):
            continue
        series = []
        for k, ser in enumerate(plot.findall(f"{C_NS}ser"), 1):
            tx = ser.find(f"{C_NS}tx")
            name = _oneline(" ".join((v.text or "") for v in tx.iter() if v.tag in (f"{C_NS}v", f"{A_NS}t"))) if tx is not None else ""
            cat = ser.find(f"{C_NS}cat")
            if cat is None:
                cat = ser.find(f"{C_NS}xVal")
            val = ser.find(f"{C_NS}val")
            if val is None:
                val = ser.find(f"{C_NS}yVal")

            def pts(el):
                d = {}
                if el is None:
                    return d
                for pt in el.iter(f"{C_NS}pt"):
                    v = pt.find(f"{C_NS}v")
                    d.setdefault(int(pt.get("idx", "0")), (v.text or "") if v is not None else "")
                return d
            series.append((name or f"Chuỗi {k}", pts(cat), pts(val)))
        if not series:
            continue
        idxs = sorted(set().union(*[set(c) | set(v) for _, c, v in series]))
        rows = [["Danh mục"] + [s[0] for s in series]]
        for i in idxs:
            label = next((c[i] for _, c, _ in series if i in c), str(i + 1))
            rows.append([label] + [v.get(i, "") for _, _, v in series])
        tbl = md_table(rows) if len(rows) > 1 else "_(không có dữ liệu đệm)_"
        blocks.append(f"**Biểu đồ ({tag}){': ' + title if title else ''}**\n\n{tbl}")
    return "\n\n".join(blocks)


def smartart_texts(blob: bytes) -> list[str]:
    try:
        root = _safe_xml(blob)
    except ET.ParseError:
        return []
    return [_oneline(t.text) for t in root.iter(f"{A_NS}t") if t.text and t.text.strip()]


def _zip_extras(ctx: Ctx, p: Path, chart_prefix: str, media_prefix: str | None, diagram_prefix: str | None) -> str:
    """Biểu đồ/SmartArt/ảnh mà bộ chuyển đổi chính không đưa vào luồng chữ."""
    out = []
    try:
        with zipfile.ZipFile(p) as z:
            names = sorted(z.namelist(), key=lambda n: [int(x) if x.isdigit() else x for x in re.split(r"(\d+)", n)])
            charts = [n for n in names if n.startswith(chart_prefix) and re.search(r"/chart\d*\.xml$", n)]
            cmd = [chart_xml_to_md(z.read(n)) for n in charts]
            cmd = [c for c in cmd if c]
            if cmd:
                out.append("## Biểu đồ trong tệp\n\n" + "\n\n".join(cmd))
            if diagram_prefix:
                for n in names:
                    if n.startswith(diagram_prefix) and re.search(r"/data\d*\.xml$", n):
                        texts = smartart_texts(z.read(n))
                        if texts:
                            out.append("## SmartArt\n\n" + "\n".join(f"- {t}" for t in texts))
            if media_prefix:
                refs = []
                for n in names:
                    if n.startswith(media_prefix) and not n.endswith("/"):
                        data = z.read(n)
                        rel = ctx.save_media(data, _img_ext(data, n.rsplit(".", 1)[-1]), Path(n).stem)
                        if rel:
                            refs.append(f"![{Path(n).name}]({rel})")
                if refs:
                    out.append("## Hình ảnh trong tệp (vị trí gốc không xác định)\n\n" + "\n\n".join(refs))
    except (zipfile.BadZipFile, OSError, KeyError):
        pass
    return "\n\n".join(out)


def _app_props(p: Path) -> dict:
    try:
        with zipfile.ZipFile(p) as z:
            root = _safe_xml(z.read("docProps/app.xml"))
    except (KeyError, zipfile.BadZipFile, OSError, ET.ParseError):
        return {}
    ns = "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}"
    out = {}
    for k in ("Pages", "Slides"):
        v = (root.findtext(ns + k) or "").strip()
        if v.isdigit():
            out[k] = int(v)
    return out


# ════════════════════════════════════════════════════════════════════════════
# DOCX (mammoth của markitdown + ảnh thật + phần ngoài luồng chính)
# ════════════════════════════════════════════════════════════════════════════
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
DOCX_STYLE_MAP = "p[style-name='Title'] => h1:fresh\np[style-name='Subtitle'] => h2:fresh"


def _docx_paragraphs(xml_bytes: bytes) -> list[str]:
    root = _safe_xml(xml_bytes)
    paras: list[list[str]] = []

    def walk(el, cur):
        for ch in el:
            if ch.tag == MC_FALLBACK:
                continue
            if ch.tag == W_NS + "p":
                buf: list[str] = []
                paras.append(buf)
                walk(ch, buf)
            elif ch.tag == W_NS + "t":
                if cur is not None:
                    cur.append(ch.text or "")
            elif ch.tag == W_NS + "tab":
                if cur is not None:
                    cur.append(" ")
            else:
                walk(ch, cur)
    walk(root, None)
    return [_oneline("".join(b)) for b in paras if "".join(b).strip()]


def _norm_for_cmp(s: str) -> str:
    s = re.sub(r"\[\[\d+\]\]\([^)]*\)", "", s)  # chú thích mammoth [[1]](#footnote-1)
    s = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", s)
    s = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", s)
    return re.sub(r"\W+", "", unicodedata.normalize("NFC", s).lower())


def _docx_font_hint(z: zipfile.ZipFile) -> str | None:
    blob = b""
    for n in ("word/document.xml", "word/styles.xml", "word/fontTable.xml"):
        try:
            blob += z.read(n)
        except KeyError:
            pass
    if re.search(rb'w:(?:ascii|hAnsi|cs)="\.Vn', blob):
        return "tcvn3"
    if re.search(rb'w:(?:ascii|hAnsi|cs)="VNI[- ]', blob):
        return "vni"
    return None


def conv_docx(ctx: Ctx, p: Path) -> str:
    import mammoth
    try:
        from markitdown.converter_utils.docx.pre_process import pre_process_docx
    except ImportError:  # API nội bộ markitdown thay đổi → bỏ qua tiền xử lý công thức toán
        pre_process_docx = None
    counter = [0]

    def handle(image):
        counter[0] += 1
        with image.open() as f:
            data = f.read()
        ext = _img_ext(data, (image.content_type or "image/png").split("/")[-1])
        rel = ctx.save_media(data, ext, f"image_{counter[0]:03d}")
        return {"src": rel or "", "alt": (image.alt_text or "").strip()}

    try:
        with open(p, "rb") as f:
            stream = pre_process_docx(f) if pre_process_docx else f
            res = mammoth.convert_to_html(stream, style_map=DOCX_STYLE_MAP,
                                          convert_image=mammoth.images.img_element(handle))
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as e:
        raise Unreadable(MSG["corrupt"].format(detail=f"DOCX: {e}"))
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(res.value, "html.parser")
    _prep_tables(soup)
    md = _html_to_md(str(soup))
    ctx.chain("markitdown/mammoth(docx)")
    extras = []
    try:
        with zipfile.ZipFile(p) as z:
            names = z.namelist()
            ctx.font_hint = _docx_font_hint(z)
            # Phần chữ mammoth bỏ sót (text box, khung, shape) — đối chiếu với document.xml
            body = _docx_paragraphs(z.read("word/document.xml"))
            mdn = _norm_for_cmp(md)
            missing, seen = [], set()
            for t in body[:20000]:
                k = _norm_for_cmp(t)
                if len(k) >= 4 and k not in mdn and k not in seen:
                    seen.add(k)
                    missing.append(t)
            if missing:
                extras.append("## Nội dung ngoài luồng chính (text box, khung, hình vẽ)\n\n" + "\n\n".join(missing))
                ctx.warn(f"{len(missing)} đoạn chữ nằm trong text box/khung không theo luồng chính: đã thêm ở cuối source.md.")
            hf, seen_hf = [], set()
            for n in sorted(names):
                if re.fullmatch(r"word/(header|footer)\d*\.xml", n):
                    for t in _docx_paragraphs(z.read(n)):
                        if t not in seen_hf and not re.fullmatch(r"[\d\s/.-]*", t):
                            seen_hf.add(t)
                            hf.append(t)
            if hf:
                extras.append("## Đầu trang / chân trang (header/footer)\n\n" + "\n\n".join(hf))
            if "word/comments.xml" in names:
                root = _safe_xml(z.read("word/comments.xml"))
                cm = []
                for c in root.iter(W_NS + "comment"):
                    txt = _oneline(" ".join(t.text or "" for t in c.iter(W_NS + "t")))
                    if txt:
                        cm.append(f"- **{c.get(W_NS + 'author') or '?'}**: {txt}")
                if cm:
                    extras.append("## Bình luận (comments)\n\n" + "\n".join(cm))
    except (zipfile.BadZipFile, KeyError, ET.ParseError):
        pass
    more = _zip_extras(ctx, p, "word/charts/", None, "word/diagrams/")
    if more:
        extras.append(more)
    props = _app_props(p)
    if "Pages" in props:
        ctx.man["pages"] = props["Pages"]
    return "\n\n".join([md] + extras)


# ════════════════════════════════════════════════════════════════════════════
# XLSX (openpyxl: bảng giá trị + bảng công thức; LibreOffice tính lại nếu thiếu giá trị đệm)
# ════════════════════════════════════════════════════════════════════════════
def _formula_text(v) -> str | None:
    if isinstance(v, str) and v.startswith("=") and len(v) > 1:
        return v
    t = getattr(v, "text", None)
    if t is not None and "formula" in type(v).__name__.lower():
        return str(t) if str(t).startswith("=") else "=" + str(t)
    return None


def _fmt_cell(v) -> str:
    import datetime as dt
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v.is_integer() and abs(v) < 1e15:
            return str(int(v))
        return f"{v:.15g}"
    if isinstance(v, dt.datetime):
        s = v.isoformat(sep=" ")
        return s[:-9] if s.endswith(" 00:00:00") else s
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    return str(v)


def _get(rows, r, c):
    if r < len(rows) and c < len(rows[r]):
        return rows[r][c]
    return None


def _load_sheets(p: Path, data_only: bool, read_only: bool):
    import openpyxl
    # truyền file object: openpyxl từ chối theo PHẦN ĐUÔI (vd tệp tải về tên .bin) dù nội dung là XLSX thật
    wb = openpyxl.load_workbook(open(p, "rb"), data_only=data_only, read_only=read_only)
    out = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        merged = []
        if not read_only:
            try:
                merged = [str(r) for r in ws.merged_cells.ranges]
            except AttributeError:
                pass
        out.append({"name": ws.title, "state": getattr(ws, "sheet_state", "visible"), "rows": rows, "merged": merged})
    try:
        wb.close()
    except AttributeError:
        pass
    return out


def conv_xlsx(ctx: Ctx, p: Path, allow_recalc: bool = True) -> str:
    from openpyxl.utils import get_column_letter
    big = p.stat().st_size > 15 * 1024 * 1024
    try:
        sf = _load_sheets(p, data_only=False, read_only=big)
        sv = _load_sheets(p, data_only=True, read_only=big)
    except Exception as e:  # noqa: BLE001
        raise Unreadable(MSG["corrupt"].format(detail=f"XLSX: {type(e).__name__}: {e}"))
    ctx.chain("openpyxl(values+formulas)")
    need_recalc = any(_formula_text(fv) and _get(v["rows"], r, c) is None
                      for f, v in zip(sf, sv) for r, row in enumerate(f["rows"]) for c, fv in enumerate(row))
    if need_recalc and allow_recalc:
        try:
            rp = soffice_convert(ctx, p, "xlsx", "xlsx", ".xlsx")
            recalced = {s["name"]: s for s in _load_sheets(rp, data_only=True, read_only=big)}
            for s in sv:
                if s["name"] in recalced:
                    s["rows"] = recalced[s["name"]]["rows"]
            ctx.chain("soffice-recalc")
            ctx.warn("Tệp chưa lưu giá trị công thức (thường do tạo bằng script): đã tính lại bằng LibreOffice.")
        except IngestError:
            ctx.warn("Tệp chưa lưu giá trị công thức và không tính lại được (thiếu LibreOffice): ô công thức hiển thị "
                     "công thức thay cho giá trị. Cài LibreOffice: brew install --cask libreoffice.")
    parts = []
    for i, (f, v) in enumerate(zip(sf, sv), 1):
        name = f["name"]
        hdr = f"## Sheet {i}: {name}" + ("" if f["state"] == "visible" else " (ẩn)")
        rows_f, rows_v = f["rows"], v["rows"]
        maxr = maxc = 0
        for rows in (rows_f, rows_v):
            for r, row in enumerate(rows):
                for c, x in enumerate(row):
                    if x is not None and x != "":
                        maxr, maxc = max(maxr, r + 1), max(maxc, c + 1)
        if maxr == 0:
            parts.append(f"{hdr}\n\n_(Trang tính trống)_")
            continue
        table = [["#"] + [get_column_letter(c + 1) for c in range(maxc)]]
        formulas = []
        for r in range(maxr):
            row = [str(r + 1)]
            for c in range(maxc):
                val, fx = _get(rows_v, r, c), _formula_text(_get(rows_f, r, c))
                row.append(fx if (val is None and fx) else _fmt_cell(val))
                if fx:
                    formulas.append([f"{get_column_letter(c + 1)}{r + 1}", f"`{fx}`", _fmt_cell(val)])
            if any(x for x in row[1:]):
                table.append(row)
        block = [hdr, f"<!-- sheet {i}: {name} | vùng dữ liệu A1:{get_column_letter(maxc)}{maxr} -->", md_table(table)]
        if formulas:
            block += [f"### Công thức — {name}", md_table([["Ô", "Công thức", "Giá trị đã tính"]] + formulas)]
        if f["merged"]:
            block.append("Ô gộp: " + ", ".join(f["merged"][:60]) + (" …" if len(f["merged"]) > 60 else ""))
        parts.append("\n\n".join(block))
    ctx.man["sheets"] = len(sf)
    extras = _zip_extras(ctx, p, "xl/charts/", "xl/media/", None)
    return "\n\n".join(parts + ([extras] if extras else []))


# ════════════════════════════════════════════════════════════════════════════
# PPTX (python-pptx: "## Slide N", chữ, bảng, ảnh, biểu đồ, SmartArt, ghi chú)
# ════════════════════════════════════════════════════════════════════════════
def _sorted_shapes(shapes):
    def key(s):
        try:
            return (int((s.top or 0) / 91440 / 2), s.left or 0)
        except Exception:  # noqa: BLE001
            return (0, 0)
    return sorted(list(shapes), key=key)


def _pptx_shape(ctx: Ctx, slide, shp, sidx: int, out: list[str]):
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
    try:
        st = shp.shape_type
    except Exception:  # noqa: BLE001
        st = None
    if st == MSO_SHAPE_TYPE.GROUP:
        for ch in _sorted_shapes(shp.shapes):
            _pptx_shape(ctx, slide, ch, sidx, out)
        return
    if shp.has_table:
        rows = [[c.text for c in r.cells] for r in shp.table.rows]
        t = md_table(rows)
        if t:
            out.append(t)
        return
    if shp.has_chart:
        try:
            md = chart_xml_to_md(shp.chart_part.blob)
        except Exception:  # noqa: BLE001
            md = ""
        out.append(md or "**Biểu đồ** _(không đọc được dữ liệu)_")
        return
    el = shp._element
    gd = el.find(f".//{A_NS}graphicData")
    if gd is not None and gd.get("uri", "").endswith("/diagram"):
        rel_ids = el.find(f".//{DGM_NS}relIds")
        try:
            texts = smartart_texts(slide.part.related_part(rel_ids.get(f"{R_NS}dm")).blob)
        except Exception:  # noqa: BLE001
            texts = []
        if texts:
            out.append("\n".join(f"- {t}" for t in texts))
        return
    if el.tag.endswith("}pic"):
        try:
            img = shp.image
            descr = next((e.get("descr") for e in el.iter() if e.tag.endswith("}cNvPr")), "") or ""
            rel = ctx.save_media(img.blob, img.ext, f"slide{sidx:03d}_img")
            if rel:
                out.append(f"![{_oneline(descr)}]({rel})")
        except Exception:  # noqa: BLE001
            pass
        return
    if shp.has_text_frame:
        is_body = False
        try:
            is_body = shp.is_placeholder and shp.placeholder_format.type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
        except Exception:  # noqa: BLE001
            pass
        for para in shp.text_frame.paragraphs:
            t = _oneline(para.text.replace("\v", " "))
            if not t:
                continue
            lvl = para.level or 0
            out.append(("  " * lvl + "- " + t) if (is_body or lvl > 0) else t)
        return
    if st == MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT:
        try:
            out.append(f"_[Đối tượng nhúng: {shp.ole_format.prog_id}]_")
        except Exception:  # noqa: BLE001
            out.append("_[Đối tượng nhúng]_")


def conv_pptx(ctx: Ctx, p: Path) -> str:
    from pptx import Presentation
    try:
        prs = Presentation(str(p))
    except Exception as e:  # noqa: BLE001  (vd .ppsx/.potx: content-type khác)
        if not find_soffice():
            raise Unreadable(MSG["corrupt"].format(detail=f"PPTX: {e}"))
        prs = Presentation(str(soffice_convert(ctx, p, "pptx", "pptx", ".pptx")))
        ctx.chain("soffice:normalize-pptx")
    parts = []
    n_notes = 0
    for idx, slide in enumerate(prs.slides, 1):
        hidden = slide._element.get("show") == "0"
        lines = [f"## Slide {idx}" + (" (ẩn)" if hidden else "")]
        title_id = None
        try:
            ts = slide.shapes.title
        except Exception:  # noqa: BLE001
            ts = None
        if ts is not None and ts.has_text_frame and ts.text_frame.text.strip():
            lines.append(f"### {_oneline(ts.text_frame.text)}")
            title_id = ts.shape_id
        body: list[str] = []
        for shp in _sorted_shapes(slide.shapes):
            if title_id is not None and shp.shape_id == title_id:
                continue
            try:
                _pptx_shape(ctx, slide, shp, idx, body)
            except Exception as e:  # noqa: BLE001
                ctx.warn(f"Slide {idx}: bỏ qua một shape lỗi ({type(e).__name__}).")
        lines += body
        if slide.has_notes_slide:
            ns = slide.notes_slide
            tf = ns.notes_text_frame
            note = tf.text.strip() if tf is not None else ""
            if not note:
                note = "\n".join(s.text_frame.text for s in ns.shapes
                                 if s.has_text_frame and not s.is_placeholder and s.text_frame.text.strip())
            if note.strip():
                n_notes += 1
                lines.append("### Ghi chú diễn giả (notes)\n" + note.strip().replace("\v", "\n"))
        parts.append("\n\n".join(lines))
    ctx.man["slides"] = len(prs.slides)
    ctx.chain("python-pptx")
    if n_notes:
        ctx.man["notes"] = n_notes
    return "\n\n".join(parts)


# ════════════════════════════════════════════════════════════════════════════
# PDF (pdfplumber/markitdown từng trang + marker trang + ảnh PyMuPDF + OCR trang scan)
# ════════════════════════════════════════════════════════════════════════════
def _image_coverage(page) -> float:
    import pymupdf
    area = abs(page.rect) or 1.0
    cov = 0.0
    try:
        for info in page.get_image_info():
            r = pymupdf.Rect(info["bbox"]) & page.rect
            if not r.is_empty:
                cov += abs(r)
    except Exception:  # noqa: BLE001
        return 0.0
    return min(1.0, cov / area)


def _vector_heavy(page) -> bool:
    try:
        fn = getattr(page, "get_cdrawings", None) or page.get_drawings
        return len(fn()) > 200
    except Exception:  # noqa: BLE001
        return False


def _plumber_page_md(pl_page) -> str:
    try:
        tables = [t for t in pl_page.find_tables() if t.bbox and len(t.rows) >= 2]
    except Exception:  # noqa: BLE001
        tables = []
    if not tables:
        return pl_page.extract_text() or ""
    boxes = [t.bbox for t in tables]

    def outside(obj):
        try:
            cx, cy = (obj["x0"] + obj["x1"]) / 2, (obj["top"] + obj["bottom"]) / 2
        except (KeyError, TypeError):
            return True
        return not any(b[0] <= cx <= b[2] and b[1] <= cy <= b[3] for b in boxes)

    blocks = []
    for ln in pl_page.filter(outside).extract_text_lines():
        blocks.append((ln["top"], ln["text"]))
    for t in tables:
        md = md_table(t.extract())
        if md:
            blocks.append((t.bbox[1], "\n" + md + "\n"))
    blocks.sort(key=lambda b: b[0])
    return "\n".join(b[1] for b in blocks)


def _pdf_alternatives(p: Path, page, pno: int) -> list[tuple[str, str]]:
    cands = []
    exe = shutil.which("pdftotext")
    if exe:
        try:
            r = subprocess.run([exe, "-layout", "-enc", "UTF-8", "-f", str(pno), "-l", str(pno), str(p), "-"],
                               capture_output=True, timeout=60)
            if r.returncode == 0:
                txt = "\n".join(ln.rstrip() for ln in r.stdout.decode("utf-8", "replace").splitlines())
                cands.append(("pdftotext-layout", txt.strip("\f\n")))
        except (subprocess.TimeoutExpired, OSError):
            pass
    try:
        cands.append(("pymupdf-sort", page.get_text(sort=True)))
    except Exception:  # noqa: BLE001
        pass
    try:
        from pdfminer.high_level import extract_text
        cands.append(("pdfminer", extract_text(str(p), page_numbers=[pno - 1])))
    except Exception:  # noqa: BLE001
        pass
    return cands


def _pdf_page_images(ctx: Ctx, doc, page, pno: int, seen: set) -> list[str]:
    import pymupdf
    refs = []
    for info in page.get_images(full=True):
        xref, w, h = info[0], info[2], info[3]
        if xref in seen:
            continue
        seen.add(xref)
        if w < 48 or h < 48:
            continue
        try:
            d = doc.extract_image(xref)
            data, ext = d.get("image"), (d.get("ext") or "png").lower()
            if ext not in ("png", "jpg", "jpeg"):
                pix = pymupdf.Pixmap(doc, xref)
                if pix.colorspace is not None and pix.colorspace.n >= 4:
                    pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                data, ext = pix.tobytes("png"), "png"
        except Exception:  # noqa: BLE001
            continue
        rel = ctx.save_media(data, ext, f"page{pno:03d}_img")
        if rel:
            refs.append(f"![Hình trang {pno}]({rel})")
    return refs


def conv_pdf(ctx: Ctx, p: Path) -> str:
    import pymupdf
    try:
        doc = pymupdf.open(str(p))
    except Exception as e:  # noqa: BLE001
        raise Unreadable(MSG["corrupt"].format(detail=f"PDF: {e}"))
    if doc.needs_pass:
        raise Unreadable(MSG["pdf_password"])
    n = doc.page_count
    ctx.man["pages"] = n
    if n == 0:
        raise Unreadable(MSG["no_content"].format(detail="PDF 0 trang"))
    try:
        from markitdown.converters._pdf_converter import _extract_form_content_from_words as md_form
    except Exception:  # noqa: BLE001  (API nội bộ markitdown có thể đổi)
        md_form = None
    pl = None
    try:
        import pdfplumber
        pl = pdfplumber.open(str(p))
        ctx.chain("pdfplumber/markitdown(per-page)")
    except Exception:  # noqa: BLE001
        ctx.chain("pymupdf(text)")
        ctx.warn("pdfplumber không mở được PDF: dùng PyMuPDF để trích chữ (có thể mất cấu trúc bảng).")
    parts: list = []
    ocr_items = []
    seen_xref: set = set()
    damaged_pages, fixed_pages = [], []
    ocr_dir = ctx.out / "ocr_pages"
    for i in range(n):
        pno = i + 1
        page = doc[i]
        raw = page.get_text()
        letters = sum(c.isalnum() for c in raw)
        bad = sum(1 for c in raw if c == "\ufffd" or 0xE000 <= ord(c) <= 0xF8FF)
        cov = _image_coverage(page)
        reason = None
        if letters < 20:
            if cov >= 0.05 or _vector_heavy(page):
                reason = "no_text_layer"
        elif bad / max(1, letters + bad) > 0.3:
            reason = "garbage_text_layer"
        elif cov >= 0.6 and letters < 80:
            reason = "mostly_image"
        if reason:
            ocr_dir.mkdir(parents=True, exist_ok=True)
            w_in, h_in = page.rect.width / 72, page.rect.height / 72
            dpi = max(72, min(200, int(4000 / max(w_in, h_in, 0.1))))
            png = ocr_dir / f"page_{pno:03d}.png"
            page.get_pixmap(dpi=dpi).save(str(png))
            ocr_items.append((pno, png, reason, raw if letters >= 20 else ""))
            parts.append(("ocr", pno))
            continue
        if letters < 20:
            parts.append(f"<!-- page {pno} -->\n_(trang trống)_")
            continue
        content = None
        if pl is not None:
            try:
                plp = pl.pages[i]
                if md_form is not None:
                    try:
                        content = md_form(plp)
                    except Exception:  # noqa: BLE001
                        content = None
                if content is None:
                    content = _plumber_page_md(plp)
                plp.close()
            except Exception:  # noqa: BLE001
                content = None
        if not content or not content.strip():
            content = page.get_text(sort=True)
        name, best, score = choose_least_damaged(("pdfplumber", content), [])
        if score >= 0.06:
            damaged_pages.append(pno)
            name, best, score2 = choose_least_damaged(("pdfplumber", content), _pdf_alternatives(p, page, pno))
            if name != "pdfplumber":
                fixed_pages.append((pno, name))
                content = best
        imgs = _pdf_page_images(ctx, doc, page, pno, seen_xref)
        parts.append(f"<!-- page {pno} -->\n" + _space_tables(content.strip()) + ("\n\n" + "\n\n".join(imgs) if imgs else ""))
    if pl is not None:
        pl.close()
    if fixed_pages:
        ctx.warn("Font dự phòng làm rơi dấu tiếng Việt ở trang " + ", ".join(str(x) for x, _ in fixed_pages[:30]) +
                 f": đã dùng bộ trích khác ({', '.join(sorted({nm for _, nm in fixed_pages}))}) cho các trang này.")
        ctx.chain("diacritic-fallback")
    unfixed = [x for x in damaged_pages if x not in {f for f, _ in fixed_pages}]
    if unfixed:
        ctx.warn("Nghi rơi dấu tiếng Việt (font nhúng lỗi) ở trang " + ", ".join(map(str, unfixed[:30])) +
                 ": đối chiếu với ảnh trang (render bằng PyMuPDF) trước khi trích dẫn.")
    blocks = _ocr_blocks(ctx, ocr_items) if ocr_items else {}
    out = []
    for item in parts:
        if isinstance(item, tuple):
            out.append(f"<!-- page {item[1]} -->\n" + blocks[item[1]])
        else:
            out.append(item)
    doc.close()
    return "\n\n".join(out)


# ════════════════════════════════════════════════════════════════════════════
# EPUB (chương theo spine, ảnh thật, chặn DRM)
# ════════════════════════════════════════════════════════════════════════════
XMLENC = "{http://www.w3.org/2001/04/xmlenc#}"
FONT_OBFUSCATION = {"http://www.idpf.org/2008/embedding", "http://ns.adobe.com/pdf/enc#RC"}


def _epub_has_drm(z: zipfile.ZipFile, names: set) -> bool:
    if "META-INF/rights.xml" in names:
        return True
    if "META-INF/encryption.xml" not in names:
        return False
    try:
        root = _safe_xml(z.read("META-INF/encryption.xml"))
    except ET.ParseError:
        return True
    for ed in root.iter(f"{XMLENC}EncryptedData"):
        m = ed.find(f"{XMLENC}EncryptionMethod")
        alg = m.get("Algorithm", "") if m is not None else ""
        ref = ed.find(f".//{XMLENC}CipherReference")
        uri = (ref.get("URI", "") if ref is not None else "").lower()
        if alg in FONT_OBFUSCATION or uri.endswith((".ttf", ".otf", ".woff", ".woff2")):
            continue
        return True
    return False


def conv_epub(ctx: Ctx, p: Path) -> str:
    from bs4 import BeautifulSoup
    try:
        z = zipfile.ZipFile(p)
    except zipfile.BadZipFile as e:
        raise Unreadable(MSG["corrupt"].format(detail=f"EPUB: {e}"))
    with z:
        names = set(z.namelist())
        if _epub_has_drm(z, names):
            raise Unreadable(MSG["epub_drm"])
        try:
            croot = _safe_xml(z.read("META-INF/container.xml"))
            rf = next(e for e in croot.iter() if e.tag.endswith("rootfile"))
            opf_path = rf.get("full-path")
            opf = _safe_xml(z.read(opf_path))
        except (KeyError, StopIteration, ET.ParseError) as e:
            raise Unreadable(MSG["corrupt"].format(detail=f"EPUB thiếu/hỏng container hoặc OPF: {e}"))
        opf_dir = posixpath.dirname(opf_path)
        manifest = {}
        for it in opf.iter():
            if it.tag.endswith("}item") or it.tag == "item":
                href = posixpath.normpath(posixpath.join(opf_dir, unquote(it.get("href", ""))))
                manifest[it.get("id")] = (href, it.get("media-type", ""))
        spine = [ir.get("idref") for ir in opf.iter() if ir.tag.endswith("itemref") or ir.tag == "itemref"]
        meta = {}
        for tag in ("title", "creator", "language"):
            el = next((e for e in opf.iter() if e.tag.endswith("}" + tag)), None)
            if el is not None and (el.text or "").strip():
                meta[tag] = el.text.strip()
        head = []
        if meta.get("title"):
            head.append(f"**Tiêu đề:** {meta['title']}")
        if meta.get("creator"):
            head.append(f"**Tác giả:** {meta['creator']}")
        parts = ["  \n".join(head)] if head else []
        chapters, unreadable = 0, 0
        for idref in spine:
            href, mt = manifest.get(idref, (None, ""))
            if not href or ("html" not in mt and not href.lower().endswith((".xhtml", ".html", ".htm"))):
                continue
            try:
                raw = z.read(href)
            except KeyError:
                ctx.warn(f"EPUB thiếu tệp chương {href}.")
                continue
            if b"\x00" in raw[:2048]:
                unreadable += 1
                continue
            soup = BeautifulSoup(_decode_html(raw), "html.parser")
            base = posixpath.dirname(href)

            def resolver(src, base=base):
                if src.startswith("data:"):
                    return _data_uri(src)
                target = posixpath.normpath(posixpath.join(base, unquote(src.split("#")[0])))
                if target in names:
                    data = z.read(target)
                    return data, _img_ext(data, target.rsplit(".", 1)[-1])
                return None, None
            _rewrite_html_images(ctx, soup, resolver, f"epub_{Path(href).stem}")
            _prep_tables(soup)
            md = _html_to_md(str(soup)).strip()
            if md:
                chapters += 1
                parts.append(f"<!-- chapter {chapters}: {href} -->\n{md}")
        if chapters == 0 and unreadable:
            raise Unreadable(MSG["epub_drm"])
    ctx.man["chapters"] = chapters
    ctx.chain("epub-spine+markitdown(html)")
    return "\n\n".join(parts)


# ════════════════════════════════════════════════════════════════════════════
# HTML / Markdown / văn bản / CSV / Outlook MSG / ảnh
# ════════════════════════════════════════════════════════════════════════════
def conv_html(ctx: Ctx, p: Path) -> str:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(_decode_html(p.read_bytes()), "html.parser")
    _rewrite_html_images(ctx, soup, lambda s: _local_image(p.parent, s), "html_img")
    _prep_tables(soup)
    title = soup.title.get_text(strip=True) if soup.title else ""
    md = _html_to_md(str(soup))
    ctx.chain("markitdown(html)")
    if title and title not in md[:500]:
        md = f"# {title}\n\n{md}"
    return md


MD_IMG_RE = re.compile(r'!\[([^\]]*)\]\(\s*<?([^)\s>]+)>?(?:\s+"[^"]*")?\s*\)')


def conv_markdown(ctx: Ctx, p: Path) -> str:
    text, enc, legacy = decode_bytes(p.read_bytes())
    ctx.man["text_encoding"] = enc
    if legacy:
        ctx.man["legacy_encoding"] = legacy

    def repl(m):
        data, ext = _local_image(p.parent, m.group(2))
        rel = ctx.save_media(data, ext, Path(m.group(2)).stem or "image") if data else None
        return f"![{m.group(1)}]({rel})" if rel else m.group(0)
    ctx.chain(f"text:{enc}")
    return MD_IMG_RE.sub(repl, text)


def conv_text(ctx: Ctx, p: Path) -> str:
    text, enc, legacy = decode_bytes(p.read_bytes())
    ctx.man["text_encoding"] = enc
    if legacy:
        ctx.man["legacy_encoding"] = legacy
    if legacy == "cp1258":
        ctx.warn("Tệp không phải UTF-8: đã giải mã theo Windows-1258 (CP1258). Kiểm tra lại dấu tiếng Việt.")
    ctx.chain(f"text:{enc}")
    return text


def conv_csv(ctx: Ctx, p: Path) -> str:
    text, enc, legacy = decode_bytes(p.read_bytes())
    ctx.man["text_encoding"] = enc
    if legacy:
        ctx.man["legacy_encoding"] = legacy
    delim = "\t" if p.suffix.lower() == ".tsv" else None
    if delim is None:
        try:
            delim = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delim = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    ctx.chain(f"csv(stdlib,{enc},delim={delim!r})")
    return f"<!-- csv: {len(rows)} dòng, phân cách {delim!r} -->\n\n" + md_table(rows)


def conv_msg(ctx: Ctx, p: Path) -> str:
    from markitdown import MarkItDown
    src = p
    if p.suffix.lower() != ".msg":
        src = ctx.tmp / "input.msg"
        shutil.copyfile(p, src)
    ctx.chain("markitdown(outlook-msg)")
    return MarkItDown().convert(str(src)).text_content or ""


def conv_image(ctx: Ctx, p: Path, fmt: str) -> str:
    from PIL import Image, ImageOps
    src = p
    if fmt == "heic":
        try:
            import pillow_heif  # type: ignore
            pillow_heif.register_heif_opener()
        except ImportError:
            if sys.platform == "darwin" and shutil.which("sips"):
                conv = ctx.tmp / "heic.png"
                subprocess.run(["sips", "-s", "format", "png", str(p), "--out", str(conv)], capture_output=True, timeout=120)
                if not conv.exists():
                    raise Unreadable(MSG["corrupt"].format(detail="HEIC không chuyển được bằng sips"))
                src = conv
                ctx.chain("sips:heic->png")
            else:
                raise MissingDependency(MSG["heic"])
    try:
        im = Image.open(src)
        im.load()
    except Exception as e:  # noqa: BLE001
        raise Unreadable(MSG["corrupt"].format(detail=f"ảnh: {e}"))
    ocr_dir = ctx.out / "ocr_pages"
    ocr_dir.mkdir(parents=True, exist_ok=True)
    n = min(getattr(im, "n_frames", 1) if fmt == "tiff" else 1, 500)
    items = []
    for k in range(n):
        im.seek(k)
        fr = ImageOps.exif_transpose(im) if k == 0 else im.copy()
        if fr.mode not in ("RGB", "L", "RGBA"):
            fr = fr.convert("RGB")
        if max(fr.size) > 6000:
            fr.thumbnail((6000, 6000))
        png = ocr_dir / f"page_{k + 1:03d}.png"
        fr.save(png)
        items.append((k + 1, png, "image_input", ""))
    ctx.man["pages"] = n
    ctx.chain(f"image:{fmt}->png")
    blocks = _ocr_blocks(ctx, items)
    return "\n\n".join(f"<!-- page {k} -->\n{blocks[k]}" for k, _, _, _ in items)


def _pandoc_fallback(ctx: Ctx, p: Path, fmt: str) -> str | None:
    exe = shutil.which("pandoc")
    if not exe:
        return None
    media_tmp = ctx.tmp / "pandoc_media"
    out_md = ctx.tmp / "pandoc.md"
    try:
        r = subprocess.run([exe, str(p), "-f", fmt, "-t", "gfm", "--wrap=none", f"--extract-media={media_tmp}",
                            "-o", str(out_md)], capture_output=True, text=True, timeout=300)
    except (subprocess.TimeoutExpired, OSError):
        return None
    if r.returncode != 0 or not out_md.exists():
        return None
    md = out_md.read_text("utf-8", "replace")
    if media_tmp.exists():
        for f in sorted(media_tmp.rglob("*")):
            if f.is_file():
                rel = ctx.save_media(f.read_bytes(), f.suffix, f.stem)
                if rel:
                    md = md.replace(str(f), rel)
    ctx.chain(f"pandoc({fmt}->gfm)")
    return md


# ════════════════════════════════════════════════════════════════════════════
# Điều phối
# ════════════════════════════════════════════════════════════════════════════
def _convert_legacy_vector_images(ctx: Ctx, md: str) -> str:
    """EMF/WMF (ảnh vector Windows) → PNG bằng LibreOffice để Agent xem được."""
    vec = [r for r in ctx.media_paths() if r.lower().endswith((".emf", ".wmf"))]
    if not vec:
        return md
    if not find_soffice():
        ctx.warn(f"{len(vec)} ảnh EMF/WMF (vector Windows) chưa chuyển sang PNG (thiếu LibreOffice).")
        return md
    for rel in vec[:50]:
        src = ctx.out / rel
        try:
            png = soffice_convert(ctx, src, "png", "emf", src.suffix)
        except IngestError:
            continue
        dst = src.with_suffix(".png")
        shutil.copyfile(png, dst)
        md = md.replace(f"({rel})", f"(media/{dst.name})")
    ctx.chain("soffice:emf/wmf->png")
    return md


def _convert(ctx: Ctx, fmt: str, info: dict) -> str:
    src = ctx.src
    if fmt == "empty":
        raise Unreadable(MSG["empty"])
    if info.get("encrypted") or fmt == "ooxml_encrypted":
        raise Unreadable(MSG["office_password"])
    if fmt == "pdf":
        return conv_pdf(ctx, src)
    if fmt == "docx":
        return conv_docx(ctx, src)
    if fmt in ("doc", "odt", "rtf"):
        try:
            p = soffice_convert(ctx, src, "docx", fmt, "." + fmt)
        except MissingDependency:
            if fmt in ("odt", "rtf"):
                md = _pandoc_fallback(ctx, src, fmt)
                if md is not None:
                    ctx.warn("Không có LibreOffice: đã đọc bằng pandoc (bảng phức tạp có thể ở dạng HTML). "
                             "Cài LibreOffice để có kết quả tốt hơn: brew install --cask libreoffice.")
                    return md
            raise
        ctx.chain(f"soffice:{fmt}->docx")
        return conv_docx(ctx, p)
    if fmt == "xlsx":
        return conv_xlsx(ctx, src)
    if fmt in ("xls", "ods", "xlsb"):
        try:
            p = soffice_convert(ctx, src, "xlsx", fmt, "." + fmt)
        except MissingDependency:
            if fmt == "xls":
                from markitdown import MarkItDown
                tmp = ctx.tmp / "input.xls"
                shutil.copyfile(src, tmp)
                ctx.chain("markitdown(xls:xlrd)")
                ctx.warn("Không có LibreOffice: đọc XLS bằng xlrd (chỉ có giá trị, KHÔNG có bảng công thức). "
                         "Cài LibreOffice: brew install --cask libreoffice.")
                return MarkItDown().convert(str(tmp)).text_content or ""
            raise
        ctx.chain(f"soffice:{fmt}->xlsx")
        return conv_xlsx(ctx, p, allow_recalc=False)
    if fmt == "pptx":
        return conv_pptx(ctx, src)
    if fmt in ("ppt", "odp"):
        p = soffice_convert(ctx, src, "pptx", fmt, "." + fmt)
        ctx.chain(f"soffice:{fmt}->pptx")
        return conv_pptx(ctx, p)
    if fmt == "epub":
        return conv_epub(ctx, src)
    if fmt == "html":
        return conv_html(ctx, src)
    if fmt == "csv":
        return conv_csv(ctx, src)
    if fmt == "markdown":
        return conv_markdown(ctx, src)
    if fmt == "text":
        return conv_text(ctx, src)
    if fmt == "msg":
        return conv_msg(ctx, src)
    if fmt in IMAGE_FORMATS:
        return conv_image(ctx, src, fmt)
    if fmt == "corrupt":
        raise Unreadable(MSG["corrupt"].format(detail=info.get("detail", "")))
    raise Unreadable(MSG["unsupported"].format(detail=info.get("detail") or fmt))


def _postprocess(ctx: Ctx, md: str) -> str:
    md = md.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    md = unicodedata.normalize("NFC", md)
    kind = detect_legacy(md, hint=ctx.font_hint or ctx.man.get("legacy_encoding"))
    if kind == "tcvn3":
        md, changed = tcvn3_to_unicode(md)
        if changed:
            ctx.man["legacy_encoding"] = "tcvn3"
            ctx.chain("tcvn3->unicode")
            ctx.warn(f"Văn bản dùng mã TCVN3/ABC (font .VnTime…): đã tự chuyển {changed} dòng sang Unicode. "
                     "Kiểm tra lại chữ HOA có dấu và tên riêng trước khi trích dẫn.")
    elif kind == "vni":
        ctx.man["legacy_encoding"] = "vni"
        ctx.warn("Văn bản dùng mã VNI (font VNI-Times…): dấu tiếng Việt hiển thị sai và CHƯA được chuyển tự động. "
                 "Hãy hỏi người dùng bản Unicode, hoặc nhờ họ chuyển mã bằng Unikey (Công cụ chuyển mã: VNI Windows → "
                 "Unicode) rồi gửi lại.")
    md = unicodedata.normalize("NFC", md)
    md = re.sub(r"[ \t]+\n", "\n", md)
    md = re.sub(r"\n{4,}", "\n\n\n", md)
    return md.strip() + "\n"


def _new_manifest(src: Path, out: Path) -> dict:
    try:
        size = src.stat().st_size
    except OSError:
        size = None
    return {
        "tool": "doc_ingest", "version": __version__,
        "input": str(src.resolve() if src.exists() else src), "input_name": src.name, "input_bytes": size,
        "out_dir": str(out.resolve()),
        "detected_format": None, "converter_chain": [],
        "pages": None, "sheets": None, "slides": None, "chapters": None,
        "chars": 0, "tables": 0, "images": 0,
        "warnings": [], "pages_needing_ocr": [], "ocr_pages": [],
        "legacy_encoding": None, "text_encoding": None,
        "status": None, "exit_code": None, "message": None,
        "outputs": {"source_md": "source.md", "media_dir": None, "ocr_dir": None, "manifest": "manifest.json"},
        "truncated": False,
    }


def _prepare_out(out: Path):
    out.mkdir(parents=True, exist_ok=True)
    prev = out / "manifest.json"
    ours = False
    if prev.exists():
        try:
            ours = json.loads(prev.read_text("utf-8")).get("tool") == "doc_ingest"
        except (ValueError, OSError, AttributeError):
            ours = False
    for name in ("source.md", "source_full.md"):
        if ours:
            (out / name).unlink(missing_ok=True)
    if ours:
        prev.unlink(missing_ok=True)
        for d in ("media", "ocr_pages"):
            shutil.rmtree(out / d, ignore_errors=True)


def _write_outputs(out: Path, man: dict, md: str, max_chars: int | None):
    man["chars"] = len(md)
    if max_chars and len(md) > max_chars:
        (out / "source_full.md").write_text(md, encoding="utf-8")
        cut = md.rfind("\n", 0, max_chars)
        if cut < max_chars * 0.8:
            cut = max_chars
        md = (md[:cut].rstrip() + f"\n\n<!-- TRUNCATED: source.md cắt ở {cut:,} / {man['chars']:,} ký tự theo --max-chars. "
              "Toàn văn: source_full.md -->\n")
        man["truncated"] = True
        man["outputs"]["source_full_md"] = "source_full.md"
        man["warnings"].append(f"Đã cắt source.md còn {cut:,}/{man['chars']:,} ký tự (--max-chars {max_chars}). "
                               "Toàn văn ở source_full.md — đọc tiếp nếu cần.")
    (out / "source.md").write_text(md, encoding="utf-8")
    (out / "manifest.json").write_text(json.dumps(man, ensure_ascii=False, indent=2), encoding="utf-8")


def _ingest_via_subprocess(vp: Path, src: Path, out: Path, ocr: bool, max_chars: int | None) -> dict:
    cmd = [str(vp), str(Path(__file__).resolve()), str(src), "--out", str(out), "--json"]
    if not ocr:
        cmd.append("--no-ocr")
    if max_chars:
        cmd += ["--max-chars", str(max_chars)]
    env = dict(os.environ, **{REEXEC_ENV: "1", "PYTHONIOENCODING": "utf-8"})
    r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", env=env)
    try:
        return json.loads(r.stdout)
    except ValueError:
        man = _new_manifest(src, out)
        man.update(status="missing_dependency", exit_code=EXIT_MISSING_DEP,
                   message=_missing_python_message(_missing_modules()), error_detail=(r.stderr or "")[-1500:])
        return man


def ingest(path, out_dir, ocr: bool = True, max_chars: int | None = None) -> dict:
    """Chuyển một tệp → <out_dir>/source.md + manifest.json (+ media/, ocr_pages/). Trả về manifest (dict).
    Không ném lỗi với tệp hỏng/khoá: xem manifest['status'] / ['exit_code'] / ['message']."""
    t0 = time.time()
    src, out = Path(path).expanduser(), Path(out_dir).expanduser()
    missing = _missing_modules()
    if missing:
        vp = _venv_python()
        if vp and not _running_in_repo_venv() and not os.environ.get(REEXEC_ENV):
            return _ingest_via_subprocess(vp, src, out, ocr, max_chars)
    _prepare_out(out)
    man = _new_manifest(src, out)
    ctx = Ctx(src, out, man, ocr)
    md = ""
    try:
        if missing:
            raise MissingDependency(_missing_python_message(missing))
        if not src.exists():
            raise Unreadable(MSG["not_found"].format(path=src))
        if src.is_dir():
            raise Unreadable(MSG["is_dir"].format(path=src))
        fmt, info = detect_format(src)
        man["detected_format"] = fmt
        md = _convert(ctx, fmt, info)
        md = _convert_legacy_vector_images(ctx, md)
        md = _postprocess(ctx, md)
        if len(re.sub(r"<!--.*?-->|\s", "", md, flags=re.S)) == 0 and not man["pages_needing_ocr"]:
            raise Unreadable(MSG["no_content"].format(detail=fmt))
        code = EXIT_PARTIAL if man["pages_needing_ocr"] else EXIT_OK
        if code == EXIT_PARTIAL:
            pg = man["pages_needing_ocr"]
            total = man.get("pages") or len(pg)
            man["message"] = (f"{len(pg)}/{total} trang cần OCR bằng thị giác: đọc ảnh trong ocr_pages/ "
                              "(tài liệu scan dài → dùng skill boc-tach-pdf). Bản nháp OCR trong source.md CHƯA kiểm chứng.")
            ctx.warn(man["message"])
    except IngestError as e:
        code = e.code
        man["message"] = e.message
        if e.detail:
            man["error_detail"] = e.detail
        label = "THIẾU PHỤ THUỘC" if code == EXIT_MISSING_DEP else "KHÔNG ĐỌC ĐƯỢC"
        md = f"<!-- doc_ingest: {label} -->\n> [!CAUTION] {label}: {e.message}\n"
    except Exception as e:  # noqa: BLE001
        code = EXIT_UNREADABLE
        man["message"] = MSG["corrupt"].format(detail=f"{type(e).__name__}: {e}")
        man["error_detail"] = traceback.format_exc(limit=6)
        md = f"<!-- doc_ingest: KHÔNG ĐỌC ĐƯỢC -->\n> [!CAUTION] {man['message']}\n"
    finally:
        ctx.cleanup()
    man["status"], man["exit_code"] = STATUS_BY_CODE[code], code
    man["images"] = ctx.image_count
    man["tables"] = _count_tables(md) if code in (EXIT_OK, EXIT_PARTIAL) else 0
    if (out / "media").is_dir():
        man["outputs"]["media_dir"] = "media"
    if (out / "ocr_pages").is_dir():
        man["outputs"]["ocr_dir"] = "ocr_pages"
    if len(md) > BIG_TEXT_CHARS and code in (EXIT_OK, EXIT_PARTIAL):
        man["warnings"].append(f"Văn bản dài {len(md):,} ký tự: đọc source.md theo từng đoạn (theo marker "
                               "<!-- page N --> / ## Slide / ## Sheet / <!-- chapter -->), không nạp một lần.")
    man["elapsed_sec"] = round(time.time() - t0, 2)
    _write_outputs(out, man, md, max_chars)
    return man


# ════════════════════════════════════════════════════════════════════════════
# CLI
# ════════════════════════════════════════════════════════════════════════════
def _summary_line(man: dict) -> str:
    tag = {"ok": "OK", "partial": "PARTIAL", "unreadable": "KHÔNG ĐỌC ĐƯỢC", "missing_dependency": "THIẾU PHỤ THUỘC"}
    name = man.get("input_name") or man.get("input")
    head = f"[{tag.get(man.get('status'), man.get('status'))}] {name}"
    if man.get("status") in ("ok", "partial"):
        counts = [f"{man.get('detected_format')}", f"{man.get('chars', 0):,} ký tự", f"{man.get('tables', 0)} bảng",
                  f"{man.get('images', 0)} ảnh"]
        for k, lab in (("pages", "trang"), ("sheets", "sheet"), ("slides", "slide"), ("chapters", "chương")):
            if man.get(k):
                counts.append(f"{man[k]} {lab}")
        line = f"{head} → {Path(man['out_dir']) / 'source.md'} | " + ", ".join(counts)
    else:
        line = f"{head}: {man.get('message')}"
    for w in man.get("warnings", []):
        if w != man.get("message"):
            line += f"\n  ! {w}"
    if man.get("status") == "partial":
        line += f"\n  ! {man.get('message')}"
    return line


def main(argv: list[str] | None = None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(
        prog="doc_ingest.py",
        description="Chuyển tài liệu (DOC/DOCX/XLS/XLSX/PPT/PPTX/PDF/EPUB/ODT/ODS/ODP/RTF/HTML/CSV/TXT/ảnh) → "
                    "source.md + manifest.json. Mã thoát: 0 ok, 3 cần OCR thị giác, 2 không đọc được, 4 thiếu phụ thuộc.")
    ap.add_argument("inputs", nargs="+", help="tệp đầu vào (một hoặc nhiều)")
    ap.add_argument("--out", required=True, help="thư mục đầu ra (nhiều tệp → mỗi tệp một thư mục con)")
    ap.add_argument("--json", action="store_true", help="in manifest JSON ra stdout")
    ap.add_argument("--no-ocr", action="store_true", help="không chạy OCR nháp (vẫn render PNG trang scan)")
    ap.add_argument("--max-chars", type=int, default=None,
                    help="cắt source.md ở N ký tự (toàn văn giữ ở source_full.md; chỉ cảnh báo, không đổi mã thoát)")
    args = ap.parse_args(argv)

    missing = _missing_modules()
    if missing and not _running_in_repo_venv() and not os.environ.get(REEXEC_ENV):
        vp = _venv_python()
        if vp:
            env = dict(os.environ, **{REEXEC_ENV: "1"})
            return subprocess.call([str(vp), str(Path(__file__).resolve()), *argv], env=env)

    out = Path(args.out).expanduser()
    ocr = not args.no_ocr
    results = []
    if len(args.inputs) == 1:
        results.append(ingest(args.inputs[0], out, ocr=ocr, max_chars=args.max_chars))
    else:
        out.mkdir(parents=True, exist_ok=True)
        used: set[str] = set()
        for inp in args.inputs:
            slug, k = slugify(Path(inp).name), 1
            base = slug
            while slug in used:
                k += 1
                slug = f"{base}-{k}"
            used.add(slug)
            results.append(ingest(inp, out / slug, ocr=ocr, max_chars=args.max_chars))
    code = max((m.get("exit_code", EXIT_UNREADABLE) for m in results), key=lambda c: SEVERITY.get(c, 2))
    if len(results) > 1:
        index = {"tool": "doc_ingest", "version": __version__, "exit_code": code, "status": STATUS_BY_CODE[code],
                 "results": [{"input": m.get("input"), "folder": Path(m["out_dir"]).name, "status": m.get("status"),
                              "exit_code": m.get("exit_code"), "detected_format": m.get("detected_format"),
                              "chars": m.get("chars"), "message": m.get("message")} for m in results]}
        (out / "index.json").write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")
    if args.json:
        payload = results[0] if len(results) == 1 else {"exit_code": code, "status": STATUS_BY_CODE[code], "results": results}
        print(json.dumps(payload, ensure_ascii=False, indent=2))
    else:
        for m in results:
            print(_summary_line(m))
        if len(results) > 1:
            print(f"Tổng hợp: {out / 'index.json'} (mã thoát {code})")
    return code


if __name__ == "__main__":
    sys.exit(main())
