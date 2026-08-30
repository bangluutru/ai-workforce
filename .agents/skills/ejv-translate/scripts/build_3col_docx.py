#!/usr/bin/env python3
"""Build professional 3-column comparative DOCX for EJV Trilingual Translation."""

import argparse
import json
import os
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="999999"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="D0D0D0"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def format_text(val):
    if val is None:
        return ""
    if isinstance(val, list):
        return "\n".join([f"• {str(item)}" for item in val])
    return str(val).strip()


def build_3col_docx(blocks: list, output_path: Path, doc_title: str = "TÀI LIỆU DỊCH THUẬT TAM NGỮ (VIỆT - ANH - NHẬT)"):
    doc = Document()

    # Set Landscape Orientation for 3 columns (A4 Landscape)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    # Normal Style
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Arial'
    normal_font.size = Pt(8.5)
    normal_font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run("BẢN ĐỐI CHIẾU DỊCH THUẬT TAM NGỮ (VIỆT - ANH - NHẬT)\nTRILINGUAL COMPARATIVE TRANSLATION (VIETNAMESE - ENGLISH - JAPANESE)\n三言語対照翻訳（ベトナム語・英語・日本語）")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run(doc_title)
    sub_run.bold = True
    sub_run.font.size = Pt(10.5)
    sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()  # Spacer

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    col_widths = [Cm(8.8), Cm(8.8), Cm(8.8)]

    # Header row
    hdr_cells = table.rows[0].cells
    headers = [
        ("🇻🇳 TIẾNG VIỆT (GỐC)", "1A365D"),
        ("🇬🇧 ENGLISH (TRANSLATION)", "1E40AF"),
        ("🇯🇵 日本語 (翻訳)", "1E293B")
    ]

    for i, (text, color_hex) in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], color_hex)
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Repeat header row across pages
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))

    for idx, block in enumerate(blocks):
        btype = block.get("type", "p")
        vn_text = format_text(block.get("vn", ""))
        en_text = format_text(block.get("en", ""))
        ja_text = format_text(block.get("ja", ""))

        if not vn_text and not en_text and not ja_text:
            continue

        row = table.add_row()
        cells = row.cells

        # Styling depending on block type
        is_heading = btype in ["h1", "h2", "h3"]
        bg_color = "F1F5F9" if is_heading else ("F8FAFC" if idx % 2 == 0 else "FFFFFF")

        for i, text in enumerate([vn_text, en_text, ja_text]):
            cells[i].width = col_widths[i]
            set_cell_background(cells[i], bg_color)
            set_cell_margins(cells[i], top=80 if is_heading else 60,
                             bottom=80 if is_heading else 60,
                             left=120, right=120)

            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)

            if is_heading:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(text)
                r.bold = True
                r.font.size = Pt(9.5 if btype == "h1" else 9.0)
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(text)
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    doc.save(str(output_path))
    print(f"✅ Generated 3-Column Comparative DOCX: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Export EJV JSON to 3-column comparative DOCX.")
    parser.add_argument("--input", required=True, type=Path, help="Input EJV JSON file")
    parser.add_argument("--output", required=True, type=Path, help="Output DOCX file path")
    parser.add_argument("--title", default="TÀI LIỆU DỊCH THUẬT TAM NGỮ", help="Document title")
    args = parser.parse_args()

    with open(args.input, "r", encoding="utf-8") as f:
        blocks = json.load(f)

    build_3col_docx(blocks, args.output, doc_title=args.title)


if __name__ == "__main__":
    main()
