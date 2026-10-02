#!/usr/bin/env python3
"""
ingest.py — Chuẩn bị tài liệu nguồn cho Đọc Sâu: lưu văn bản nguồn thành FILE CỤC BỘ (để kiểm chứng trích dẫn)
và chia thành các đoạn (chunk) đọc tuần tự, kèm bảng độ phủ.

Đầu vào: .pdf / .docx / .md / .txt / .html. Với URL: Agent tải trang bằng công cụ đọc web, lưu nội dung chữ vào
một file .md/.txt trong <process_dir>/sources/ rồi mới chạy ingest (evidence_verifier không nhận URL).

Usage:
  python3 ingest.py <file_nguon> --process-dir <process_dir> [--chunk-chars 6000] [--as main|ref]
     --as main (mặc định): tài liệu chính -> <process_dir>/source.txt + chunks/ + manifest.json
     --as ref            : tài liệu đối chiếu (Level 4) -> <process_dir>/sources/<tên>.txt (không chia chunk)
Output (main):
  <process_dir>/source.txt                 văn bản thuần (nguồn để kiểm chứng trích dẫn)
  <process_dir>/chunks/chunk_NNN.md        từng đoạn ~6000 ký tự, cắt ở ranh giới đoạn văn, có nhãn trang/mục
  <process_dir>/manifest.json              danh sách chunk + vị trí (trang PDF / tiêu đề gần nhất)
  <process_dir>/coverage.md                bảng độ phủ để chép vào mục 0 của báo cáo
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path


def read_source(path: Path) -> list[tuple[str, str]]:
    """Trả về [(nhãn vị trí, văn bản)]; PDF theo trang, file khác 1 khối."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
        doc = fitz.open(str(path))
        pages = [(f"tr.{i + 1}", p.get_text("text")) for i, p in enumerate(doc)]
        empty = sum(1 for _, t in pages if len(t.strip()) < 30)
        if pages and empty / len(pages) > 0.5:
            print(f"[WARN] {empty}/{len(pages)} trang không có lớp chữ (PDF scan). Chạy skill 'boc-tach-pdf' trước, "
                  f"rồi ingest file .md kết quả.")
        return pages
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return [("", "\n".join(parts))]
    if ext in (".html", ".htm"):
        raw = path.read_text(encoding="utf-8", errors="ignore")
        raw = re.sub(r"(?is)<(script|style|nav|footer|header)[^>]*>.*?</\1>", " ", raw)
        raw = re.sub(r"(?i)<br\s*/?>|</p>|</h\d>|</li>|</div>", "\n", raw)
        raw = re.sub(r"<[^>]+>", " ", raw)
        import html
        return [("", html.unescape(raw))]
    return [("", path.read_text(encoding="utf-8", errors="ignore"))]


def clean(text: str) -> str:
    text = unicodedata.normalize("NFC", text).replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # nối từ bị ngắt dòng ở PDF tiếng Anh
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def main() -> int:
    ap = argparse.ArgumentParser(description="Lưu nguồn cục bộ + chia chunk cho Đọc Sâu")
    ap.add_argument("source")
    ap.add_argument("--process-dir", required=True)
    ap.add_argument("--chunk-chars", type=int, default=6000)
    ap.add_argument("--as", dest="role", choices=["main", "ref"], default="main")
    a = ap.parse_args()

    src = Path(a.source).expanduser()
    if re.match(r"^https?://", a.source):
        print("[FAIL] Không nhận URL. Tải nội dung trang, lưu thành file trong <process_dir>/sources/ rồi chạy lại.")
        return 2
    if not src.is_file():
        print(f"[FAIL] Không thấy file {src}")
        return 2
    pdir = Path(a.process_dir).expanduser()
    pdir.mkdir(parents=True, exist_ok=True)
    parts = [(loc, clean(t)) for loc, t in read_source(src)]
    full = "\n\n".join(t for _, t in parts if t)
    if len(full) < 200:
        print("[FAIL] Văn bản trích được quá ngắn (< 200 ký tự). Kiểm tra file nguồn (PDF scan? trang yêu cầu đăng nhập?).")
        return 2

    if a.role == "ref":
        out = pdir / "sources" / (re.sub(r"[^\w.-]+", "_", src.stem) + ".txt")
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(full, encoding="utf-8")
        print(f"[OK] Nguồn đối chiếu: {out} ({len(full):,} ký tự)")
        return 0

    (pdir / "source.txt").write_text(full, encoding="utf-8")
    # chia chunk theo đoạn văn, giữ nhãn trang/tiêu đề
    chunks, cur, cur_len, cur_loc, heading = [], [], 0, None, ""
    for loc, text in parts:
        for para in re.split(r"\n\s*\n", text):
            para = para.strip()
            if not para:
                continue
            if re.match(r"^(#{1,6}\s|chương\s|chapter\s|phần\s|part\s|\d+(\.\d+)*\.?\s+[A-ZÀ-Ỹ])", para, re.I) and len(para) < 120:
                heading = para.lstrip("# ").strip()
            if cur_loc is None:
                cur_loc = (loc, heading)
            if cur and cur_len + len(para) > a.chunk_chars:
                chunks.append({"text": "\n\n".join(cur), "start": cur_loc, "end": (loc, heading)})
                cur, cur_len, cur_loc = [], 0, (loc, heading)
            cur.append(para)
            cur_len += len(para)
    if cur:
        chunks.append({"text": "\n\n".join(cur), "start": cur_loc, "end": (parts[-1][0], heading)})

    cdir = pdir / "chunks"
    cdir.mkdir(exist_ok=True)
    for old in cdir.glob("chunk_*.md"):
        old.unlink()
    manifest = []
    for i, c in enumerate(chunks, 1):
        locs = [c["start"][0]] + ([c["end"][0]] if c["end"][0] != c["start"][0] else [])
        where = " - ".join(x for x in locs if x) or "-"
        label = c["start"][1] or "(đầu tài liệu)"
        (cdir / f"chunk_{i:03d}.md").write_text(
            f"<!-- chunk {i}/{len(chunks)} | vị trí: {where} | mục: {label} -->\n\n{c['text']}\n", encoding="utf-8")
        manifest.append({"chunk": i, "file": f"chunks/chunk_{i:03d}.md", "location": where, "section": label,
                         "chars": len(c["text"])})
    (pdir / "manifest.json").write_text(json.dumps({"source_file": str(src), "source_txt": str(pdir / "source.txt"),
                                                     "total_chars": len(full), "chunks": manifest},
                                                    ensure_ascii=False, indent=2), encoding="utf-8")
    rows = ["| Chunk | Vị trí | Mục | Trạng thái | Ý chính (1 câu) |", "|---|---|---|---|---|"]
    rows += [f"| {m['chunk']} | {m['location']} | {m['section'][:50]} | Chưa đọc | |" for m in manifest]
    (pdir / "coverage.md").write_text("\n".join(rows) + "\n", encoding="utf-8")
    print(f"[OK] {len(full):,} ký tự -> {len(chunks)} chunk tại {cdir}")
    print(f"[OK] Nguồn kiểm chứng trích dẫn: {pdir / 'source.txt'}")
    print(f"[OK] Bảng độ phủ: {pdir / 'coverage.md'} (cập nhật 'Đã đọc' + ý chính sau mỗi chunk)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
