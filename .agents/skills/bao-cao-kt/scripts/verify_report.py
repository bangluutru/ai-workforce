#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
verify_report.py - Quality Gate tự động cho file Excel của bao-cao-kt.

1. Recalc bằng LibreOffice headless (tự tìm soffice kể cả /Applications/LibreOffice.app).
2. FAIL nếu có ô lỗi (#REF!, #DIV/0!, #NAME?, #VALUE!, #N/A, Err:...).
3. FAIL nếu ô KPI trên Dashboard rỗng hoặc bằng 0.
4. FAIL nếu ô công thức không có giá trị sau recalc (công thức hỏng).
5. Đối chiếu chéo (nếu có -i data.json): mọi mã chỉ tiêu tính bằng Python phải khớp Excel (sai số <= 1 đồng).
6. Kiểm tra dòng "Chênh lệch với P&L" của bảng theo kỳ phải = 0 (lệch -> FAIL, trừ khi --allow-recon-diff
   SAU KHI đã gắn cờ [CẦN XÁC MINH] trong báo cáo).
7. --render-dir: xuất PDF + PNG từng trang (xlsx và pptx nếu có) để Agent MỞ ẢNH RA XEM.

Mã thoát: 0 = PASS, 1 = FAIL, 2 = lỗi đầu vào / thiếu LibreOffice.
"""

import argparse
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from office_utils import recalc_xlsx, render_png  # noqa: E402
from pnl_model import DataError, compute, load_data  # noqa: E402

ERR_PREFIX = ("#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "Err:")


def _resolve(wb, ref):
    sheet, cell = ref.rsplit("!", 1)
    return wb[sheet.strip("'")][cell].value


def main():
    ap = argparse.ArgumentParser(description="Quality Gate cho Excel bao-cao-kt")
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--input", "-i", help="data.json để đối chiếu chéo Python vs Excel")
    ap.add_argument("--pptx", help="Deck slide cần render để soát bằng mắt")
    ap.add_argument("--render-dir", help="Thư mục xuất PNG review")
    ap.add_argument("--allow-zero-kpi", action="store_true", help="Chỉ dùng khi KPI = 0 là số thật đã xác nhận")
    ap.add_argument("--allow-recon-diff", action="store_true",
                    help="Chỉ dùng khi đã gắn cờ [CẦN XÁC MINH] cho chênh lệch bảng theo kỳ")
    args = ap.parse_args()

    src = Path(args.xlsx).expanduser()
    if not src.is_file():
        print(f"❌ Không tìm thấy {src}", file=sys.stderr)
        sys.exit(2)
    try:
        recalced = recalc_xlsx(src)
    except RuntimeError as e:
        print(f"❌ {e}", file=sys.stderr)
        sys.exit(2)

    wf = openpyxl.load_workbook(src)
    wv = openpyxl.load_workbook(recalced, data_only=True)
    fails, warns = [], []

    for ws in wv.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith(ERR_PREFIX):
                    fails.append(f"Ô lỗi {ws.title}!{c.coordinate} = {v}")
                f = wf[ws.title][c.coordinate].value
                if isinstance(f, str) and f.startswith("=") and v is None and '""' not in f:
                    fails.append(f"Công thức không ra giá trị {ws.title}!{c.coordinate}: {f}")

    if "_map" not in wv.sheetnames:
        fails.append("Thiếu sheet _map (file không do build_excel_dashboard.py tạo?)")
        mapping = []
    else:
        mapping = [r for r in wv["_map"].iter_rows(min_row=2, values_only=True) if r[0]]

    kpi_seen = 0
    for key, ref, kind in mapping:
        v = _resolve(wv, ref)
        if kind == "kpi":
            kpi_seen += 1
            if v in (None, ""):
                fails.append(f"KPI rỗng: {key} ({ref})")
            elif isinstance(v, (int, float)) and v == 0 and not args.allow_zero_kpi:
                fails.append(f"KPI bằng 0: {key} ({ref}) - kiểm tra tham chiếu hoặc số liệu")
        elif kind == "recon":
            if isinstance(v, (int, float)) and abs(v) >= 1:
                msg = f"Bảng theo kỳ lệch P&L: {key} = {v:,.0f}"
                (warns if args.allow_recon_diff else fails).append(msg)
    if mapping and kpi_seen == 0:
        fails.append("Không có KPI nào trong _map")

    if args.input:
        try:
            data = load_data(args.input)
        except DataError as e:
            print(f"❌ LỖI DỮ LIỆU: {e}", file=sys.stderr)
            sys.exit(2)
        model = compute(data)
        checked = 0
        for key, ref, kind in mapping:
            if kind != "pnl":
                continue
            _, code, per = key.split(".")
            exp = model.get(code, {}).get(per)
            got = _resolve(wv, ref)
            if exp is None and got in (None, ""):
                continue
            checked += 1
            if not isinstance(got, (int, float)) or exp is None or abs(got - exp) > (1e-6 if code in ("GM", "NM") else 1):
                fails.append(f"Lệch Python vs Excel tại {key}: Python={exp} Excel={got}")
        print(f"• Đối chiếu chéo Python vs Excel: {checked} ô")

    if args.render_dir:
        rd = Path(args.render_dir).expanduser()
        for f in [src] + ([Path(args.pptx).expanduser()] if args.pptx else []):
            try:
                _, pngs = render_png(f, rd)
                print(f"• Render {f.name}: {len(pngs)} trang -> {rd}")
                for p in pngs:
                    print(f"    {p}")
            except Exception as e:  # render lỗi không che lỗi số liệu
                fails.append(f"Không render được {f.name}: {e}")
        print("  ⚠️  BẮT BUỘC mở từng PNG bằng công cụ xem ảnh, soát: chữ tràn, số ####, biểu đồ cắt trang, KPI trống.")

    for w in warns:
        print(f"⚠️  {w}")
    if fails:
        print(f"❌ QUALITY GATE FAIL ({len(fails)} lỗi):")
        for f in fails:
            print(f"   - {f}")
        sys.exit(1)
    print("✅ QUALITY GATE PASS: không có ô lỗi, KPI có giá trị, công thức khớp mô hình.")


if __name__ == "__main__":
    main()
