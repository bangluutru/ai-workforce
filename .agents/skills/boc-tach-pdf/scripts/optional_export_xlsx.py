#!/usr/bin/env python3
"""
optional_export_xlsx.py — Xuất TẤT CẢ bảng Markdown trong MERGED.md ra Excel (mỗi bảng 1 sheet).

- Số kiểu Việt Nam (mặc định --locale vi): '.' = phân cách nghìn, ',' = thập phân
      "4.200" -> 4200 ; "3.750,5" -> 3750.5 ; "12.450.000.000" -> 12450000000 ; "15%" -> 0.15
  --locale en: ',' = nghìn, '.' = thập phân.
- Ô có [CẦN XÁC MINH] hoặc không parse được giữ nguyên dạng chữ và được tô vàng.
- Dòng "Tổng / Tổng cộng / Cộng / Total / 合計": ghi CÔNG THỨC SỐNG =SUM(...) cho các cột số,
  đồng thời so với con số OCR trên giấy; lệch -> ghi chú [CẦN XÁC MINH] ở cột cuối + tô đỏ.
- Bảng trải qua nhiều trang (cùng tiêu đề cột) được gộp thành một sheet.

Usage:
    python3 optional_export_xlsx.py <MERGED.md> [--out <file.xlsx>] [--locale vi|en]
Exit: 0 = OK; 3 = có tổng lệch hoặc ô cần xác minh (vẫn xuất file); 1 = không có bảng.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TOTAL_RE = re.compile(r"^\**\s*(tổng\s*cộng|tổng\s*số|tổng|cộng|total|合計|小計)\b", re.I)
SEP_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
VI_NUM = re.compile(r"^[-+]?\d{1,3}(\.\d{3})*(,\d+)?$|^[-+]?\d+(,\d+)?$")
EN_NUM = re.compile(r"^[-+]?\d{1,3}(,\d{3})*(\.\d+)?$|^[-+]?\d+(\.\d+)?$")


def parse_number(raw: str, locale: str = "vi"):
    """Trả về (giá_trị, là_phần_trăm) hoặc (None, False) nếu không phải số chắc chắn."""
    s = raw.strip().replace("**", "").replace(" ", " ").strip()
    if not s or "CẦN XÁC MINH" in s:
        return None, False
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(" ", "")
    pct = s.endswith("%")
    s = s.rstrip("%")
    if locale == "vi":
        if not VI_NUM.match(s):
            return None, False
        val = float(s.replace(".", "").replace(",", "."))
    else:
        if not EN_NUM.match(s):
            return None, False
        val = float(s.replace(",", ""))
    if neg:
        val = -val
    if pct:
        val = val / 100.0
    if val.is_integer() and not pct:
        val = int(val)
    return val, pct


def split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", line)]


def find_tables(md: str) -> list[dict]:
    tables, cur, page = [], None, None
    lines = md.splitlines()
    for i, raw in enumerate(lines):
        m = re.match(r"<!--\s*page:\s*(\d+)\s*-->", raw.strip())
        if m:
            page = int(m.group(1))
        s = raw.strip()
        if s.startswith("|") and i + 1 < len(lines) and SEP_RE.match(lines[i + 1].strip()) and cur is None:
            cur = {"header": split_row(s), "rows": [], "page": page, "line": i + 1}
            continue
        if cur is not None:
            if SEP_RE.match(s):
                continue
            if s.startswith("|"):
                cur["rows"].append(split_row(s))
                continue
            tables.append(cur); cur = None
    if cur is not None:
        tables.append(cur)
    # gộp bảng nối trang (cùng header, liên tiếp)
    merged = []
    for t in tables:
        if merged and merged[-1]["header"] == t["header"] and t["page"] and merged[-1]["page"] and t["page"] - merged[-1]["page"] <= 1:
            merged[-1]["rows"] += t["rows"]
        else:
            merged.append(t)
    return merged


def export(md_path: Path, out: Path, locale: str) -> int:
    md = md_path.read_text(encoding="utf-8")
    tables = find_tables(md)
    if not tables:
        print("[FAIL] Không tìm thấy bảng Markdown nào.")
        return 1
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    warn_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    bad_fill = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
    issues = 0

    for ti, t in enumerate(tables, 1):
        ws = wb.create_sheet(f"Bang_{ti}" + (f"_tr{t['page']}" if t["page"] else ""))
        ncol = max([len(t["header"])] + [len(r) for r in t["rows"]])
        header = t["header"] + [""] * (ncol - len(t["header"]))
        ws.append(header + ["Ghi chú kiểm tra"])
        for c in range(1, ncol + 2):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True); cell.fill = hdr_fill; cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        block_start = 2
        for r_i, row in enumerate(t["rows"], start=2):
            row = row + [""] * (ncol - len(row))
            label = " ".join(row[:3])
            is_total = any(TOTAL_RE.match(c.strip()) for c in row[:3] if c.strip())
            note = []
            for c_i, raw in enumerate(row, start=1):
                cell = ws.cell(row=r_i, column=c_i)
                cell.border = border
                val, pct = parse_number(raw, locale)
                if is_total and val is not None and r_i > block_start:
                    col = get_column_letter(c_i)
                    above = [ws.cell(row=k, column=c_i).value for k in range(block_start, r_i)]
                    nums = [v for v in above if isinstance(v, (int, float))]
                    cell.value = f"=SUM({col}{block_start}:{col}{r_i - 1})"
                    cell.number_format = "#,##0.##"
                    calc = sum(nums)
                    if nums and abs(calc - val) > max(1e-6, abs(val) * 1e-9):
                        note.append(f"[CẦN XÁC MINH] cột {header[c_i - 1] or col}: giấy ghi {raw} nhưng SUM = {calc:,.4g}")
                        cell.fill = bad_fill
                        cell.comment = Comment(f"Bản scan ghi: {raw}", "boc-tach-pdf")
                    continue
                if val is not None:
                    cell.value = val
                    cell.number_format = "0.0%" if pct else ("#,##0" if isinstance(val, int) else "#,##0.##")
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.value = raw.replace("**", "")
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    if "CẦN XÁC MINH" in raw or (re.search(r"\d", raw) and re.fullmatch(r"[\d.,%\s()+-]+", raw.strip() or "x")):
                        cell.fill = warn_fill
                        note.append(f"Ô {get_column_letter(c_i)}{r_i}: số không đúng định dạng {locale} -> giữ dạng chữ")
            if is_total:
                block_start = r_i + 1
                for c_i in range(1, ncol + 1):
                    ws.cell(row=r_i, column=c_i).font = Font(bold=True)
            if note:
                issues += len(note)
                ws.cell(row=r_i, column=ncol + 1, value="; ".join(note)).fill = warn_fill
            _ = label
        for c in range(1, ncol + 2):
            width = max(len(str(ws.cell(row=r, column=c).value or "")) for r in range(1, ws.max_row + 1))
            ws.column_dimensions[get_column_letter(c)].width = min(60, max(10, width + 2))
        ws.freeze_panes = "A2"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"[OK] {len(tables)} bảng -> {out}")
    if issues:
        print(f"[WARN] {issues} vấn đề cần xác minh (xem cột 'Ghi chú kiểm tra').")
        return 3
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="Xuất mọi bảng Markdown ra Excel (số VN, công thức SUM sống)")
    ap.add_argument("merged_md")
    ap.add_argument("--out", default=None)
    ap.add_argument("--locale", choices=["vi", "en"], default="vi")
    a = ap.parse_args()
    md = Path(a.merged_md)
    if not md.exists():
        print(f"[FAIL] Không tìm thấy {md}")
        return 1
    out = Path(a.out) if a.out else md.with_suffix(".xlsx")
    return export(md, out, a.locale)


if __name__ == "__main__":
    sys.exit(main())
