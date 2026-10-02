"""
02_structure.py — Layer 2: dựng lại khối header văn bản hành chính NĐ 30/2020 thành bảng 2 cột không viền.

Thiết kế an toàn (v4):
  - Đọc header theo TỪNG DÒNG (00_pandoc giữ ngắt dòng cứng trong khối header) và giữ nguyên in đậm của Agent.
  - Nhận diện tổng quát: cơ quan chủ quản + cơ quan ban hành (UBND TỈNH .../SỞ ..., BỘ .../CỤC ...),
    quốc hiệu/tiêu ngữ (mọi biến thể đặt dấu), số ký hiệu, "<địa danh bất kỳ>, ngày ... tháng ... năm ...",
    V/v (công văn), tên loại (TỜ TRÌNH, NGHỊ QUYẾT, QUYẾT ĐỊNH, ... ), trích yếu, Kính gửi.
  - KHÔNG BAO GIỜ XÓA NỘI DUNG: nếu trong vùng header có dòng không nhận diện được -> BỎ QUA tái cấu trúc
    (giữ nguyên DOCX) và báo rõ dòng đó. Sau khi dựng, kiểm tra mọi dòng cũ đều có mặt trong header mới.
"""
import re
import sys
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from ooxml_helpers import (add_separator_line, apply_table_grid, apply_tcW, make_run, normalize_table_xml,
                           remove_table_borders, set_cell_no_border, set_table_fixed_width)

BODY_FONT = "Times New Roman"
TITLES = ("TO TRINH", "QUYET DINH", "NGHI QUYET", "NGHI DINH", "THONG TU", "CHI THI", "QUY CHE", "QUY DINH",
          "THONG CAO", "THONG BAO", "HUONG DAN", "CHUONG TRINH", "KE HOACH", "PHUONG AN", "DE AN", "DU AN",
          "BAO CAO", "BIEN BAN", "HOP DONG", "CONG DIEN", "BAN GHI NHO", "BAN THOA THUAN", "GIAY", "PHIEU",
          "THU CONG", "LUAT", "PHAP LENH")
PARENT_AGENCY = re.compile(r"^(UY BAN NHAN DAN|UBND|BO |TONG CUC|CHINH PHU|HOI DONG NHAN DAN|HDND|TINH UY|THANH UY|"
                           r"DAI HOC|TAP DOAN|TONG CONG TY)")


def base(text):
    text = unicodedata.normalize("NFC", text).replace("đ", "d").replace("Đ", "D")
    text = "".join(c for c in unicodedata.normalize("NFD", text) if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", text).strip().upper()


def is_upper_line(s):
    letters = [c for c in s if c.isalpha()]
    return bool(letters) and sum(1 for c in letters if c.isupper()) / len(letters) >= 0.85


def para_lines(para):
    """[(text, bold, italic)] theo từng dòng (tách tại <w:br/>)."""
    lines, cur, b, it = [], [], False, False
    for run in para.runs:
        parts = run.text.split("\n")
        for i, part in enumerate(parts):
            if i > 0:
                lines.append(("".join(cur), b, it)); cur, b, it = [], False, False
            if part.strip():
                cur.append(part)
                b = b or bool(run.bold)
                it = it or bool(run.italic)
            else:
                cur.append(part)
    lines.append(("".join(cur), b, it))
    return [(re.sub(r"\s+", " ", t).strip(), bb, ii) for t, bb, ii in lines if t.strip()]


def classify(text, state):
    b = base(text)
    if re.search(r"CONG HOA XA HOI CHU NGHIA", b) or (b == "VIET NAM" and state.get("last") == "qh"):
        return "qh"
    if re.search(r"DOC LAP\s*[-–—]?\s*TU DO\s*[-–—]?\s*HANH PHUC", b):
        return "tn"
    if re.match(r"^SO\s*:", b):
        return "so"
    if re.search(r",\s*NGAY\s+\S+\s+THANG\s+\S+\s+NAM\s+\d{4}", b):
        return "date"
    if re.match(r"^(V/V|VE VIEC)\b", b) and not state.get("title"):
        return "vv"
    if re.match(r"^KINH GUI\b", b):
        return "kg"
    if is_upper_line(text) and any(b == t or b.startswith(t + " ") for t in TITLES) and len(b) < 60 \
            and state.get("agency"):
        return "title"
    if state.get("title") and not state.get("kg"):
        if is_upper_line(text) and state.get("trich_yeu"):
            return "authority"  # VD: "CHỦ TỊCH ỦY BAN NHÂN DÂN TỈNH ..." dưới trích yếu của Quyết định
        if len(text) < 250 and (state.get("_bold") or re.match(r"^(VE|V/V)\b", b)) and state.get("last") in ("title", "trich_yeu"):
            return "trich_yeu"
    if is_upper_line(text) and len(text) <= 80 and not state.get("title"):
        return "agency"
    return None


def build_nd30_header(docx_path):
    doc = Document(str(docx_path))
    paras = doc.paragraphs
    zone, info, state = [], {"agency": [], "qh": [], "tn": [], "so": [], "date": [], "vv": [], "title": [],
                              "trich_yeu": [], "authority": [], "kg": []}, {}
    unknown = []
    for idx, para in enumerate(paras[:30]):
        text = para.text.strip()
        if not text:
            zone.append(idx)
            continue
        b = base(text)
        if re.match(r"^(CAN CU|THUC HIEN)\b", b) or (len(text) > 300 and state.get("title")):
            break
        lines = para_lines(para)
        state["_bold"] = lines[0][1] if lines else False
        if lines and classify(lines[0][0], dict(state)) is None and (
                state.get("title") or (state.get("qh") and (state.get("date") or state.get("so")))):
            break  # đoạn thân bài đầu tiên: hết vùng header
        kinds = []
        for t, bold, ital in lines:
            state["_bold"] = bold
            k = classify(t, state)
            kinds.append(k)
            if k is None:
                unknown.append(t)
                continue
            info[k].append((t, bold, ital))
            state["last"] = k
            state[k] = True
        zone.append(idx)
        if "kg" in kinds:
            break
        if None in kinds:
            break

    if not info["qh"] or not info["agency"]:
        print("[INFO] Không đủ thông tin header NĐ 30 (thiếu quốc hiệu hoặc tên cơ quan). Giữ nguyên.")
        doc.save(str(docx_path))
        return False
    if unknown:
        print("[WARN] Vùng header có dòng KHÔNG nhận diện được -> GIỮ NGUYÊN (không tái cấu trúc để tránh mất chữ):")
        for u in unknown[:5]:
            print(f"    • {u[:100]}")
        doc.save(str(docx_path))
        return False

    old_text = " ".join(paras[i].text for i in zone)
    paras_text = [t for i in zone for t, _, _ in para_lines(paras[i])]
    first_after = paras[zone[-1] + 1]._element if zone[-1] + 1 < len(paras) else None
    body = doc.element.body
    anchor = paras[zone[0]]._element
    placeholder = anchor.makeelement(anchor.tag, {})
    anchor.addprevious(placeholder)
    for i in zone:
        body.remove(paras[i]._element)

    # ---- bảng header 2 cột
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(table)
    for r in range(2):
        for c in range(2):
            set_cell_no_border(table.cell(r, c))
    set_table_fixed_width(table, pct="5000")
    left_w, right_w = 3969, 5670
    apply_table_grid(table, [left_w, right_w])

    def cell_par(cell, align, before=0, after=0):
        p = cell.paragraphs[0]
        p.paragraph_format.alignment = align
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        return p

    agency = info["agency"]
    if not any(bd for _, bd, _ in agency):  # Agent không đánh dấu đậm: áp quy tắc NĐ 30
        if len(agency) >= 2 and PARENT_AGENCY.match(base(agency[0][0])) and not PARENT_AGENCY.match(base(agency[-1][0])):
            agency = [(t, False, i) for t, _, i in agency[:-1]] + [(agency[-1][0], True, agency[-1][2])]
        else:
            agency = [(t, True, i) for t, _, i in agency]
    c = table.cell(0, 0); apply_tcW(c, left_w)
    p = cell_par(c, WD_ALIGN_PARAGRAPH.CENTER)
    for k, (t, bd, _) in enumerate(agency):
        if k:
            p.add_run("\n")
        make_run(p, t, bold=bd, size=13 if bd else 12)
    add_separator_line(c, width_cm=3.0)

    c = table.cell(0, 1); apply_tcW(c, right_w)
    p = cell_par(c, WD_ALIGN_PARAGRAPH.CENTER)
    make_run(p, " ".join(t for t, _, _ in info["qh"]), bold=True, size=12)
    tn = info["tn"][0][0] if info["tn"] else ""
    if tn:
        p.add_run("\n")
        make_run(p, tn, bold=True, size=13)
    add_separator_line(c, width_cm=5.0)

    c = table.cell(1, 0); apply_tcW(c, left_w)
    p = cell_par(c, WD_ALIGN_PARAGRAPH.CENTER, before=6)
    for k, (t, _, _) in enumerate(info["so"]):
        if k:
            p.add_run("\n")
        make_run(p, t, size=13)
    for t, _, _ in info["vv"]:
        p.add_run("\n")
        make_run(p, t, size=12)

    c = table.cell(1, 1); apply_tcW(c, right_w)
    p = cell_par(c, WD_ALIGN_PARAGRAPH.RIGHT, before=6)
    for k, (t, _, _) in enumerate(info["date"]):
        if k:
            p.add_run("\n")
        make_run(p, t, italic=True, size=13)
    normalize_table_xml(table)

    new_elems = [table._tbl]

    def new_par(text, bold=False, size=13, before=0, after=6):
        np_ = doc.add_paragraph()
        np_.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        np_.paragraph_format.space_before = Pt(before)
        np_.paragraph_format.space_after = Pt(after)
        np_.paragraph_format.first_line_indent = Cm(0)
        make_run(np_, text, bold=bold, size=size)
        new_elems.append(np_._element)
        return np_

    if info["title"]:
        new_par(" ".join(t for t, _, _ in info["title"]), bold=True, size=14, before=18, after=0)
        if info["trich_yeu"]:
            new_par(" ".join(t for t, _, _ in info["trich_yeu"]), bold=True, size=14, after=0)
        sep = doc.add_paragraph()
        sep.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sep.paragraph_format.first_line_indent = Cm(0)
        sep.paragraph_format.space_after = Pt(6)
        r = sep.add_run("─" * 14); r.font.size = Pt(8); r.font.color.rgb = RGBColor(0, 0, 0)
        new_elems.append(sep._element)
        for t, _, _ in info["authority"]:
            new_par(t, bold=True, size=14, before=6, after=6)
    for t, _, _ in info["kg"]:
        new_par(t, size=14, before=6, after=12)

    for el in new_elems:
        placeholder.addprevious(el)
    body.remove(placeholder)

    # ---- kiểm tra bảo toàn nội dung
    new_text = " ".join(
        "".join(n.text or "" for n in el.iter() if n.tag.endswith("}t")) for el in new_elems)
    norm = lambda s: re.sub(r"\s+", "", s)  # noqa: E731
    lost = [t for k in info for t, _, _ in info[k] if norm(t) not in norm(new_text)]
    lost += [ln for ln in (paras_text or []) if norm(ln) and norm(ln) not in norm(new_text)]
    if lost:
        print(f"[FAIL] Tái cấu trúc làm mất chữ: {lost[:3]} -> hủy, giữ nguyên DOCX.")
        return False
    doc.save(str(docx_path))
    print(f"[OK] Header NĐ 30 đã tái cấu trúc ({len(zone)} đoạn cũ; cơ quan: {len(agency)} dòng; "
          f"tên loại: {bool(info['title'])}; trích yếu: {bool(info['trich_yeu'])}; V/v: {bool(info['vv'])}).")
    _ = (old_text, first_after)
    return True


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python3 02_structure.py <input_docx>")
        sys.exit(1)
    print("[INFO] Bước 2: Tái cấu trúc Header NĐ 30...")
    build_nd30_header(Path(sys.argv[1]))
    print("[OK] Xong Layer 2.")
