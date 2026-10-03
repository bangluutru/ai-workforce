#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
export_tax_sheet.py — Xuất bảng tính thuế TNCN Excel bằng CÔNG THỨC SỐNG, theo đúng năm thuế.

Nguyên tắc (sửa lỗi bản cũ từng tự điền "đã khấu trừ 10 triệu", "thu nhập năm = Gross x 12", "BHXH 11+12 năm"):
  • KHÔNG có số liệu mặc định của người dùng. Sheet nào thiếu dữ liệu đầu vào thì KHÔNG tạo (hoặc báo lỗi).
  • Mọi định mức (giảm trừ, trần bảo hiểm, biểu thuế) lấy từ tax_calculator.RULES của năm --year và đặt ở
    sheet "Thông số" — công thức các sheet khác tham chiếu tới đó, sửa 1 ô là cả bảng cập nhật.
  • Đơn vị y tế/giáo dục/hưu trí giống tax_calculator.py: --medical/--education/--pension là số BÌNH QUÂN THÁNG.
  • Ngưỡng kết luận quyết toán giống tax_calculator.calculate_annual_settlement.

    python3 export_tax_sheet.py --output <xlsx> --year 2026 --region I --gross 40000000 --dependents 2
    python3 export_tax_sheet.py --output <xlsx> --year 2025 --dependents 1 \
        --settlement-income 300000000 --settlement-insurance 31500000 --tax-withheld 9000000
    python3 export_tax_sheet.py --output <xlsx> --include-bhxh --bhxh-before-2014 3 --bhxh-from-2014 6 --mbqtl 9000000 --months-off 14
"""

import argparse
import os
import sys
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import tax_calculator as tc  # noqa: E402

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:
    print("Lỗi: thiếu openpyxl. Cài: pip3 install openpyxl", file=sys.stderr)
    sys.exit(2)

BIG = 1e15  # "vô cực" cho bậc thuế cuối trong công thức Excel

F = "Segoe UI"
font_title = Font(name=F, size=15, bold=True, color="1B365D")
font_sub = Font(name=F, size=9, italic=True, color="555555")
font_section = Font(name=F, size=11, bold=True, color="FFFFFF")
font_bold = Font(name=F, size=10, bold=True, color="000000")
font_reg = Font(name=F, size=10, color="1C1C1C")
font_input = Font(name=F, size=10, bold=True, color="002060")
font_formula = Font(name=F, size=10, bold=True, color="006100")
font_total = Font(name=F, size=11, bold=True, color="9C0006")
fill_section = PatternFill(start_color="2E5B88", end_color="2E5B88", fill_type="solid")
fill_result = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
fill_input = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
fill_hl = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
fill_total = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
thin = Side(border_style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
border_total = Border(top=thin, bottom=Side(border_style="double", color="1B365D"))
VND = '#,##0 "VNĐ"'
NUM = '#,##0'
PCT = '0.0%'


class Sheet:
    """Ghi tuần tự từng dòng nhãn | giá trị | ghi chú, nhớ địa chỉ ô theo khoá để công thức tham chiếu."""

    def __init__(self, ws, title, subtitle):
        self.ws, self.r, self.ref = ws, 4, {}
        ws["A1"], ws["A1"].font = title, font_title
        ws["A2"], ws["A2"].font = subtitle, font_sub
        for col, w in zip("ABCD", (52, 24, 60, 4)):
            ws.column_dimensions[col].width = w
        ws.sheet_properties.pageSetUpPr.fitToPage = True   # in/PDF: vừa 1 trang ngang, không tách cột ghi chú
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.page_setup.orientation = "landscape"

    def section(self, text, result=False):
        self.r += 1 if self.r > 4 else 0
        ws = self.ws
        ws.cell(row=self.r, column=1, value=text).font = font_section
        ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=3)
        for c in range(1, 4):
            ws.cell(row=self.r, column=c).fill = fill_result if result else fill_section
        self.r += 1

    def row(self, key, label, value, note="", kind="formula", fmt=VND):
        ws, r = self.ws, self.r
        ws.cell(row=r, column=1, value=label).font = font_bold if kind in ("input", "total") else font_reg
        c = ws.cell(row=r, column=2, value=value)
        c.number_format, c.alignment = fmt, Alignment(horizontal="right")
        c.border = border_total if kind == "total" else border
        if kind == "input":
            c.font, c.fill = font_input, fill_input
        elif kind == "total":
            c.font, c.fill = font_total, fill_total
        else:
            c.font = font_formula
        if fmt == "@":   # kết luận dạng chữ: trải sang cột C để không bị cắt
            c.alignment = Alignment(horizontal="left")
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        elif note:
            ws.cell(row=r, column=3, value=note).font = font_sub
        if key:
            self.ref[key] = f"$B${r}"
        self.r += 1
        return f"B{r}"


def build_params(wb, year, region):
    """Sheet 'Thông số': định mức của năm thuế. Trả về dict khoá → tham chiếu tuyệt đối."""
    R = tc.RULES[year]
    ws = wb.active
    ws.title = "Thông số"
    s = Sheet(ws, f"THÔNG SỐ PHÁP LÝ ÁP DỤNG: NĂM THUẾ {year}",
              f"Căn cứ: {R['legal']}. Ô vàng là định mức; nếu văn bản mới thay đổi, sửa tại đây (đối chiếu văn bản gốc).")
    s.section("A. GIẢM TRỪ GIA CẢNH VÀ KHOẢN GIẢM TRỪ KHÁC")
    s.row("personal", "Giảm trừ bản thân (VNĐ/tháng)", R["personal"], "", "input")
    s.row("dependent", "Giảm trừ mỗi người phụ thuộc (VNĐ/tháng)", R["dependent"], "", "input")
    s.row("pension_max", "Hưu trí tự nguyện tối đa (VNĐ/tháng)", R["pension_max_monthly"], "", "input")
    s.row("medical_max", "Chi phí y tế được trừ tối đa (VNĐ/năm)", R["medical_max_annual"],
          "0 = năm thuế này chưa có khoản giảm trừ y tế" if not R["medical_max_annual"] else "", "input")
    s.row("education_max", "Chi phí giáo dục được trừ tối đa (VNĐ/năm)", R["education_max_annual"],
          "0 = năm thuế này chưa có khoản giảm trừ giáo dục" if not R["education_max_annual"] else "", "input")
    s.section("B. BẢO HIỂM BẮT BUỘC (phần người lao động)")
    s.row("ins_cap", "Trần lương đóng BHXH, BHYT (VNĐ/tháng)", R["insurance_cap"],
          "20 x lương cơ sở. Nếu lương cơ sở đổi giữa năm, dùng mức của tháng tính", "input")
    if region:
        s.row("bhtn_cap", f"Trần lương đóng BHTN vùng {region} (VNĐ/tháng)", tc.REGION_MIN_WAGE[year][region] * 20,
              f"20 x lương tối thiểu vùng {region} ({tc.REGION_MIN_WAGE[year][region]:,.0f}đ)", "input")
    else:
        s.row("bhtn_cap", "Trần lương đóng BHTN (VNĐ/tháng)", 0,
              "[CẦN XÁC MINH] chưa biết vùng → 0 = tính BHTN trên toàn bộ lương", "input")
    s.row("r_bhxh", "Tỷ lệ BHXH", tc.RATE_BHXH, "", "input", PCT)
    s.row("r_bhyt", "Tỷ lệ BHYT", tc.RATE_BHYT, "", "input", PCT)
    s.row("r_bhtn", "Tỷ lệ BHTN", tc.RATE_BHTN, "", "input", PCT)
    s.section(f"C. BIỂU THUẾ LŨY TIẾN TỪNG PHẦN THEO THÁNG ({len(R['rates'])} bậc)")
    lower = 0
    for i, (lim, rate) in enumerate(zip(R["limits"], R["rates"]), start=1):
        s.row(f"lo{i}", f"Bậc {i}: cận dưới (VNĐ/tháng)", lower, "", "input")
        s.row(f"hi{i}", f"Bậc {i}: cận trên (VNĐ/tháng)", BIG if lim == float("inf") else lim,
              "bậc cuối: không giới hạn" if lim == float("inf") else "", "input")
        s.row(f"rate{i}", f"Bậc {i}: thuế suất", rate, "", "input", PCT)
        lower = lim
    s.ref["n_brackets"] = len(R["rates"])
    return {k: f"'Thông số'!{v}" if isinstance(v, str) else v for k, v in s.ref.items()}


def bracket_rows(s, P, tntt, annual):
    """Một dòng/bậc: =MAX(0, MIN(TNTT, cận trên x k) - cận dưới x k) x thuế suất (k = 12 nếu tính năm)."""
    k = "*12" if annual else ""
    cells = []
    for i in range(1, P["n_brackets"] + 1):
        cells.append(s.row(None, f"Bậc {i}", f"=MAX(0,MIN({tntt},{P[f'hi{i}']}{k})-{P[f'lo{i}']}{k})*{P[f'rate{i}']}"))
    return cells


def build_monthly(wb, P, a):
    s = Sheet(wb.create_sheet("Lương tháng"), f"THUẾ TNCN TIỀN LƯƠNG THÁNG & LƯƠNG NET, NĂM {a.year}",
              "Công thức sống: sửa ô vàng để tính lại. Định mức lấy từ sheet 'Thông số'.")
    s.section("A. ĐẦU VÀO (người dùng cung cấp)")
    g = s.row("g", "Lương Gross tháng", a.gross, "Tổng thu nhập chịu thuế từ tiền lương trong tháng", "input")
    d = s.row("d", "Số người phụ thuộc đã đăng ký", a.dependents, "", "input", NUM)
    m = s.row("m", "Chi phí y tế (bình quân VNĐ/tháng)", a.medical, "Cùng đơn vị với tax_calculator.py --medical", "input")
    e = s.row("e", "Chi phí giáo dục (bình quân VNĐ/tháng)", a.education, "", "input")
    p = s.row("p", "Hưu trí tự nguyện đã đóng (VNĐ/tháng)", a.pension, "", "input")
    s.section("B. BẢO HIỂM BẮT BUỘC NGƯỜI LAO ĐỘNG ĐÓNG")
    b1 = s.row(None, "BHXH", f"=MIN({g},{P['ins_cap']})*{P['r_bhxh']}")
    b2 = s.row(None, "BHYT", f"=MIN({g},{P['ins_cap']})*{P['r_bhyt']}")
    b3 = s.row(None, "BHTN", f"=IF({P['bhtn_cap']}>0,MIN({g},{P['bhtn_cap']}),{g})*{P['r_bhtn']}")
    ins = s.row("ins", "Tổng bảo hiểm", f"=SUM({b1}:{b3})", "", "total")
    s.section("C. GIẢM TRỪ")
    c1 = s.row(None, "Giảm trừ bản thân", f"={P['personal']}")
    s.row(None, "Giảm trừ người phụ thuộc", f"={d}*{P['dependent']}")
    s.row(None, "Giảm trừ y tế", f"=MIN({m},{P['medical_max']}/12)")
    s.row(None, "Giảm trừ giáo dục", f"=MIN({e},{P['education_max']}/12)")
    c5 = s.row(None, "Giảm trừ hưu trí tự nguyện", f"=MIN({p},{P['pension_max']})")
    ded = s.row(None, "Tổng giảm trừ (gồm bảo hiểm)", f"={ins}+SUM({c1}:{c5})", "", "total")
    s.section("D. THUẾ THEO BIỂU LŨY TIẾN TỪNG PHẦN")
    tntt = s.row("tntt", "Thu nhập tính thuế", f"=MAX(0,{g}-{ded})")
    br = bracket_rows(s, P, tntt, annual=False)
    tax = s.row("tax", "THUẾ TNCN THÁNG", f"=ROUND(SUM({br[0]}:{br[-1]}),0)", "", "total")
    s.section("E. KẾT QUẢ", result=True)
    s.row("net", "LƯƠNG NET (Gross - bảo hiểm - thuế)", f"={g}-{ins}-{tax}", "", "total")
    s.row(None, "Thuế / Gross", f"=IF({g}>0,{tax}/{g},0)", "", fmt=PCT)
    return s


def build_settlement(wb, P, a):
    s = Sheet(wb.create_sheet("Quyết toán năm"), f"QUYẾT TOÁN THUẾ TNCN TIỀN LƯƠNG, NĂM {a.year}",
              "Số liệu lấy từ chứng từ khấu trừ / tờ khai của người dùng. Không có ô nào được tự điền.")
    s.section("A. ĐẦU VÀO (theo chứng từ)")
    inc = s.row(None, "Tổng thu nhập chịu thuế cả năm (mọi nơi chi trả)", a.settlement_income, "", "input")
    ins = s.row(None, "Bảo hiểm bắt buộc đã đóng cả năm", a.settlement_insurance, "Theo chứng từ khấu trừ / VssID", "input")
    mon = s.row(None, "Số tháng tính giảm trừ bản thân", a.months, "Quyết toán cả năm: 12", "input", NUM)
    dm = s.row(None, "Tổng số tháng-người phụ thuộc", a.dependent_months if a.dependent_months is not None else a.dependents * a.months,
               "VD 2 NPT đủ năm = 24; NPT đăng ký từ tháng 7 = 6", "input", NUM)
    m = s.row(None, "Chi phí y tế cả năm", a.medical * 12, "", "input")
    e = s.row(None, "Chi phí giáo dục cả năm", a.education * 12, "", "input")
    p = s.row(None, "Hưu trí tự nguyện cả năm", a.pension * 12, "", "input")
    wh = s.row(None, "Thuế TNCN đã tạm khấu trừ/tạm nộp", a.tax_withheld, "Tổng các chứng từ khấu trừ", "input")
    s.section("B. GIẢM TRỪ CẢ NĂM")
    c1 = s.row(None, "Giảm trừ bản thân", f"={P['personal']}*{mon}")
    s.row(None, "Giảm trừ người phụ thuộc", f"={P['dependent']}*{dm}")
    s.row(None, "Giảm trừ y tế", f"=MIN({m},{P['medical_max']})")
    s.row(None, "Giảm trừ giáo dục", f"=MIN({e},{P['education_max']})")
    c5 = s.row(None, "Giảm trừ hưu trí tự nguyện", f"=MIN({p},{P['pension_max']}*{mon})")
    ded = s.row(None, "Tổng giảm trừ (gồm bảo hiểm)", f"={ins}+SUM({c1}:{c5})", "", "total")
    s.section("C. THUẾ CẢ NĂM (biểu tháng x 12)")
    tntt = s.row(None, "Thu nhập tính thuế cả năm", f"=MAX(0,{inc}-{ded})")
    br = bracket_rows(s, P, tntt, annual=True)
    tax = s.row(None, "THUẾ TNCN PHẢI NỘP CẢ NĂM", f"=ROUND(SUM({br[0]}:{br[-1]}),0)", "", "total")
    s.section("D. KẾT LUẬN", result=True)
    diff = s.row(None, "Chênh lệch = đã khấu trừ - phải nộp", f"={wh}-{tax}", "Dương: nộp thừa | Âm: thiếu")
    s.row(None, "KẾT LUẬN", f'=IF({diff}>0,"NỘP THỪA "&TEXT({diff},"#,##0")&"đ: đề nghị hoàn hoặc bù trừ kỳ sau",'
          f'IF({diff}<-50000,"PHẢI NỘP THÊM "&TEXT(-{diff},"#,##0")&"đ",'
          f'IF({diff}<0,"Số nộp thêm ≤ 50.000đ: thuộc diện miễn nộp (đối chiếu điều kiện)","CÂN BẰNG")))', "", "total", "@")
    return s


def build_bhxh(wb, a):
    s = Sheet(wb.create_sheet("BHXH một lần"), "ƯỚC TÍNH BHXH MỘT LẦN & SO SÁNH BẢO LƯU",
              "Luật BHXH số 41/2024/QH15 (hiệu lực 01/07/2025). Ước tính tham khảo; số chính thức do cơ quan BHXH xác định.")
    s.section("A. ĐẦU VÀO (theo sổ BHXH / VssID)")
    y1 = s.row(None, "Số năm đóng trước 2014", a.bhxh_before_2014, "", "input", '0.0')
    y2 = s.row(None, "Số năm đóng từ 2014", a.bhxh_from_2014, "", "input", '0.0')
    ty = s.row(None, "Tổng số năm đóng", f"={y1}+{y2}", "", fmt='0.0')
    mb = s.row(None, "Mức bình quân tiền lương đóng BHXH (đã điều chỉnh trượt giá)", a.mbqtl, "", "input")
    mo = s.row(None, "Số tháng đã dừng đóng", a.months_off, "", "input", NUM)
    s.section("B. ĐIỀU KIỆN RÚT MỘT LẦN THÔNG THƯỜNG")
    s.row(None, "Dừng đóng đủ 12 tháng", f'=IF({mo}>=12,"ĐẠT","CHƯA ĐẠT")', "", fmt="@")
    s.row(None, "Tổng thời gian đóng dưới 20 năm", f'=IF({ty}<20,"ĐẠT","KHÔNG ĐẠT")',
          "Áp dụng cho người tham gia trước 01/07/2025", fmt="@")
    s.section("C. SỐ TIỀN NẾU RÚT MỘT LẦN")
    s1 = s.row(None, "Số tháng hưởng giai đoạn trước 2014 (1,5 tháng/năm)", f"={y1}*1.5", "", fmt='0.0')
    s2 = s.row(None, "Số tháng hưởng giai đoạn từ 2014 (2 tháng/năm)", f"={y2}*2", "", fmt='0.0')
    lump = s.row(None, "TRỢ CẤP MỘT LẦN ƯỚC TÍNH", f"=({s1}+{s2})*{mb}", "", "total")
    s.section("D. NẾU BẢO LƯU ĐỂ HƯỞNG LƯƠNG HƯU (tối thiểu 15 năm)")
    s.row(None, "Đủ số năm đóng tối thiểu", f'=IF({ty}>=15,"ĐỦ 15 NĂM","CHƯA ĐỦ 15 NĂM")', "", fmt="@")
    rm = s.row(None, "Tỷ lệ hưởng (nam)", f"=IF({ty}<15,0,IF({ty}>=20,MIN(0.75,0.45+INT({ty}-20)*0.02),0.40+INT({ty}-15)*0.01))",
               "Điều 66 Luật 41/2024: 20 năm = 45%, +2%/năm; nam đóng 15–<20 năm: 40% + 1%/năm", fmt=PCT)
    rf = s.row(None, "Tỷ lệ hưởng (nữ)", f"=IF({ty}<15,0,MIN(0.75,0.45+INT({ty}-15)*0.02))", "15 năm = 45%, +2%/năm, tối đa 75%", fmt=PCT)
    s.row(None, "Lương hưu tháng ước tính (nam)", f"={mb}*{rm}")
    s.row(None, "Lương hưu tháng ước tính (nữ)", f"={mb}*{rf}")
    s.row(None, "Thời gian hưởng lương hưu để bằng trợ cấp 1 lần (nữ) (tháng)",
          f'=IF({mb}*{rf}>0,{lump}/({mb}*{rf}),"-")', "", fmt='0.0')
    return s


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--output", required=True, help="Đường dẫn .xlsx (trong ~/Downloads/AIWF_Output/...)")
    ap.add_argument("--year", type=int, default=2026, choices=sorted(tc.RULES), help="Năm thuế")
    ap.add_argument("--region", choices=["I", "II", "III", "IV"], help="Vùng lương tối thiểu → trần BHTN")
    ap.add_argument("--gross", type=float, help="Lương Gross tháng → sheet 'Lương tháng'")
    ap.add_argument("--dependents", type=int, help="Số người phụ thuộc (BẮT BUỘC khi có --gross hoặc quyết toán)")
    ap.add_argument("--medical", type=float, default=0.0, help="Chi phí y tế BÌNH QUÂN THÁNG (như tax_calculator)")
    ap.add_argument("--education", type=float, default=0.0, help="Chi phí giáo dục BÌNH QUÂN THÁNG")
    ap.add_argument("--pension", type=float, default=0.0, help="Hưu trí tự nguyện VNĐ/tháng")
    ap.add_argument("--settlement-income", type=float, help="Tổng thu nhập chịu thuế cả năm → sheet 'Quyết toán năm'")
    ap.add_argument("--settlement-insurance", type=float, help="Bảo hiểm bắt buộc đã đóng cả năm (bắt buộc khi quyết toán)")
    ap.add_argument("--tax-withheld", "--settlement-tax-withheld", dest="tax_withheld", type=float,
                    help="Thuế đã tạm khấu trừ cả năm (bắt buộc khi quyết toán)")
    ap.add_argument("--months", type=int, default=12, help="Số tháng tính giảm trừ bản thân khi quyết toán")
    ap.add_argument("--dependent-months", type=int, help="Tổng tháng-người phụ thuộc (mặc định = NPT x số tháng)")
    ap.add_argument("--include-bhxh", action="store_true", help="Thêm sheet BHXH một lần (cần đủ 4 tham số bên dưới)")
    ap.add_argument("--bhxh-before-2014", type=float)
    ap.add_argument("--bhxh-from-2014", type=float)
    ap.add_argument("--mbqtl", type=float)
    ap.add_argument("--months-off", type=int)
    a = ap.parse_args()

    errors = []
    if a.gross is None and a.settlement_income is None and not a.include_bhxh:
        errors.append("Không có dữ liệu nào để xuất: cần --gross và/hoặc --settlement-income và/hoặc --include-bhxh.")
    if (a.gross is not None or a.settlement_income is not None) and a.dependents is None:
        errors.append("Thiếu --dependents (hỏi người dùng; 0 nếu không có người phụ thuộc).")
    if a.settlement_income is not None:
        if a.settlement_insurance is None:
            errors.append("Quyết toán cần --settlement-insurance (bảo hiểm đã đóng cả năm, theo chứng từ).")
        if a.tax_withheld is None:
            errors.append("Quyết toán cần --tax-withheld (thuế đã khấu trừ theo chứng từ). Không được tự giả định.")
    if a.include_bhxh and None in (a.bhxh_before_2014, a.bhxh_from_2014, a.mbqtl, a.months_off):
        errors.append("--include-bhxh cần đủ --bhxh-before-2014 --bhxh-from-2014 --mbqtl --months-off (theo sổ BHXH/VssID).")
    if errors:
        for e in errors:
            print(f"✗ {e}", file=sys.stderr)
        sys.exit(1)
    if a.gross is not None and not a.region and a.gross > 20 * tc.REGION_MIN_WAGE[a.year]["IV"]:
        print("⚠️ Chưa có --region: BHTN tính trên toàn bộ lương (có thể thừa). Ghi rõ giả định trong báo cáo.", file=sys.stderr)

    wb = openpyxl.Workbook()
    P = build_params(wb, a.year, a.region)
    sheets = []
    if a.gross is not None:
        build_monthly(wb, P, a); sheets.append("Lương tháng")
    if a.settlement_income is not None:
        build_settlement(wb, P, a); sheets.append("Quyết toán năm")
    if a.include_bhxh:
        build_bhxh(wb, a); sheets.append("BHXH một lần")
    wb.active = 1 if len(wb.sheetnames) > 1 else 0

    out = Path(a.output).expanduser().resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"✅ Đã xuất {out}")
    print(f"   Năm thuế {a.year} · sheet: Thông số, {', '.join(sheets)}")
    print("   Kiểm tra: mở bằng LibreOffice/Excel hoặc chạy verify_tax_sheet.py để so với tax_calculator.py")


if __name__ == "__main__":
    main()
