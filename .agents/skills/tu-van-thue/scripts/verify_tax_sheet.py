#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
verify_tax_sheet.py — Cổng kiểm định file Excel do export_tax_sheet.py xuất ra.

Tính lại toàn bộ công thức bằng LibreOffice headless, rồi so từng kết quả (thuế tháng, Net, thuế năm,
chênh lệch quyết toán) với tax_calculator.py cùng đầu vào. Lệch > 1đ, ô lỗi (#VALUE!, #REF!...) hoặc ô trống → exit 1.

    python3 verify_tax_sheet.py --xlsx <file.xlsx>
"""
import argparse
import glob
import os
import shutil
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import tax_calculator as tc  # noqa: E402
import openpyxl  # noqa: E402

SOFFICE = [shutil.which("soffice"), shutil.which("libreoffice"), "/Applications/LibreOffice.app/Contents/MacOS/soffice"]


def recalc(path):
    exe = next((p for p in SOFFICE if p and os.path.exists(p)), None)
    if not exe:
        sys.exit("✗ Không tìm thấy LibreOffice (soffice) để tính lại công thức. Cài: brew install --cask libreoffice")
    out = tempfile.mkdtemp(prefix="taxverify_")
    subprocess.run([exe, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", out, path],
                   check=True, capture_output=True, timeout=180)
    return glob.glob(os.path.join(out, "*.xlsx"))[0]


def table(ws):
    return {str(ws.cell(row=r, column=1).value or "").strip(): ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    a = ap.parse_args()
    raw = openpyxl.load_workbook(a.xlsx)
    calc = openpyxl.load_workbook(recalc(a.xlsx), data_only=True)
    fails = []

    for ws in calc.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#"):
                    fails.append(f"{ws.title}!{c.coordinate} lỗi công thức {c.value}")
        for row in raw[ws.title].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and calc[ws.title][c.coordinate].value is None:
                    fails.append(f"{ws.title}!{c.coordinate} công thức không ra giá trị")

    title = str(raw["Thông số"]["A1"].value)
    year = int(title.rsplit(" ", 1)[-1])
    tc.set_tax_year(year)
    P = table(calc["Thông số"])
    bhtn_cap = next((v for k, v in P.items() if k.startswith("Trần lương đóng BHTN")), 0) or 0
    checked = []

    def cmp(label, got, want):
        checked.append(label)
        if got is None or abs(float(got) - float(want)) > 1:
            fails.append(f"{label}: Excel {got} ≠ tax_calculator {want:,.0f}")

    if "Lương tháng" in calc.sheetnames:
        T = table(calc["Lương tháng"])
        r = tc.calculate_monthly_salary_tax(T["Lương Gross tháng"], int(T["Số người phụ thuộc đã đăng ký"]),
                                            T["Hưu trí tự nguyện đã đóng (VNĐ/tháng)"], T["Chi phí y tế (bình quân VNĐ/tháng)"],
                                            T["Chi phí giáo dục (bình quân VNĐ/tháng)"], bhtn_cap)
        cmp("Tổng bảo hiểm tháng", T["Tổng bảo hiểm"], r["insurance"]["total"])
        cmp("Thu nhập tính thuế tháng", T["Thu nhập tính thuế"], r["taxable_income"])
        cmp("Thuế TNCN tháng", T["THUẾ TNCN THÁNG"], round(r["tax"]["total_tax"]))
        cmp("Lương Net", T["LƯƠNG NET (Gross - bảo hiểm - thuế)"], r["net_salary"])

    if "Quyết toán năm" in calc.sheetnames:
        T = table(calc["Quyết toán năm"])
        months, dep_months = int(T["Số tháng tính giảm trừ bản thân"]), T["Tổng số tháng-người phụ thuộc"]
        r = tc.calculate_annual_settlement(T["Tổng thu nhập chịu thuế cả năm (mọi nơi chi trả)"], T["Bảo hiểm bắt buộc đã đóng cả năm"],
                                           0, T["Chi phí y tế cả năm"], T["Chi phí giáo dục cả năm"], T["Hưu trí tự nguyện cả năm"],
                                           T["Thuế TNCN đã tạm khấu trừ/tạm nộp"], months)
        # tax_calculator tính NPT theo người x tháng làm việc; sheet cho nhập tháng-người → tính lại phần NPT
        ti = max(0.0, r["annual_gross_income"] - r["deductions"]["total_deductions"] - dep_months * tc.DEDUCTION_DEPENDENT)
        want = round(sum(max(0.0, min(ti, lim * 12) - lo * 12) * rt for lo, lim, rt in
                         zip([0] + tc.LIMITS_MONTHLY[:-1], tc.LIMITS_MONTHLY, tc.RATES)))
        cmp("Thuế TNCN cả năm", T["THUẾ TNCN PHẢI NỘP CẢ NĂM"], want)
        diff = T["Thuế TNCN đã tạm khấu trừ/tạm nộp"] - want
        cmp("Chênh lệch quyết toán", T["Chênh lệch = đã khấu trừ - phải nộp"], diff)
        verdict = str(T["KẾT LUẬN"])
        expect = "NỘP THỪA" if diff > 0 else "PHẢI NỘP THÊM" if diff < -50_000 else "miễn nộp" if diff < 0 else "CÂN BẰNG"
        if expect not in verdict:
            fails.append(f"Kết luận quyết toán '{verdict}' không khớp chênh lệch {diff:,.0f}")

    print(f"Kiểm định {a.xlsx} — năm thuế {year}")
    for c in checked:
        print(f"  ✓ so khớp {c}" if not any(f.startswith(c) for f in fails) else f"  ✗ {c}")
    for f in fails:
        print(f"  ✗ {f}")
    print("✅ ĐẠT" if not fails else f"✗ KHÔNG ĐẠT ({len(fails)} lỗi) — sửa trước khi giao")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
