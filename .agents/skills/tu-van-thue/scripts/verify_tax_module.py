#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""verify_tax_module.py: tính lại file do export_tax_module.py tạo (LibreOffice) và đối chiếu với engine Python.

    python3 verify_tax_module.py <file.xlsx>      # mã thoát 0 = khớp, 1 = lệch/lỗi công thức, 2 = lỗi chạy
"""
from __future__ import annotations

import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl  # noqa: E402
from verify_tax_sheet import recalc  # noqa: E402  (dùng lại hàm tính lại LibreOffice của skill)


def main(argv=None) -> int:
    argv = argv if argv is not None else sys.argv[1:]
    if len(argv) != 1:
        print(__doc__)
        return 2
    src = argv[0]
    meta = json.loads(openpyxl.load_workbook(src).properties.description or "{}")
    if "expected" not in meta:
        print("✗ File không có dữ liệu kỳ vọng (không phải xuất từ export_tax_module.py)")
        return 2
    wb = openpyxl.load_workbook(recalc(src), data_only=True)
    ws = wb["Tính"]
    bad = 0
    for k, want in meta["expected"].items():
        got = ws[meta["cells"][k]].value
        ok = got is not None and abs(float(got) - float(want)) < 1e-6
        print(f"{'✅' if ok else '❌'} {k}: Excel={got} engine={want}")
        bad += 0 if ok else 1
    errs = [c.coordinate for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("#")]
    if errs:
        print("❌ Ô lỗi công thức:", errs)
        bad += 1
    print("✅ ĐẠT" if not bad else f"✗ {bad} lỗi")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
