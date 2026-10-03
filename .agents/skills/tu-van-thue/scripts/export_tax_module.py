#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""export_tax_module.py: xuất bảng tính Excel công thức sống cho các module thuế mới.

Sheet "Thông số": tham số có hiệu lực tại --on, kèm văn bản căn cứ và trạng thái kiểm chứng.
Sheet "Tính": ô đầu vào (nền vàng) và công thức tham chiếu sang "Thông số". Sửa tham số là bảng tự cập nhật.
Giá trị kỳ vọng từ engine Python được ghi vào thuộc tính tệp để verify_tax_module.py đối chiếu sau khi tính lại.

    python3 export_tax_module.py --module tndn --output out.xlsx --revenue 5e9 --expenses 4e9 --prior-revenue 5e9
    python3 export_tax_module.py --module hkd --output out.xlsx --activity phan_phoi --revenue 2e9
    python3 export_tax_module.py --module gtgt --output out.xlsx --revenue 1e9 --input-vat 30e6 --reduction
    python3 export_tax_module.py --module ttdb --output out.xlsx --item bia --price 1e9
    python3 export_tax_module.py --module nha-thau --output out.xlsx --item dich_vu --revenue 1e9
    python3 export_tax_module.py --module cham-nop --output out.xlsx --amount 1e8 --due 2026-04-30 --paid 2026-05-05
"""
from __future__ import annotations

import argparse
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    print("Lỗi: thiếu openpyxl. Cài: pip3 install openpyxl", file=sys.stderr)
    sys.exit(2)

from tax_engines import fct, gtgt, hkd, penalties, tndn, ttdb  # noqa: E402
from rules_loader import ParamError, load, parse_date  # noqa: E402

INPUT_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
HEAD_FILL = PatternFill(start_color="2E5B88", end_color="2E5B88", fill_type="solid")
RESULT_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
BOLD = Font(bold=True)
WHITE = Font(bold=True, color="FFFFFF")
NUM = "#,##0"
PCT = "0.0##%"


class Book:
    def __init__(self, on: str):
        self.wb = openpyxl.Workbook()
        self.pa = self.wb.active
        self.pa.title = "Thông số"
        self.ca = self.wb.create_sheet("Tính")
        self.on = on
        self.prow, self.crow = 2, 2
        self.refs: dict = {}
        self.keys: dict = {}
        for i, h in enumerate(["Tham số", "Giá trị", "Văn bản căn cứ", "Trạng thái"], 1):
            c = self.pa.cell(1, i, h); c.font, c.fill = WHITE, HEAD_FILL
        for i, h in enumerate(["Chỉ tiêu", "Giá trị", "Ghi chú"], 1):
            c = self.ca.cell(1, i, h); c.font, c.fill = WHITE, HEAD_FILL
        self.pa.column_dimensions["A"].width = 40
        self.pa.column_dimensions["C"].width = 60
        self.ca.column_dimensions["A"].width = 52
        self.ca.column_dimensions["B"].width = 24
        self.ca.column_dimensions["C"].width = 70

    def param(self, module: str, name: str, label: str, pct: bool = False) -> str:
        p = load(module).get(name, self.on)
        r = self.prow
        self.pa.cell(r, 1, label)
        c = self.pa.cell(r, 2, p.value); c.number_format = PCT if pct else NUM
        self.pa.cell(r, 3, p.citation())
        self.pa.cell(r, 4, p.status)
        self.prow += 1
        return f"'Thông số'!$B${r}"

    def row(self, label: str, value, note: str = "", inp: bool = False, key: str = "", fmt: str = NUM, bold: bool = False) -> str:
        r = self.crow
        self.ca.cell(r, 1, label).font = BOLD if bold else Font()
        c = self.ca.cell(r, 2, value)
        c.number_format = fmt
        if inp:
            c.fill = INPUT_FILL
        if bold:
            c.font, c.fill = BOLD, RESULT_FILL
        self.ca.cell(r, 3, note).alignment = Alignment(wrap_text=True)
        if key:
            self.keys[key] = f"B{r}"
        self.crow += 1
        return f"B{r}"

    def finish(self, path: str, expected: dict, disclaimer: str):
        self.row("", "")
        self.ca.cell(self.crow, 1, disclaimer).alignment = Alignment(wrap_text=True)
        self.wb.properties.description = json.dumps({"expected": expected, "cells": self.keys, "on": self.on}, ensure_ascii=False)
        self.wb.save(path)


def build(a) -> Book:
    bk = Book(a.on)
    mod = a.module
    if mod == "tndn":
        ex_lim = bk.param("tndn", "nguong_mien_tndn", "Ngưỡng miễn thuế (doanh thu)")
        small = bk.param("tndn", "nguong_dn_nho", "Ngưỡng DN nhỏ")
        mid = bk.param("tndn", "nguong_dn_vua", "Ngưỡng DN vừa")
        r_s = bk.param("tndn", "thue_suat_dn_nho", "Thuế suất DN nhỏ", True)
        r_m = bk.param("tndn", "thue_suat_dn_vua", "Thuế suất DN vừa", True)
        r_l = bk.param("tndn", "thue_suat_chuan", "Thuế suất chuẩn", True)
        rev = bk.row("Doanh thu", a.revenue, inp=True)
        oth = bk.row("Thu nhập khác", a.other_income, inp=True)
        exp = bk.row("Chi phí đã ghi nhận", a.expenses, inp=True)
        nd = bk.row("Chi phí không được trừ", a.non_deductible, inp=True)
        exi = bk.row("Thu nhập miễn thuế", a.exempt, inp=True)
        loss = bk.row("Lỗ kết chuyển còn được trừ", a.loss, inp=True)
        basis = bk.row("Tổng doanh thu kỳ liền kề trước", a.prior_revenue if a.prior_revenue is not None else a.revenue + a.other_income,
                       "Dùng chọn thuế suất và xét miễn thuế", inp=True)
        tx = bk.row("Thu nhập chịu thuế trước lỗ", f"={rev}+{oth}-{exp}+{nd}-{exi}", key="thu_nhap_chiu_thue_truoc_lo")
        rate = bk.row("Thuế suất áp dụng", f"=IF({basis}<={ex_lim},0,IF({basis}<={small},{r_s},IF({basis}<={mid},{r_m},{r_l})))", fmt=PCT, key="thue_suat")
        bk.row("Thu nhập tính thuế", f"=MAX(0,{tx}-IF({tx}>0,MIN({tx},{loss}),0))", key="thu_nhap_tinh_thue")
        bk.row("Thuế TNDN phải nộp", f"=ROUND(B{bk.crow - 1}*{rate},0)", bold=True, key="thue_tndn")
        exp_v = tndn.compute(a.revenue, a.other_income, a.expenses, a.non_deductible, a.exempt, a.loss, a.prior_revenue, a.on).values
        expected = {k: exp_v[k] for k in ("thu_nhap_chiu_thue_truoc_lo", "thue_suat", "thue_tndn")}
        expected["thu_nhap_tinh_thue"] = exp_v.get("thu_nhap_tinh_thue", 0)
    elif mod == "hkd":
        ex_lim = bk.param("hkd", "nguong_mien_thue", "Ngưỡng miễn thuế (doanh thu năm)")
        t_rate = bk.param("hkd", f"tncn_dt_{a.activity}", f"Tỷ lệ TNCN theo doanh thu ({a.activity})", True)
        g_rate = bk.param("gtgt", f"ty_le_truc_tiep_{a.gtgt_group}", f"Tỷ lệ GTGT trực tiếp ({a.gtgt_group})", True)
        rev = bk.row("Doanh thu năm", a.revenue, inp=True)
        bk.row("Doanh thu chịu thuế (vượt ngưỡng)", f"=IF({rev}>{ex_lim},{rev},0)", "Không đủ ngưỡng thì không nộp (kê khai theo quy định)", key="doanh_thu_chiu_thue")
        base = f"B{bk.crow - 1}"
        bk.row("Thuế GTGT", f"=ROUND({base}*{g_rate},0)", key="gtgt")
        g = f"B{bk.crow - 1}"
        bk.row("Thuế TNCN (phương pháp doanh thu)", f"=ROUND(MAX(0,{base}-IF({base}>0,{ex_lim},0))*{t_rate},0)", "Phần doanh thu vượt ngưỡng nhân tỷ lệ", key="tncn")
        bk.row("Tổng thuế", f"={g}+B{bk.crow - 1}", bold=True, key="tong_thue")
        h = hkd.compute({a.activity: a.revenue}, None, a.on, False).values
        expected = {"gtgt": h["gtgt"], "tncn": h["tncn"], "tong_thue": h["tong_thue"]}
    elif mod == "gtgt":
        std = bk.param("gtgt", "thue_suat_chuan", "Thuế suất chuẩn", True)
        red = bk.param("gtgt", "thue_suat_giam", "Thuế suất giảm", True) if a.reduction else std
        rev = bk.row("Doanh thu chưa thuế (chịu thuế 10%)", a.revenue, inp=True)
        inv = bk.row("Thuế GTGT đầu vào được khấu trừ", a.input_vat, inp=True)
        carry = bk.row("Thuế đầu vào kỳ trước chuyển sang", a.carry, inp=True)
        bk.row("Thuế GTGT đầu ra", f"=ROUND({rev}*{red},0)", key="thue_dau_ra")
        out = f"B{bk.crow - 1}"
        bk.row("Thuế phải nộp", f"=MAX(0,{out}-{inv}-{carry})", bold=True, key="thue_phai_nop")
        e = gtgt.deduction([{"name": "A", "amount": a.revenue, "rate": "10", "reduction_eligible": a.reduction}], a.input_vat, a.carry, a.on).values
        expected = {"thue_dau_ra": e["thue_dau_ra"], "thue_phai_nop": e["thue_phai_nop"]}
    elif mod == "ttdb":
        rate = bk.param("ttdb", a.item, f"Thuế suất {a.item}", True)
        price = bk.row("Giá tính thuế (chưa TTĐB, BVMT, GTGT)", a.price, inp=True)
        bk.row("Thuế TTĐB", f"=ROUND({price}*{rate},0)", bold=True, key="thue_ttdb")
        expected = {"thue_ttdb": ttdb.compute(a.item, a.price, 0, a.on).values["thue_ttdb"]}
    elif mod == "nha-thau":
        rate = bk.param("fct", a.item, f"Tỷ lệ TNDN nhà thầu ({a.item})", True)
        rev = bk.row("Doanh thu tính thuế", a.revenue, inp=True)
        bk.row("Thuế TNDN nhà thầu", f"=ROUND({rev}*{rate},0)", bold=True, key="thue_tndn_nha_thau")
        expected = {"thue_tndn_nha_thau": fct.compute(a.item, a.revenue, a.on).values["thue_tndn_nha_thau"]}
    elif mod == "cham-nop":
        rate = bk.param("quan_ly_thue", "muc_cham_nop_ngay", "Mức chậm nộp mỗi ngày", True)
        amt = bk.row("Số thuế chậm nộp", a.amount, inp=True)
        due = bk.row("Hạn nộp (YYYY-MM-DD)", a.due, inp=True, fmt="@")
        paid = bk.row("Ngày thực nộp (YYYY-MM-DD)", a.paid, inp=True, fmt="@")
        bk.row("Số ngày chậm", f"=MAX(0,DATEVALUE({paid})-DATEVALUE({due})-1)", "Từ ngày liền sau hạn đến ngày liền trước ngày nộp", key="so_ngay_cham")
        days = f"B{bk.crow - 1}"
        bk.row("Tiền chậm nộp", f"=ROUND({amt}*{rate}*{days},0)", bold=True, key="tien_cham_nop")
        v = penalties.late_payment(a.amount, a.due, a.paid).values
        expected = {"so_ngay_cham": v["so_ngay_cham"], "tien_cham_nop": v["tien_cham_nop"]}
    else:
        raise ValueError(f"module không hợp lệ: {mod}")
    bk._expected = expected
    return bk


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--module", required=True, choices=["tndn", "hkd", "gtgt", "ttdb", "nha-thau", "cham-nop"])
    ap.add_argument("--output", required=True)
    ap.add_argument("--on", default="2026-12-31")
    ap.add_argument("--revenue", type=float, default=0)
    ap.add_argument("--other-income", type=float, default=0)
    ap.add_argument("--expenses", type=float, default=0)
    ap.add_argument("--non-deductible", type=float, default=0)
    ap.add_argument("--exempt", type=float, default=0)
    ap.add_argument("--loss", type=float, default=0)
    ap.add_argument("--prior-revenue", type=float)
    ap.add_argument("--activity", default="phan_phoi", help="hkd: phan_phoi|dich_vu|san_xuat|khac|noi_dung_so")
    ap.add_argument("--gtgt-group", default=None, help="hkd: nhóm tỷ lệ GTGT trực tiếp (phan_phoi|dich_vu|san_xuat|khac); mặc định theo activity")
    ap.add_argument("--input-vat", type=float, default=0)
    ap.add_argument("--carry", type=float, default=0)
    ap.add_argument("--reduction", action="store_true")
    ap.add_argument("--item")
    ap.add_argument("--price", type=float, default=0)
    ap.add_argument("--amount", type=float, default=0)
    ap.add_argument("--due")
    ap.add_argument("--paid")
    a = ap.parse_args(argv)
    if a.module == "hkd" and not a.gtgt_group:
        a.gtgt_group = a.activity if a.activity in ("phan_phoi", "dich_vu", "san_xuat") else "khac"
    if a.module in ("ttdb", "nha-thau") and not a.item:
        print("✗ Thiếu --item", file=sys.stderr)
        return 1
    if a.module == "cham-nop" and not (a.due and a.paid):
        print("✗ Thiếu --due/--paid", file=sys.stderr)
        return 1
    try:
        parse_date(a.on)
        bk = build(a)
    except (ParamError, ValueError) as ex:
        print(f"✗ {ex}", file=sys.stderr)
        return 1
    bk.finish(a.output, bk._expected,
              "Tư vấn sơ bộ, không phải kết luận của cơ quan thuế. Tham số có trạng thái CORROBORATED cần [CẦN XÁC MINH] trước khi dùng.")
    print(f"✅ Đã xuất {a.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
