#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
fetch_vn_source.py — Tải văn bản quy phạm pháp luật Việt Nam từ nguồn CHÍNH THỐNG và lưu thành
file text trong <research_dir>/sources/ để evidence_verifier.py đối chiếu nguyên văn.

Chỉ đọc (HTTP GET) trang công khai; không gửi dữ liệu. Không gọi LLM API.

Cách dùng (chạy từ thư mục workspace AIWF):
  # Trang văn bản trên Công báo (tự tìm file .docx/.pdf đính kèm):
  python3 .agents/skills/tu-van-phap-luat/scripts/fetch_vn_source.py \
      --url "https://congbao.chinhphu.vn/van-ban/luat-so-76-2025-qh15-45505.htm" \
      --out-dir "<research_dir>/sources" --name luat-76-2025-qh15

  # File người dùng cung cấp (khi nguồn web bị chặn 403):
  python3 .agents/skills/tu-van-phap-luat/scripts/fetch_vn_source.py \
      --file "~/Downloads/59-2020-QH14.pdf" --out-dir "<research_dir>/sources" --name luat-59-2020-qh14

Mã thoát:
  0 = đã lưu <out-dir>/<name>.txt
  3 = BỊ CHẶN (403/Cloudflare/captcha)  → áp dụng giao thức 403 trong SKILL.md
  4 = Trang chỉ là khung JavaScript / không có nội dung Điều → thử nguồn chính thống khác
  1 = lỗi khác
"""

import argparse
import hashlib
import html
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126 Safari/537.36"
BLOCK_MARKERS = ("Just a moment", "Attention Required", "cf-browser-verification", "Enable JavaScript and cookies",
                 "captcha", "Access denied")
OFFICIAL_HOSTS = ("congbao.chinhphu.vn", "congbaocdn.chinhphu.vn", "vanban.chinhphu.vn", "datafiles.chinhphu.vn",
                  "cdnchinhphu.vn", "vbpl.vn", "chinhphu.vn", ".gov.vn", "quochoi.vn")


def fetch(url: str, timeout: int = 60):
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept-Language": "vi,en;q=0.8"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, r.headers.get("Content-Type", ""), r.read(), r.geturl()
    except urllib.error.HTTPError as e:
        return e.code, e.headers.get("Content-Type", "") if e.headers else "", e.read() if e.fp else b"", url


def html_to_text(raw: str) -> str:
    raw = re.sub(r"(?is)<(script|style|noscript).*?</\1>", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>|</h\d>|</li>", "\n", raw)
    txt = html.unescape(re.sub(r"<[^>]+>", " ", raw))
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in txt.splitlines()]
    return "\n".join(ln for ln in lines if ln)


def file_to_text(path: Path) -> str:
    suffix = path.suffix.lower()
    head = path.read_bytes()[:8]
    if suffix == ".pdf" or head.startswith(b"%PDF"):
        if shutil.which("pdftotext"):
            out = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True)
            if out.returncode == 0 and out.stdout.strip():
                return out.stdout
        try:
            import fitz  # PyMuPDF
            with fitz.open(str(path)) as doc:
                return "\n".join(page.get_text() for page in doc)
        except Exception as e:  # noqa: BLE001
            raise RuntimeError(f"Không trích được chữ từ PDF ({e}). PDF scan cần OCR: dùng skill boc-tach-pdf.")
    if suffix == ".docx" or head.startswith(b"PK"):
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if suffix == ".doc":
        if shutil.which("textutil"):  # macOS
            out = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(path)], capture_output=True, text=True)
            if out.returncode == 0:
                return out.stdout
        for tool in ("antiword", "catdoc"):
            if shutil.which(tool):
                out = subprocess.run([tool, str(path)], capture_output=True, text=True)
                if out.returncode == 0:
                    return out.stdout
        raise RuntimeError("File .doc cần textutil (macOS) hoặc antiword/catdoc; hoặc tải bản .pdf/.docx.")
    return path.read_text(encoding="utf-8", errors="ignore")


def find_attachments(page_html: str, base: str):
    links = re.findall(r'href="([^"]+)"', page_html)
    out = []
    for raw in links:
        u = html.unescape(raw)
        low = urllib.parse.unquote(u).lower()
        for ext, rank in ((".docx", 0), (".pdf", 1), (".doc", 2)):
            if low.endswith(ext):
                out.append((rank, urllib.parse.urljoin(base, u), ext))
                break
    out.sort(key=lambda x: x[0])
    return out


def looks_blocked(status: int, body: bytes) -> bool:
    if status in (401, 403, 429, 503):
        return True
    head = body[:5000].decode("utf-8", errors="ignore")
    return any(m.lower() in head.lower() for m in BLOCK_MARKERS) and len(body) < 20000


def has_articles(text: str) -> bool:
    return len(re.findall(r"(?m)^\W*Điều\s+\d+\s*[\.:]", text)) >= 1


def save(out_dir: Path, name: str, text: str, meta: dict) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{name}.txt"
    header = "".join(f"# {k}: {v}\n" for k, v in meta.items())
    path.write_text(header + "\n" + text, encoding="utf-8")
    return path


def main() -> int:
    ap = argparse.ArgumentParser(description="Tải/chuyển VBQPPL chính thống thành text cho evidence_verifier")
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--url", help="Trang văn bản (congbao.chinhphu.vn, vanban.chinhphu.vn, vbpl.vn) hoặc link .pdf/.docx")
    src.add_argument("--file", help="File .pdf/.docx/.doc người dùng cung cấp")
    ap.add_argument("--out-dir", required=True, help="<research_dir>/sources")
    ap.add_argument("--name", required=True, help="Tên file, VD luat-59-2020-qh14")
    args = ap.parse_args()

    now = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
    out_dir = Path(args.out_dir).expanduser()

    if args.file:
        p = Path(args.file).expanduser()
        if not p.is_file():
            print(f"❌ Không thấy file: {p}", file=sys.stderr)
            return 1
        text = file_to_text(p)
        meta = {"SOURCE": f"user-file:{p.name}", "FETCHED_AT": now,
                "SHA256": hashlib.sha256(p.read_bytes()).hexdigest(), "OFFICIAL": "người dùng cung cấp - ghi rõ trong báo cáo"}
        path = save(out_dir, args.name, text, meta)
        print(f"✅ Đã lưu {path} ({len(text)} ký tự, {len(re.findall(r'Điều\s+\d+', text))} lần 'Điều N')")
        return 0 if has_articles(text) else 4

    url = args.url
    host = urllib.parse.urlparse(url).netloc
    if not any(h in host for h in OFFICIAL_HOSTS):
        print(f"⚠️  {host} không phải nguồn chính thống. Chỉ dùng để tìm số hiệu; trích dẫn phải lấy từ nguồn chính thống.",
              file=sys.stderr)

    status, ctype, body, final = fetch(url)
    if looks_blocked(status, body):
        print(f"⛔ BLOCKED ({status}) {url}\n   → Thử nguồn chính thống khác (congbao.chinhphu.vn, vanban.chinhphu.vn); "
              f"nếu vẫn không được: xin người dùng file PDF/DOCX (--file) hoặc gắn [CẦN XÁC MINH] cho mọi trích dẫn.",
              file=sys.stderr)
        return 3
    if status != 200:
        print(f"❌ HTTP {status} {url}", file=sys.stderr)
        return 1

    file_url = final
    with tempfile.TemporaryDirectory() as td:
        if "html" in ctype.lower() or body[:200].lstrip().lower().startswith((b"<!doctype", b"<html")):
            page = body.decode("utf-8", errors="ignore")
            text = html_to_text(page)
            if not has_articles(text):
                atts = find_attachments(page, final)
                for _, att_url, ext in atts:
                    st, _ct, data, fin = fetch(att_url)
                    if st == 200 and data and not looks_blocked(st, data):
                        tmp = Path(td) / f"doc{ext}"
                        tmp.write_bytes(data)
                        try:
                            text = file_to_text(tmp)
                        except RuntimeError as e:
                            print(f"⚠️  {e}", file=sys.stderr)
                            continue
                        file_url, body = att_url, data
                        if has_articles(text):
                            break
        else:
            ext = ".pdf" if body.startswith(b"%PDF") else (".docx" if body.startswith(b"PK") else ".bin")
            tmp = Path(td) / f"doc{ext}"
            tmp.write_bytes(body)
            text = file_to_text(tmp)

    if not has_articles(text):
        print(f"⛔ NO_CONTENT: {url} không có nội dung 'Điều N' (trang JavaScript hoặc không có file đính kèm). "
              f"Thử trang Công báo của cùng văn bản hoặc xin người dùng file.", file=sys.stderr)
        return 4

    meta = {"SOURCE_URL": url, "FILE_URL": file_url, "FETCHED_AT": now,
            "SHA256": hashlib.sha256(body).hexdigest()}
    path = save(out_dir, args.name, text, meta)
    print(f"✅ Đã lưu {path} ({len(text)} ký tự). Dùng làm 'source' trong claims.json.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
