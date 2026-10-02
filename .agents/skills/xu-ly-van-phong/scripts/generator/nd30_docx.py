#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
nd30_docx.py - Sinh văn bản hành chính DOCX đúng thể thức Nghị định 30/2020/NĐ-CP (Track 1, đen trắng).

Đầu vào: file JSON theo schema trong standards/nd30.md (mục "Cấu trúc JSON").
Hỗ trợ loai: cong_van, quyet_dinh, to_trinh, thong_bao, bao_cao, ke_hoach, giay_moi, bien_ban (và
"ten_loai" tùy ý cho loại khác).

Thể thức được dựng sẵn (Agent KHÔNG tự vẽ lại):
- A4, lề trên/dưới 20 mm, trái 30 mm, phải 20 mm; Times New Roman, chữ đen.
- Header 2 hàng x 2 cột (tên cơ quan | Quốc hiệu - Tiêu ngữ ; Số, ký hiệu (+ V/v) | Địa danh, ngày tháng),
  đường kẻ dưới tên cơ quan ban hành dài 1/3-1/2 dòng chữ, dưới Tiêu ngữ dài bằng dòng chữ
  (đo bằng font thật, không dùng chuỗi "____" và không co giãn ký tự).
- Tên loại + trích yếu + đường kẻ dưới trích yếu; căn cứ in nghiêng tự thêm ";" và "." dòng cuối.
- "Điều N." in đậm; nội dung căn đều hai lề, lùi đầu dòng 1,27 cm; kết thúc "./.".
- Khối Nơi nhận | Chữ ký trong 1 hàng bảng không cho tách trang, đoạn cuối nội dung giữ cùng trang.
- Số trang từ trang 2, giữa lề trên, không hiện ở trang 1.

Thiếu trường bắt buộc -> in lỗi và thoát mã 2 (không chèn placeholder ngầm).
"""

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
CONTENT_W_CM = 16.0          # 21 - 3 - 2
LEFT_COL_CM, RIGHT_COL_CM = 6.0, 10.0
SIG_LEFT_CM, SIG_RIGHT_CM = 7.0, 9.0
BODY_INDENT_CM = 1.27        # NĐ 30: lùi đầu dòng 1 cm hoặc 1,27 cm
PBDR_SUCCESSORS = ('w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')  # thứ tự con của w:pPr theo schema OOXML

TEN_LOAI = {
    "quyet_dinh": "QUYẾT ĐỊNH", "to_trinh": "TỜ TRÌNH", "thong_bao": "THÔNG BÁO",
    "bao_cao": "BÁO CÁO", "ke_hoach": "KẾ HOẠCH", "giay_moi": "GIẤY MỜI", "bien_ban": "BIÊN BẢN",
    "cong_van": None,
}
KINH_GUI_TYPES = {"cong_van", "to_trinh"}

_FONT_FILES = {
    (False, False): ["/System/Library/Fonts/Supplemental/Times New Roman.ttf", "C:/Windows/Fonts/times.ttf",
                     "/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman.ttf",
                     "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"],
    (True, False): ["/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf", "C:/Windows/Fonts/timesbd.ttf",
                    "/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman_Bold.ttf",
                    "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf"],
}


class Nd30Error(Exception):
    pass


def text_width_cm(text, size_pt, bold=False):
    """Đo độ rộng dòng chữ bằng font Times New Roman thật (hoặc Liberation Serif cùng metric)."""
    try:
        from PIL import ImageFont
        for f in _FONT_FILES[(bold, False)]:
            if Path(f).exists():
                font = ImageFont.truetype(f, 1000)
                w = font.getlength(text)
                return w / 1000 * size_pt / 72 * 2.54
    except Exception:
        pass
    factor = 0.62 if bold else 0.5  # ước lượng thận trọng khi không có font
    upper = sum(1 for c in text if c.isupper())
    em = (len(text) - upper) * factor + upper * (factor + 0.15)
    return em * size_pt / 72 * 2.54


# ----------------------------------------------------------------------------- helpers XML

def _set_run(run, size, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FONT)


def _para(container, text="", size=14, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
          first_indent=None, before=3, after=3, line_at_least=None, keep_next=False):
    p = container.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if line_at_least:
        pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        pf.line_spacing = Pt(line_at_least)
    else:
        pf.line_spacing = 1.0
    pf.first_line_indent = Cm(first_indent) if first_indent is not None else Cm(0)
    pf.keep_with_next = keep_next
    if text:
        _set_run(p.add_run(text), size, bold, italic)
    return p


def _rule(container, text, size, bold, col_w_cm, ratio):
    """Đoạn chỉ chứa đường kẻ ngang nét liền, dài = ratio x độ rộng dòng chữ, căn giữa cột."""
    length = min(col_w_cm - 0.2, text_width_cm(text, size, bold) * ratio)
    side = max(0.0, (col_w_cm - length) / 2)
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent, pf.right_indent = Cm(side), Cm(side)
    pf.space_before, pf.space_after = Pt(0), Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(5)
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    ppr.insert_element_before(bdr, *PBDR_SUCCESSORS)
    r = p.add_run("")
    _set_run(r, 2)
    return p


_USED_CELLS = set()


def _cell_para(cell, text, size, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0):
    """Đoạn đầu tiên của ô dùng lại paragraph rỗng sẵn có, các đoạn sau thêm mới."""
    key = id(cell._tc)
    if key not in _USED_CELLS:
        _USED_CELLS.add(key)
        p = cell.paragraphs[0]
    else:
        p = cell.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.0
    pf.first_line_indent = Cm(0)
    if text:
        _set_run(p.add_run(text), size, bold, italic)
    return p


def _no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "nil")
        borders.append(e)
    tbl_pr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.insert_element_before(layout, "w:tblCellMar", "w:tblLook", "w:tblCaption")
    mar = OxmlElement("w:tblCellMar")
    for edge in ("left", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), "0")
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tbl_pr.insert_element_before(mar, "w:tblLook", "w:tblCaption")


def _cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit")
    tr_pr.append(e)


def _set_widths(table, widths_cm):
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))


def _page_number_header(section):
    section.different_first_page_header_footer = True
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    _set_run(run, 13)
    for tag, attr in (("w:fldChar", {"w:fldCharType": "begin"}), ("w:instrText", None),
                      ("w:fldChar", {"w:fldCharType": "end"})):
        el = OxmlElement(tag)
        if attr:
            for k, v in attr.items():
                el.set(qn(k), v)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        run._r.append(el)


# ----------------------------------------------------------------------------- nội dung

def format_date(dia_danh, ngay):
    """'An Giang, ngày 05 tháng 01 năm 2026' - ngày < 10 và tháng 1, 2 thêm số 0 (NĐ 30)."""
    if isinstance(ngay, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", ngay):
        d = dt.date.fromisoformat(ngay)
        day = f"{d.day:02d}"
        month = f"{d.month:02d}" if d.month < 3 else str(d.month)
        return f"{dia_danh}, ngày {day} tháng {month} năm {d.year}"
    if isinstance(ngay, dict):  # văn bản chưa ký: để trống ngày
        m = ngay.get("thang")
        month = ("     " if m is None else (f"{m:02d}" if m < 3 else str(m)))
        return f"{dia_danh}, ngày      tháng {month} năm {ngay['nam']}"
    raise Nd30Error("'ngay' phải là 'YYYY-MM-DD' hoặc {\"thang\": 5, \"nam\": 2026} (ngày để trống chờ ký)")


def _end_punct(items, mid=";", last="."):
    out = []
    for i, t in enumerate(items):
        t = t.strip().rstrip(";.,")
        out.append(t + (last if i == len(items) - 1 else mid))
    return out


def validate(d):
    loai = d.get("loai")
    if not loai and d.get("ten_van_ban"):
        rev = {v: k for k, v in TEN_LOAI.items() if v}
        loai = rev.get(d["ten_van_ban"].upper(), "khac")
        d["loai"] = loai
    if not loai:
        raise Nd30Error("Thiếu 'loai' (cong_van, quyet_dinh, to_trinh, thong_bao, bao_cao, ke_hoach, ...)")
    if loai not in TEN_LOAI and not d.get("ten_loai"):
        raise Nd30Error(f"loai '{loai}' chưa có sẵn, khai báo thêm 'ten_loai' (vd 'CHƯƠNG TRÌNH')")
    req = ["co_quan_ban_hanh", "so_ky_hieu", "dia_danh", "ngay", "trich_yeu", "noi_dung", "noi_nhan",
           "chuc_vu_nguoi_ky", "ten_nguoi_ky"]
    if loai in KINH_GUI_TYPES:
        req.append("kinh_gui")
    if loai == "quyet_dinh":
        req += ["can_cu", "chu_the_ban_hanh"]
    miss = [k for k in req if not d.get(k)]
    if miss:
        raise Nd30Error("Thiếu trường bắt buộc: " + ", ".join(miss) + ". Hỏi người dùng, không tự bịa.")
    sk = d["so_ky_hieu"]
    if not re.fullmatch(r"[\d\s.]*/[A-ZĐ0-9][A-ZĐ0-9\-]*", sk):
        raise Nd30Error(f"so_ky_hieu '{sk}' sai dạng. Dạng đúng: '125/STC-NS' hoặc '   /QĐ-UBND' (chưa cấp số). "
                        "Văn bản hành chính KHÔNG ghi năm trong ký hiệu.")
    if loai == "quyet_dinh" and "/QĐ-" not in sk:
        raise Nd30Error("Quyết định phải có ký hiệu dạng '<số>/QĐ-<viết tắt cơ quan>'.")
    if loai == "cong_van" and re.search(r"/(QĐ|TTr|TB|BC|KH)-", sk):
        raise Nd30Error("Công văn không ghi tên loại trong ký hiệu: dạng '<số>/<viết tắt CQ>-<viết tắt đơn vị soạn>'.")
    return d


def build(d, out_path, size=14):
    validate(d)
    loai = d["loai"]
    line = round(size * 1.3)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)
    sec.header_distance, sec.footer_distance = Cm(1), Cm(1)
    _page_number_header(sec)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    doc.paragraphs[0]._element.getparent().remove(doc.paragraphs[0]._element) if doc.paragraphs else None

    # ---- Header 2 hàng x 2 cột
    t = doc.add_table(rows=2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _no_borders(t)
    _set_widths(t, [LEFT_COL_CM, RIGHT_COL_CM])
    for r in t.rows:
        _cant_split(r)
    L, R = t.cell(0, 0), t.cell(0, 1)
    if d.get("co_quan_chu_quan"):
        for ln in d["co_quan_chu_quan"].upper().split("\n"):
            _cell_para(L, ln, 13)
    ban_hanh = d["co_quan_ban_hanh"].upper().split("\n")
    for ln in ban_hanh:
        _cell_para(L, ln, 13, bold=True)
    _rule(L, max(ban_hanh, key=len), 13, True, LEFT_COL_CM, 0.45)

    qh = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
    qh_size = 13 if text_width_cm(qh, 13, True) <= RIGHT_COL_CM - 0.3 else 12
    tn = "Độc lập - Tự do - Hạnh phúc"
    _cell_para(R, qh, qh_size, bold=True)
    _cell_para(R, tn, 14, bold=True)
    _rule(R, tn, 14, True, RIGHT_COL_CM, 1.0)

    L2, R2 = t.cell(1, 0), t.cell(1, 1)
    _cell_para(L2, f"Số: {d['so_ky_hieu'].strip() if d['so_ky_hieu'].strip()[0] != '/' else '      ' + d['so_ky_hieu'].strip()}",
               13, before=6)
    if loai == "cong_van":
        ty = d["trich_yeu"].strip()
        ty = ty if ty.startswith("V/v") else "V/v " + ty[0].lower() + ty[1:]
        _cell_para(L2, ty, 12, before=3)
    _cell_para(R2, format_date(d["dia_danh"], d["ngay"]), 14, italic=True, before=6)

    # ---- Tên loại + trích yếu
    ten_loai = d.get("ten_loai") or TEN_LOAI.get(loai)
    if ten_loai:
        _para(doc, ten_loai.upper(), size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=18, after=0,
              keep_next=True)
        ty_lines = d["trich_yeu"].strip().split("\n")
        for i, ln in enumerate(ty_lines):
            _para(doc, ln, size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0, keep_next=True)
        r = _rule(doc, max(ty_lines, key=len), size, True, CONTENT_W_CM, 0.45)
        r.paragraph_format.keep_with_next = True
    else:
        _para(doc, "", size, before=6, after=0)

    # ---- Kính gửi
    if d.get("kinh_gui"):
        kg = d["kinh_gui"] if isinstance(d["kinh_gui"], list) else [d["kinh_gui"]]
        if len(kg) == 1:
            p = _para(doc, "", size, align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=12)
            _set_run(p.add_run("Kính gửi: "), size)
            _set_run(p.add_run(kg[0].strip().rstrip(".;")), size)
        else:
            p = _para(doc, "", size, align=WD_ALIGN_PARAGRAPH.LEFT, before=12, after=0)
            p.paragraph_format.left_indent = Cm(3.5)
            _set_run(p.add_run("Kính gửi:"), size)
            for i, item in enumerate(_end_punct(kg)):
                q = _para(doc, f"- {item}", size, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=0)
                q.paragraph_format.left_indent = Cm(5.0)
            doc.paragraphs[-1].paragraph_format.space_after = Pt(12)

    # ---- Thẩm quyền ban hành + căn cứ (quyết định)
    if loai == "quyet_dinh":
        _para(doc, d["chu_the_ban_hanh"].upper(), size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
              before=12, after=6)
        for c in _end_punct(d["can_cu"]):
            _para(doc, c, size, italic=True, first_indent=BODY_INDENT_CM, line_at_least=line)
        _para(doc, "QUYẾT ĐỊNH:", size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=6)
    elif d.get("can_cu"):
        for c in _end_punct(d["can_cu"]):
            _para(doc, c, size, italic=True, first_indent=BODY_INDENT_CM, line_at_least=line)

    # ---- Nội dung
    blocks = list(d["noi_dung"])
    body_paras = []
    for b in blocks:
        if isinstance(b, str):
            body_paras.append(_para(doc, b, size, first_indent=BODY_INDENT_CM, line_at_least=line))
        elif "dieu" in b:
            p = _para(doc, "", size, first_indent=BODY_INDENT_CM, line_at_least=line)
            head = b["dieu"].strip()
            if not head.endswith("."):
                head += "."
            _set_run(p.add_run(head + " "), size, bold=True)
            if b.get("tieu_de"):
                _set_run(p.add_run(b["tieu_de"].strip() + " "), size, bold=True)
            if b.get("text"):
                _set_run(p.add_run(b["text"]), size)
            body_paras.append(p)
            for sub in b.get("items", []):
                body_paras.append(_para(doc, sub, size, first_indent=BODY_INDENT_CM, line_at_least=line))
        elif "muc" in b:
            body_paras.append(_para(doc, b["muc"], size, bold=True, first_indent=BODY_INDENT_CM,
                                    line_at_least=line, before=6, keep_next=True))
        else:
            raise Nd30Error(f"Khối nội dung không hợp lệ: {b!r}. Dùng chuỗi, {{'muc': ...}} hoặc {{'dieu': ..., 'text': ...}}")
    last = body_paras[-1]
    txt = "".join(r.text for r in last.runs).rstrip()
    if not txt.endswith("./."):
        last.runs[-1].text = last.runs[-1].text.rstrip().rstrip(".") + "./."
    last.paragraph_format.keep_with_next = True

    # ---- Nơi nhận | Chữ ký (1 hàng không tách trang)
    sig = doc.add_table(rows=1, cols=2)
    _no_borders(sig)
    _set_widths(sig, [SIG_LEFT_CM, SIG_RIGHT_CM])
    _cant_split(sig.rows[0])
    NL, NR = sig.cell(0, 0), sig.cell(0, 1)
    _cell_para(NL, "Nơi nhận:", 12, bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=12)
    for item in _end_punct(d["noi_nhan"]):
        _cell_para(NL, f"- {item}", 11, align=WD_ALIGN_PARAGRAPH.LEFT)
    if d.get("quyen_han"):
        _cell_para(NR, d["quyen_han"].upper(), size, bold=True, before=12)
        _cell_para(NR, d["chuc_vu_nguoi_ky"].upper(), size, bold=True)
    else:
        _cell_para(NR, d["chuc_vu_nguoi_ky"].upper(), size, bold=True, before=12)
    for _ in range(4):
        _cell_para(NR, "", size)
    _cell_para(NR, d["ten_nguoi_ky"], size, bold=True)

    out = Path(out_path).expanduser().resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def main():
    ap = argparse.ArgumentParser(description="Sinh văn bản hành chính DOCX chuẩn NĐ 30/2020/NĐ-CP từ JSON")
    ap.add_argument("--input", "-i", required=True, help="JSON theo standards/nd30.md")
    ap.add_argument("--output", "-o", required=True, help="Đường dẫn .docx trong <output_dir>")
    ap.add_argument("--size", type=int, default=14, choices=[13, 14], help="Cỡ chữ nội dung (13 hoặc 14)")
    args = ap.parse_args()
    p = Path(args.input).expanduser()
    if not p.is_file():
        print(f"❌ Không tìm thấy {p}", file=sys.stderr)
        sys.exit(2)
    try:
        out = build(json.loads(p.read_text(encoding="utf-8")), args.output, args.size)
    except (Nd30Error, json.JSONDecodeError) as e:
        print(f"❌ LỖI DỮ LIỆU: {e}", file=sys.stderr)
        sys.exit(2)
    print(f"✅ Đã tạo {out}")
    print("   Bước tiếp theo BẮT BUỘC: python3 .agents/skills/xu-ly-van-phong/scripts/qa/check_nd30.py "
          f"\"{out}\" --render-dir <process_dir>/review")


if __name__ == "__main__":
    main()
