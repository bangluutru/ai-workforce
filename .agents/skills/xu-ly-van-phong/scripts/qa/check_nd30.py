#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
check_nd30.py - Kiểm tra thể thức văn bản hành chính DOCX theo Nghị định 30/2020/NĐ-CP (Phụ lục I).

Dùng cho MỌI văn bản Track 1 (do nd30_docx.py sinh hay do người dùng gửi lên để soát).
Lỗi (ERROR, thoát mã 1):
  khổ giấy / lề, font không phải Times New Roman, chữ màu, Quốc hiệu / Tiêu ngữ sai chữ, sai dấu gạch,
  sai đậm / cỡ, "Số:" sai dạng, địa danh ngày tháng sai (thiếu số 0, không nghiêng), công văn có tên loại
  "CÔNG VĂN", căn cứ không nghiêng / sai dấu câu cuối, "Điều N." không đậm, thiếu "./.", Nơi nhận sai cỡ
  chữ / kiểu chữ, em dash, khối Nơi nhận - chữ ký bị tách trang (khi render).
Cảnh báo (WARN, không chặn): tên đơn vị hành chính / cơ quan đã thay đổi từ năm 2025, năm trong ký hiệu
  văn bản hành chính, căn cứ pháp lý có dấu hiệu hết hiệu lực. Agent phải xác minh hoặc sửa.

  python3 check_nd30.py file.docx [--render-dir DIR]
"""

import argparse
import re
import subprocess
import shutil
import sys
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "extractor" / "office"))

QH = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
TN = "Độc lập - Tự do - Hạnh phúc"
EMU_PER_MM = 36000

# Cảnh báo bối cảnh hành chính sau sắp xếp năm 2025 (xác minh lại với văn bản mới nhất).
STALE_TERMS = [
    (r"\b[Hh]uyện\b", "Từ 01/7/2025 chính quyền địa phương 2 cấp (tỉnh, xã), không còn cấp huyện. Kiểm tra lại địa danh / cơ quan."),
    (r"Sở Kế hoạch và Đầu tư", "Sở Kế hoạch và Đầu tư đã hợp nhất vào Sở Tài chính (2025)."),
    (r"Sở Tài nguyên và Môi trường|Sở Nông nghiệp và Phát triển nông thôn",
     "Hai sở này đã hợp nhất thành Sở Nông nghiệp và Môi trường (2025)."),
    (r"Sở Thông tin và Truyền thông", "Sở Thông tin và Truyền thông đã sắp xếp lại (2025), kiểm tra tên cơ quan hiện hành."),
    (r"Sở Lao động\s*-\s*Thương binh và Xã hội", "Sở LĐ-TB&XH đã sắp xếp lại (2025), kiểm tra tên cơ quan hiện hành."),
    (r"Tổng cục Thuế", "Tổng cục Thuế đã đổi thành Cục Thuế (2025)."),
    (r"Luật Tổ chức chính quyền địa phương ngày 19 tháng 6 năm 2015",
     "Luật Tổ chức chính quyền địa phương 2015 đã được thay thế năm 2025, kiểm tra căn cứ hiện hành."),
]


class Report:
    def __init__(self):
        self.errors, self.warns = [], []

    def err(self, m):
        self.errors.append(m)

    def warn(self, m):
        self.warns.append(m)


def run_font(run, doc):
    if run.font.name:
        return run.font.name
    rpr = run._element.rPr
    if rpr is not None and rpr.rFonts is not None and rpr.rFonts.get(qn("w:ascii")):
        return rpr.rFonts.get(qn("w:ascii"))
    st = run._parent.style if hasattr(run._parent, "style") else None
    while st is not None:
        if st.font.name:
            return st.font.name
        st = st.base_style
    return doc.styles["Normal"].font.name


def run_size(run, para, doc):
    if run.font.size:
        return run.font.size.pt
    st = para.style
    while st is not None:
        if st.font.size:
            return st.font.size.pt
        st = st.base_style
    return doc.styles["Normal"].font.size.pt if doc.styles["Normal"].font.size else None


def all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p, None
    for t in doc.tables:
        for row in t.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                for p in cell.paragraphs:
                    yield p, cell


def text_runs(p):
    return [r for r in p.runs if r.text.strip()]


def check(path, rep):
    doc = Document(str(path))
    sec = doc.sections[0]
    w, h = sec.page_width / EMU_PER_MM, sec.page_height / EMU_PER_MM
    if abs(w - 210) > 1 or abs(h - 297) > 1:
        rep.err(f"Khổ giấy {w:.0f}x{h:.0f} mm, phải là A4 210x297 mm")
    for name, val, lo, hi in (("trên", sec.top_margin, 20, 25), ("dưới", sec.bottom_margin, 20, 25),
                              ("trái", sec.left_margin, 30, 35), ("phải", sec.right_margin, 15, 20)):
        mm = val / EMU_PER_MM
        if not lo - 0.5 <= mm <= hi + 0.5:
            rep.err(f"Lề {name} {mm:.1f} mm, NĐ 30 yêu cầu {lo}-{hi} mm")

    full_text = []
    for p, cell in all_paragraphs(doc):
        t = p.text
        full_text.append(t)
        for r in text_runs(p):
            f = run_font(r, doc)
            if f and f != "Times New Roman":
                rep.err(f"Font '{f}' tại đoạn '{t[:40]}', phải là Times New Roman")
                break
            c = r.font.color
            if c is not None and c.type is not None and c.rgb is not None and str(c.rgb) not in ("000000",):
                rep.err(f"Chữ màu {c.rgb} tại '{t[:40]}', văn bản hành chính phải màu đen")
                break
        if "—" in t:
            rep.err(f"Có em dash '—' tại '{t[:50]}'")
    joined = "\n".join(full_text)

    paras = list(all_paragraphs(doc))

    def find(pred):
        return [(p, c) for p, c in paras if pred(p.text.strip())]

    # Quốc hiệu
    qh = find(lambda s: s == QH)
    if not qh:
        rep.err(f"Thiếu Quốc hiệu đúng chữ '{QH}'")
    else:
        p = qh[0][0]
        rr = text_runs(p)
        if not all(r.bold for r in rr):
            rep.err("Quốc hiệu phải in đậm")
        sz = run_size(rr[0], p, doc)
        if sz and not 12 <= sz <= 13:
            rep.err(f"Quốc hiệu cỡ {sz}, NĐ 30: 12-13")
    # Tiêu ngữ
    tn_like = find(lambda s: re.sub(r"\s*[-–—]\s*", " - ", s) == TN)
    if not tn_like:
        rep.err(f"Thiếu Tiêu ngữ '{TN}'")
    else:
        p = tn_like[0][0]
        if p.text.strip() != TN:
            rep.err(f"Tiêu ngữ '{p.text.strip()}' sai dấu: giữa các cụm từ dùng gạch nối '-' có cách chữ")
        rr = text_runs(p)
        if not all(r.bold for r in rr):
            rep.err("Tiêu ngữ phải in đậm")
        sz = run_size(rr[0], p, doc)
        if sz and not 13 <= sz <= 14:
            rep.err(f"Tiêu ngữ cỡ {sz}, NĐ 30: 13-14")

    # Số, ký hiệu
    so = find(lambda s: s.startswith("Số:") or s.startswith("Số :"))
    if not so:
        rep.err("Thiếu dòng 'Số: .../...'")
    else:
        s = so[0][0].text.strip()
        if not re.fullmatch(r"Số:\s*[\d.\s]*/[A-ZĐ0-9][A-ZĐ0-9\-]*", s):
            rep.err(f"'{s}' sai dạng số, ký hiệu (vd 'Số: 125/QĐ-UBND', không cách chữ trong ký hiệu)")
        m = re.search(r"/(\d{4})/", s)
        if m:
            rep.warn(f"Ký hiệu '{s}' có năm: chỉ văn bản quy phạm pháp luật mới ghi năm; văn bản hành chính thì bỏ")
        if re.fullmatch(r"Số:\s*0?[1-9]/.*", s) and not re.match(r"Số:\s*0[1-9]/", s):
            rep.err(f"'{s}': số nhỏ hơn 10 phải ghi thêm số 0 phía trước (01, 02...)")

    # Địa danh, ngày tháng
    dates = find(lambda s: re.search(r",\s*ngày\s.*tháng\s.*năm\s+\d{4}$", s) is not None)
    if not dates:
        rep.err("Thiếu dòng 'Địa danh, ngày ... tháng ... năm ...'")
    else:
        p = dates[0][0]
        s = p.text.strip()
        if not all(r.italic for r in text_runs(p)):
            rep.err("Dòng địa danh, ngày tháng phải in nghiêng")
        m = re.search(r"ngày\s+(\S*)\s+tháng\s+(\S*)\s+năm", s)
        if m:
            day, mon = m.group(1), m.group(2)
            if re.fullmatch(r"[1-9]", day):
                rep.err(f"'{s}': ngày nhỏ hơn 10 phải ghi 0{day}")
            if mon in ("1", "2"):
                rep.err(f"'{s}': tháng 1, 2 phải ghi 0{mon}")
            if re.fullmatch(r"0[3-9]", mon):
                rep.err(f"'{s}': chỉ tháng 1, 2 thêm số 0; ghi 'tháng {mon[1]}'")

    is_cong_van = bool(find(lambda s: s.startswith("V/v")))
    if is_cong_van and find(lambda s: s == "CÔNG VĂN"):
        rep.err("Công văn KHÔNG có tên loại 'CÔNG VĂN' (trích yếu 'V/v ...' đặt dưới số, ký hiệu)")

    # Căn cứ (trước 'QUYẾT ĐỊNH:' hoặc trước nội dung)
    body = [p for p in doc.paragraphs if p.text.strip()]
    qd_idx = next((i for i, p in enumerate(body) if p.text.strip() in ("QUYẾT ĐỊNH:", "QUYẾT NGHỊ:")), None)
    if qd_idx is not None:
        cc = []
        i = qd_idx - 1
        while i >= 0 and re.match(r"(Căn cứ|Theo đề nghị|Xét đề nghị|Theo |Xét )", body[i].text.strip()):
            cc.insert(0, body[i])
            i -= 1
        if not cc:
            rep.err("Quyết định thiếu phần căn cứ ban hành trước 'QUYẾT ĐỊNH:'")
        for k, p in enumerate(cc):
            s = p.text.strip()
            if not all(r.italic for r in text_runs(p)):
                rep.err(f"Căn cứ phải in nghiêng: '{s[:50]}'")
            want = "." if k == len(cc) - 1 else ";"
            if not s.endswith(want):
                rep.err(f"Căn cứ '{s[-40:]}' phải kết thúc bằng '{want}'")
    for p in body:
        s = p.text.strip()
        if re.match(r"Điều \d+\.", s):
            rr = text_runs(p)
            if not rr or not rr[0].bold:
                rep.err(f"'{s[:12]}' phải in đậm")
        if len(s) > 80 and p.alignment not in (WD_ALIGN_PARAGRAPH.JUSTIFY, None) and \
                not (p.style and p.style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY):
            rep.err(f"Đoạn nội dung phải căn đều hai lề: '{s[:40]}'")
            break
    if "./." not in joined:
        rep.err("Thiếu dấu kết thúc nội dung './.'")

    # Nơi nhận
    nn = find(lambda s: s.startswith("Nơi nhận"))
    if not nn:
        rep.err("Thiếu 'Nơi nhận:'")
    else:
        p, cell = nn[0]
        rr = text_runs(p)
        if not (all(r.bold for r in rr) and all(r.italic for r in rr)):
            rep.err("'Nơi nhận:' phải in nghiêng, đậm")
        sz = run_size(rr[0], p, doc)
        if sz and sz != 12:
            rep.err(f"'Nơi nhận:' cỡ {sz}, NĐ 30: 12")
        if cell is not None:
            items = [q for q in cell.paragraphs if q.text.strip() and not q.text.strip().startswith("Nơi nhận")]
            for k, q in enumerate(items):
                s = q.text.strip()
                rq = text_runs(q)
                if rq and any(r.italic for r in rq):
                    rep.err(f"Danh sách nơi nhận phải chữ đứng (không nghiêng): '{s[:40]}'")
                    break
                szq = run_size(rq[0], q, doc) if rq else None
                if szq and szq != 11:
                    rep.err(f"Danh sách nơi nhận cỡ {szq}, NĐ 30: 11")
                    break
            if items:
                if not items[-1].text.strip().endswith("."):
                    rep.err("Dòng nơi nhận cuối (Lưu: VT, ...) phải kết thúc bằng '.'")
                for q in items[:-1]:
                    if not q.text.strip().endswith(";"):
                        rep.err(f"Dòng nơi nhận '{q.text.strip()[:30]}' phải kết thúc bằng ';'")
                        break
            tr = cell._tc.getparent()
            trpr = tr.find(qn("w:trPr"))
            if trpr is None or trpr.find(qn("w:cantSplit")) is None:
                rep.warn("Hàng chứa Nơi nhận / chữ ký chưa đặt 'không tách trang' (cantSplit); kiểm tra bản render")

    for pat, msg in STALE_TERMS:
        m = re.search(pat, joined)
        if m:
            rep.warn(f"'{m.group(0)}': {msg}")
    return doc, nn


def render(path, outdir, rep, nn):
    from soffice import find_soffice, run_soffice  # noqa: E402
    if not find_soffice():
        rep.err("Không tìm thấy LibreOffice để render (đặt biến SOFFICE hoặc cài LibreOffice)")
        return []
    outdir.mkdir(parents=True, exist_ok=True)
    run_soffice(["--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)], capture_output=True)
    pdf = outdir / (path.stem + ".pdf")
    if not pdf.exists():
        rep.err("Render PDF thất bại")
        return []
    pngs = []
    try:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
        d = fitz.open(str(pdf))
        pages = [unicodedata.normalize("NFC", pg.get_text()) for pg in d]
        if nn:
            cell = nn[0][1]
            signer = None
            if cell is not None:
                row = cell._tc.getparent()
                texts = [t for t in ("".join(w.text or "" for w in x.iter(qn("w:t"))).strip() for x in row.iter(qn("w:p"))) if t]
                signer = unicodedata.normalize("NFC", texts[-1]) if texts else None
            if signer:
                pg_nn = [i for i, t in enumerate(pages) if "Nơi nhận" in t]
                pg_sg = [i for i, t in enumerate(pages) if signer in t]
                if pg_nn and pg_sg and pg_nn[-1] != pg_sg[-1]:
                    rep.err(f"Khối Nơi nhận (trang {pg_nn[-1] + 1}) và họ tên người ký (trang {pg_sg[-1] + 1}) bị tách trang")
        if shutil.which("pdftoppm"):
            subprocess.run(["pdftoppm", "-png", "-r", "80", str(pdf), str(outdir / path.stem)], check=True)
            pngs = sorted(outdir.glob(path.stem + "-*.png"))
        else:
            for i, pg in enumerate(d, 1):
                p = outdir / f"{path.stem}-{i}.png"
                pg.get_pixmap(dpi=80).save(str(p))
                pngs.append(p)
    except ImportError:
        if shutil.which("pdftoppm"):
            subprocess.run(["pdftoppm", "-png", "-r", "80", str(pdf), str(outdir / path.stem)], check=True)
            pngs = sorted(outdir.glob(path.stem + "-*.png"))
        rep.warn("Thiếu PyMuPDF nên không kiểm tra được tách trang khối chữ ký")
    return pngs


def main():
    ap = argparse.ArgumentParser(description="Kiểm tra thể thức NĐ 30/2020/NĐ-CP cho file DOCX")
    ap.add_argument("docx")
    ap.add_argument("--render-dir", help="Render PDF + PNG để soát bằng mắt và kiểm tra tách trang")
    args = ap.parse_args()
    path = Path(args.docx).expanduser()
    if not path.is_file():
        print(f"❌ Không tìm thấy {path}", file=sys.stderr)
        sys.exit(2)
    rep = Report()
    _, nn = check(path, rep)
    pngs = render(path, Path(args.render_dir).expanduser(), rep, nn) if args.render_dir else []
    for w in rep.warns:
        print(f"⚠️  WARN: {w}")
    for p in pngs:
        print(f"🖼  {p}")
    if pngs:
        print("   BẮT BUỘC mở từng ảnh: soát header 2 cột thẳng hàng, đường kẻ dưới tên cơ quan / tiêu ngữ / trích yếu, "
              "chữ ký không rơi sang trang riêng.")
    if rep.errors:
        print(f"❌ NĐ30 CHECK FAIL ({len(rep.errors)} lỗi):")
        for e in rep.errors:
            print(f"   - {e}")
        sys.exit(1)
    print(f"✅ NĐ30 CHECK PASS ({len(rep.warns)} cảnh báo cần xem)")


if __name__ == "__main__":
    main()
