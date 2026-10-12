#!/usr/bin/env python3
"""
tests/governance/test_runbook_e2e_pilot.py — Simulated runbook lifecycle integration test.

CLASSIFICATION: SIMULATED INTEGRATION TEST (NOT LIVE MULTI-AGENT EXECUTION).
This test suite validates runbook phase transitions, evidence contracts, and approval gates
in a deterministic, reproducible environment.
Real multi-agent execution status: REAL E2E: NOT VERIFIED (requires live multi-agent daemon).
Real CLI & engine integration status: TESTED & EXECUTED.
"""

import json
import subprocess
import sys
from pathlib import Path
import docx
import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
SHARED_DIR = REPO_ROOT / ".agents" / "skills" / "_shared"
sys.path.insert(0, str(SHARED_DIR))
import bootstrap  # noqa: F401,E402

from evidence_gate import (
    create_evidence_pack,
    add_step_evidence,
    validate_evidence_pack,
    record_human_approval,
    render_evidence_summary_md,
    write_evidence_pack_file,
    compute_sha256,
)
from review_router import (
    determine_risk_tier,
    audit_review_compliance,
)


class TestRunbookSimulatedLifecycleIntegration:
    """
    Kiểm thử tích hợp mô phỏng chu trình thực thi Runbook (research-repository).
    Xác minh sự chuyển đổi trạng thái giữa các phase và cổng phê duyệt.
    """

    def test_simulated_research_runbook_lifecycle(self, tmp_path, sim_verifiers):
        out_dir = tmp_path / "AIWF_Output" / "research_pilot"
        out_dir.mkdir(parents=True)

        # 0. Khởi tạo Evidence Pack cho runbook research-repository
        task_id = "TASK-PILOT-RR-01"
        pack = create_evidence_pack(
            task_id=task_id,
            runbook="research-repository",
            human_approval={"required": True}
        )
        assert pack["overall_status"] == "NOT_VERIFIED"

        # Phase 1: Research Scoping -> OBSERVED
        scoping_doc = out_dir / "search_strategy.md"
        scoping_doc.write_text("# Search Strategy\n- Query: AI Workforce\n- Sources: SSOT", encoding="utf-8")
        add_step_evidence(
            pack=pack,
            step_id="phase_1_scoping",
            role="research-synthesist",
            skill_used="doc-sau",
            artifact=scoping_doc,
            verification_method="scope_definition",
            evidence_ref=str(scoping_doc),
            status="OBSERVED",
            reviewer={"kind": "automated_test", "name": "ScopeValidator"}
        )

        # Phase 2: Ingestion & Extraction -> OBSERVED with real manifest.json
        source_doc = out_dir / "source.md"
        source_doc.write_text("# Ingested Source\nRaw extracted data blocks.", encoding="utf-8")
        manifest_file = out_dir / "manifest.json"
        manifest_file.write_text(json.dumps({"task_id": task_id, "status": "ok"}), encoding="utf-8")
        # SIMULATED trusted runner / subagent runner / human authority (test doubles, see conftest.py)
        sim_verifiers.executed_sha256.add(compute_sha256(manifest_file))
        sim_verifiers.spawned_reviewers.add("Reality Checker Subagent")
        sim_verifiers.human_signatures.add("sig_rr_dir_001")
        add_step_evidence(
            pack=pack,
            step_id="phase_2_ingest",
            role="research-synthesist",
            skill_used="doc-sau",
            artifact=source_doc,
            verification_method="doc_ingest_bridge",
            evidence_ref=str(manifest_file),
            status="OBSERVED",
            reviewer={"kind": "automated_test", "name": "doc_ingest_bridge"}
        )

        # Phase 3: Multi-Model Deep Reading -> SELF_REVIEWED
        report_doc = out_dir / "draft_deep_reading.md"
        report_doc.write_text("# Deep Reading Analysis\n- SCQA: Situation, Complication, Question, Answer.", encoding="utf-8")
        add_step_evidence(
            pack=pack,
            step_id="phase_3_analysis",
            role="research-synthesist",
            skill_used="doc-sau",
            artifact=report_doc,
            verification_method="mental_models_scqa",
            evidence_ref=str(report_doc),
            status="SELF_REVIEWED",
            reviewer={"kind": "same_context_agent", "name": "DocSau Assistant"}
        )

        # Phase 4: Reality Check -> INDEPENDENTLY_REVIEWED (with verified subagent execution context)
        audit_notes = out_dir / "audit_notes.md"
        audit_notes.write_text("# Reality Check Audit\n- 0 hallucinations found. Citations verified against source.md.", encoding="utf-8")
        audit_log = out_dir / "reality_audit.log"
        audit_log.write_text("Subagent audit executed: zero discrepancies found against source.md", encoding="utf-8")
        add_step_evidence(
            pack=pack,
            step_id="phase_4_reality_check",
            role="reality-checker",
            skill_used="doc-sau",
            artifact=audit_notes,
            verification_method="zero_hallucination_cross_check",
            evidence_ref=str(audit_log),
            status="INDEPENDENTLY_REVIEWED",
            reviewer={
                "kind": "separate_agent",
                "name": "Reality Checker Subagent",
                "subagent_id": "subagent-rr-reality-01",
                "execution_context": "forked_subagent",
            },
        )

        # Phase 5: Publishing -> Real DOCX created via python-docx
        final_docx = out_dir / "bao_cao_nghien_cuu.docx"
        doc = docx.Document()
        doc.add_heading("Báo Cáo Nghiên Cứu AIWF", level=1)
        doc.add_paragraph("Tài liệu nghiên cứu chuyên sâu được tạo lập bằng python-docx.")
        doc.save(str(final_docx))

        add_step_evidence(
            pack=pack,
            step_id="phase_5_publishing",
            role="project-shepherd",
            skill_used="doc-sau",
            artifact=final_docx,
            verification_method="legal_report_generator",
            evidence_ref=str(final_docx),
            status="OBSERVED",
            reviewer={"kind": "automated_test", "name": "Docx Inspect Engine"}
        )

        # Trước khi duyệt, trạng thái tổng thể là AWAITING_HUMAN_APPROVAL
        assert pack["overall_status"] == "AWAITING_HUMAN_APPROVAL"

        # Gate 3 Test: Phê duyệt mô phỏng (agent_simulated)
        # Giữ trạng thái ở mức APPROVAL_CLAIMED và BỊ CHẶN ở cấp độ Tier 3 release
        record_human_approval(
            pack=pack,
            approved=True,
            approver="Research Director (Simulated)",
            notes="Phê duyệt mô phỏng cho bài test chuyển đổi trạng thái",
            channel="agent_simulated"
        )
        assert pack["overall_status"] == "AWAITING_HUMAN_APPROVAL"
        assert pack["human_approval"]["verification_status"] == "SIMULATED_NOT_AUTHENTICATED"
        assert pack["human_approval"]["approval_state"] == "APPROVAL_CLAIMED"

        tier = determine_risk_tier("deep_research", blast_radius=1, is_financial_or_legal=True)
        assert tier == "tier_3_high"

        # Kiểm toán Tier 3 PHẢI TỪ CHỐI phê duyệt mô phỏng (Fail-closed)
        gov_sim_valid, gov_sim_issues = audit_review_compliance(pack, tier, "research")
        assert gov_sim_valid is False
        assert any("mô phỏng" in iss for iss in gov_sim_issues)

        # Gate 3: Khai báo người dùng chưa có token xác thực -> Dừng ở APPROVAL_CLAIMED, Tier 3 bị chặn
        record_human_approval(
            pack=pack,
            approved=True,
            approver="Research Director (User)",
            notes="Nghiệm thu sơ bộ không có token",
            channel="interactive_user_confirmation"
        )
        assert pack["human_approval"]["approval_state"] == "APPROVAL_CLAIMED"
        assert pack["human_approval"]["verification_status"] == "UNATTESTED_CLAIM"
        assert pack["overall_status"] == "AWAITING_HUMAN_APPROVAL"
        gov_unattested_valid, _ = audit_review_compliance(pack, tier, "research")
        assert gov_unattested_valid is False

        # Gate 3 Xác thực Thực tế: Phê duyệt có xác thực người dùng thật (trusted attestation)
        record_human_approval(
            pack=pack,
            approved=True,
            approver="Research Director (User)",
            notes="Nghiệm thu đạt chuẩn 100%, đầy đủ trích dẫn từ nguồn SSOT.",
            channel="interactive_user_confirmation",
            trusted_attestation={
                "verified": True,
                "provider": "trusted_hardware_token",
                "event_signature": "sig_rr_dir_001",
            }
        )
        assert pack["overall_status"] == "VERIFIED"
        assert pack["human_approval"]["status"] == "APPROVED"
        assert pack["human_approval"]["approval_state"] == "APPROVAL_AUTHENTICATED"
        assert pack["human_approval"]["verification_status"] == "AUTHENTICATED"

        # Lưu file Evidence Pack
        pack_file = out_dir / "evidence_pack.json"
        write_evidence_pack_file(pack, pack_file)

        # Kiểm tra tính toàn vẹn của Evidence Pack
        is_valid, issues = validate_evidence_pack(pack_file)
        assert is_valid is True, f"Evidence pack issues: {issues}"

        # Kiểm toán Tier 3 ĐẠT CHUẨN khi đã có xác thực người dùng thật
        gov_valid, gov_issues = audit_review_compliance(pack, tier, "research")
        assert gov_valid is True, f"Governance compliance issues: {gov_issues}"

        # Xuất bản Markdown Summary
        summary_md = render_evidence_summary_md(pack)
        assert "TASK-PILOT-RR-01" in summary_md
        assert "research-repository" in summary_md
        assert "Research Director (User)" in summary_md
        assert "APPROVED" in summary_md

    def test_real_local_execution_smoke(self, tmp_path):
        """
        Smoke test thực thi thực tế engine CLI với file và mã băm thật.
        """
        smoke_dir = tmp_path / "smoke_run"
        smoke_dir.mkdir()

        artifact_file = smoke_dir / "real_output.txt"
        artifact_file.write_text("Smoke test real file content", encoding="utf-8")
        evidence_log = smoke_dir / "pytest_run.log"
        evidence_log.write_text("pytest exit_code=0 passed=10", encoding="utf-8")

        pack = create_evidence_pack("TASK-SMOKE-01", "build-webapp")
        add_step_evidence(
            pack=pack,
            step_id="step_test",
            role="evidence-collector",
            skill_used="app-auditor",
            artifact=artifact_file,
            verification_method="pytest_suite",
            evidence_ref=str(evidence_log),
            status="TESTED",
            reviewer={"kind": "automated_test", "name": "Pytest Engine"}
        )

        pack_file = smoke_dir / "smoke_pack.json"
        write_evidence_pack_file(pack, pack_file)

        # 1. Chạy CLI validate thực tế
        cmd_validate = [
            sys.executable,
            str(SHARED_DIR / "governance" / "evidence_gate.py"),
            "--validate",
            str(pack_file),
        ]
        res_v = subprocess.run(cmd_validate, capture_output=True, text=True)
        assert res_v.returncode == 0
        assert "Evidence Pack hợp lệ 100%" in res_v.stdout

        # 2. Chạy CLI review router thực tế
        cmd_router = [
            sys.executable,
            str(SHARED_DIR / "governance" / "review_router.py"),
            "--check-pack",
            str(pack_file),
            "--tier",
            "tier_2_medium",
        ]
        res_r = subprocess.run(cmd_router, capture_output=True, text=True)
        # Không có runner tin cậy: Tier 2 fail-closed (automated_test khai báo = UNVERIFIED)
        assert res_r.returncode == 1
        assert "UNVERIFIED" in res_r.stdout

        # Tier 1 (quy trình thông thường) không bị chặn
        cmd_t1 = cmd_router[:-1] + ["tier_1_low"]
        res_t1 = subprocess.run(cmd_t1, capture_output=True, text=True)
        assert res_t1.returncode == 0
        assert "Evidence Pack đạt chuẩn quản trị chất lượng" in res_t1.stdout
