"""inspect_docx (engine dùng chung docx.inspect): thứ tự đọc, kế thừa font, ô gộp, phát hiện cơ học, đầu vào hỏng."""
import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
MODULE_DIR = ROOT / ".agents" / "skills" / "_shared" / "docx"
sys.path.insert(0, str(MODULE_DIR))
from inspect_docx import InspectError, inspect_docx  # noqa: E402

CLI = [sys.executable, str(MODULE_DIR / "inspect_docx.py")]


def _png(tmp_path, size=(300, 150)):
    p = tmp_path / "img.png"
    Image.new("RGB", size, "white").save(p)
    return p


def test_reading_order_and_counts(tmp_path):
    d = Document()
    d.add_heading("Tiêu đề", 1)
    d.add_paragraph("Trước bảng")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A"
    t.cell(1, 1).text = "D"
    d.add_paragraph("Sau bảng")
    d.add_picture(str(_png(tmp_path)), width=Mm(50))
    f = tmp_path / "a.docx"
    d.save(f)

    r = inspect_docx(f)
    kinds = [b["kind"] for b in r["blocks"]]
    texts = [b["text"] for b in r["blocks"] if b["kind"] == "paragraph" and b["text"]]
    assert kinds[:4] == ["paragraph", "paragraph", "table", "paragraph"]
    assert texts == ["Tiêu đề", "Trước bảng", "Sau bảng"]
    assert r["blocks"][0]["outline"] == 1
    assert r["stats"]["tables"] == 1 and r["stats"]["images"] == 1
    img = r["images"][0]
    assert img["px"] == [300, 150]
    assert img["placed_mm"][0] == 50.0
    assert img["effective_dpi"] == round(300 / (50 / 25.4))
    assert img["alt"] is None and r["findings"]["missing_alt"] == [img["block"]]


def test_font_inheritance_sources(tmp_path):
    d = Document()
    dd = d.styles.element.find(qn("w:docDefaults"))
    rf = dd.find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:rFonts"))
    for a in list(rf.attrib):
        del rf.attrib[a]
    rf.set(qn("w:ascii"), "Courier New")

    d.styles["Heading 1"].font.name = "Arial"
    p_default = d.add_paragraph("mặc định")
    p_style = d.add_heading("kiểu", 1)
    p_run = d.add_paragraph()
    p_run.add_run("riêng").font.name = "Times New Roman"
    f = tmp_path / "f.docx"
    d.save(f)

    runs = {b["text"]: b["runs"][0] for b in inspect_docx(f)["blocks"] if b["kind"] == "paragraph" and b["runs"]}
    assert runs["kiểu"]["font"] == "Arial" and runs["kiểu"]["font_src"] == "style"
    assert runs["riêng"]["font"] == "Times New Roman" and runs["riêng"]["font_src"] == "run"
    assert runs["mặc định"]["font"] == "Courier New" and runs["mặc định"]["font_src"] == "docDefaults"


def test_theme_font_is_resolved(tmp_path):
    f = tmp_path / "t.docx"
    d = Document()
    d.add_paragraph("chữ")
    d.save(f)
    r = inspect_docx(f)
    run = r["blocks"][0]["runs"][0]
    assert run["font"], "font phải được phân giải, không để None"
    assert run["font_src"] in {"theme", "docDefaults", "style"}


def test_merged_cells_not_double_counted(tmp_path):
    d = Document()
    t = d.add_table(rows=3, cols=3)
    t.cell(0, 0).merge(t.cell(0, 1))
    t.cell(1, 0).merge(t.cell(2, 0))
    f = tmp_path / "m.docx"
    d.save(f)

    tb = next(b for b in inspect_docx(f)["blocks"] if b["kind"] == "table")
    assert tb["rows"] == 3 and tb["cols"] == 3
    assert [c["span"] for c in tb["cells"][0]] == [2, 1]
    assert tb["cells"][1][0]["vmerge"] == "restart"
    assert tb["cells"][2][0]["vmerge"] == "continue"
    assert tb["merged"] == 3


def test_mechanical_findings(tmp_path):
    d = Document()
    d.add_heading("Một", 1)
    d.add_paragraph("Điền {{ten_khach}} vào đây — nhanh")
    d.add_paragraph("")
    d.add_paragraph("")
    d.add_paragraph("")
    d.add_heading("Ba", 3)
    d.add_paragraph("Việt Nam".replace("ệ", "ệ"))  # NFD
    d.add_paragraph("Số liệu [CẦN XÁC MINH: nguồn]")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "TODO điền số"
    f = tmp_path / "x.docx"
    d.save(f)

    fi = inspect_docx(f)["findings"]
    matches = {m["match"] for m in fi["placeholders"]}
    assert "{{ten_khach}}" in matches and "TODO" in matches  # cả trong ô bảng
    assert fi["empty_paragraph_runs"] == [{"start": 2, "count": 3}]
    assert fi["heading_gaps"] == [{"block": 5, "from": 1, "to": 3}]
    assert 1 in fi["em_dash"]
    assert 6 in fi["nfd_text"]
    assert len(fi["verify_flags"]) == 1


def test_tracked_changes_and_comments_counted(tmp_path):
    d = Document()
    p = d.add_paragraph()
    run = p.add_run("giữ")
    ins, dele = OxmlElement("w:ins"), OxmlElement("w:del")
    for el in (ins, dele):
        el.set(qn("w:id"), "1")
        el.set(qn("w:author"), "t")
        p._element.append(el)
    d.add_comment(runs=run, text="xem lại", author="t")
    f = tmp_path / "r.docx"
    d.save(f)

    rv = inspect_docx(f)["review"]
    assert rv["tracked_insertions"] == 1 and rv["tracked_deletions"] == 1 and rv["comments"] == 1


def test_sections_and_page_field(tmp_path):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = Mm(30)
    fp = s.footer.paragraphs[0]
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._element.append(fld)
    d.add_paragraph("x")
    f = tmp_path / "s.docx"
    d.save(f)

    r = inspect_docx(f)
    sec = r["sections"][0]
    assert sec["page_mm"] == [210.0, 297.0] and sec["margins_mm"]["left"] == 30.0
    assert sec["has_page_number_field"] is True
    assert r["fields"]["page"] == 1  # trường ở chân trang cũng được đếm


def test_hyperlink_collected(tmp_path):
    d = Document()
    p = d.add_paragraph("xem ")
    rid = d.part.relate_to("https://example.com", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "trang"
    r.append(t)
    h.append(r)
    p._element.append(h)
    f = tmp_path / "h.docx"
    d.save(f)
    r = inspect_docx(f)
    assert {"text": "trang", "target": "https://example.com"} in r["hyperlinks"]
    assert r["blocks"][0]["text"] == "xem trang"


def test_input_file_not_modified(tmp_path):
    d = Document()
    d.add_paragraph("x")
    f = tmp_path / "ro.docx"
    d.save(f)
    before = f.read_bytes()
    inspect_docx(f)
    assert f.read_bytes() == before


# ---- đầu vào hỏng: mã 2, không traceback ----

def _bad_files(tmp_path):
    txt = tmp_path / "not.docx"
    txt.write_text("xin chào", encoding="utf-8")
    nodoc = tmp_path / "nodoc.docx"
    with zipfile.ZipFile(nodoc, "w") as z:
        z.writestr("hello.txt", "x")
    ole = tmp_path / "ole.docx"
    ole.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64)
    return [txt, nodoc, ole, tmp_path / "missing.docx"]


def test_bad_inputs_raise_inspect_error(tmp_path):
    for f in _bad_files(tmp_path):
        with pytest.raises(InspectError):
            inspect_docx(f)


def test_bad_inputs_cli_exit_2_without_traceback(tmp_path):
    for f in _bad_files(tmp_path):
        res = subprocess.run(CLI + [str(f)], capture_output=True, text=True)
        assert res.returncode == 2, f.name
        assert "Traceback" not in res.stderr, f.name
        assert res.stderr.strip(), f.name


def test_cli_json_text_and_summary(tmp_path):
    d = Document()
    d.add_paragraph("Dòng một")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "a"
    t.cell(0, 1).text = "b"
    f = tmp_path / "c.docx"
    d.save(f)
    out = tmp_path / "o.json"

    res = subprocess.run(CLI + [str(f), "--json", str(out)], capture_output=True, text=True)
    assert res.returncode == 0 and "1 bảng" in res.stdout
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["schema"] == "aiwf.docx.inspect/1" and data["stats"]["tables"] == 1

    res = subprocess.run(CLI + [str(f), "--text"], capture_output=True, text=True)
    assert res.returncode == 0
    assert res.stdout.splitlines()[:2] == ["Dòng một", "a\tb"]


# ---- --rules (bước 2) ----

def _rules_doc(tmp_path):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = Mm(30)
    d.add_heading("Một", 1)
    r = d.add_paragraph().add_run("Điền {{ten}} — nhanh")
    r.font.name = "Arial"
    d.add_paragraph("")
    d.add_paragraph("")
    d.add_paragraph("")
    d.add_heading("Ba", 3)
    d.add_picture(str(_png(tmp_path, (100, 50))), width=Mm(100))  # ~25 dpi, không alt
    f = tmp_path / "rules.docx"
    d.save(f)
    return f


def _run_rules(f, tmp_path, rules):
    rf = tmp_path / "rules.json"
    rf.write_text(json.dumps(rules), encoding="utf-8")
    return subprocess.run(CLI + [str(f), "--rules", str(rf)], capture_output=True, text=True)


def test_rules_violations_exit_1(tmp_path):
    f = _rules_doc(tmp_path)
    res = _run_rules(f, tmp_path, {
        "fonts_allowed": ["Times New Roman"], "page_mm": [210, 297], "margins_mm": {"left": [15, 25]},
        "no_placeholders": True, "max_consecutive_empty_paragraphs": 1, "heading_no_skip": True,
        "min_image_dpi": 150, "require_alt_text": True, "forbid": ["em_dash"]})
    assert res.returncode == 1, res.stdout + res.stderr
    for name in ("fonts_allowed", "margins_mm", "no_placeholders", "max_consecutive_empty_paragraphs",
                 "heading_no_skip", "min_image_dpi", "require_alt_text", "forbid.em_dash"):
        assert f"RULE {name}:" in res.stdout, name
    assert "RULE page_mm" not in res.stdout  # khổ A4 đúng
    assert "RULES FAIL" in res.stdout


def test_rules_pass_exit_0_and_json_records_result(tmp_path):
    d = Document()
    d.add_paragraph("sạch")
    f = tmp_path / "ok.docx"
    d.save(f)
    out = tmp_path / "o.json"
    rf = tmp_path / "r.json"
    rf.write_text(json.dumps({"no_placeholders": True, "forbid": ["em_dash", "nfd_text"]}), encoding="utf-8")
    res = subprocess.run(CLI + [str(f), "--rules", str(rf), "--json", str(out)], capture_output=True, text=True)
    assert res.returncode == 0 and "RULES PASS" in res.stdout
    assert json.loads(out.read_text(encoding="utf-8"))["rules"]["passed"] is True


def test_rules_unverifiable_or_invalid_exit_2(tmp_path):
    d = Document()
    d.add_paragraph("x")
    f = tmp_path / "u.docx"
    d.save(f)
    for bad in ({"khong_co_luat_nay": 1}, {"forbid": ["khong_hop_le"]}, {"pages": {"toi_da": 2}},
                {"margins_mm": {"giua": [1, 2]}}, ["khong", "phai", "doi", "tuong"]):
        res = _run_rules(f, tmp_path, bad)
        assert res.returncode == 2, (bad, res.stdout, res.stderr)
        assert "Traceback" not in res.stderr and res.stderr.strip()
    res = subprocess.run(CLI + [str(f), "--rules", str(tmp_path / "khong_ton_tai.json")], capture_output=True, text=True)
    assert res.returncode == 2 and "Traceback" not in res.stderr


# ---- --pages (bước 4) ----

import soffice_tools  # noqa: E402  (đường dẫn _shared/office được inspect_docx thêm vào sys.path)

needs_lo = pytest.mark.skipif(not soffice_tools.find_soffice(), reason="cần LibreOffice")


def _long_doc(tmp_path, paragraphs):
    d = Document()
    for i in range(paragraphs):
        p = d.add_paragraph()
        p.add_run(f"Đoạn số {i} " + "chữ " * 80).font.name = "Times New Roman"
        p.paragraph_format.page_break_before = (i > 0)
    f = tmp_path / "long.docx"
    d.save(f)
    return f


@needs_lo
def test_pages_real_count_and_text(tmp_path):
    r = inspect_docx(_long_doc(tmp_path, 3), pages=True)
    assert r["pages"]["count"] == 3
    assert "Đoạn số 1" in r["pages"]["text"][1]


@needs_lo
def test_pages_rule_pass_and_fail_via_cli(tmp_path):
    f = _long_doc(tmp_path, 3)
    ok = _run_rules(f, tmp_path, {"pages": {"min": 3, "max": 3}})
    assert ok.returncode == 0, ok.stdout + ok.stderr
    bad = _run_rules(f, tmp_path, {"pages": {"max": 2}})
    assert bad.returncode == 1 and "RULE pages:" in bad.stdout and "3 trang" in bad.stdout


def test_pages_without_libreoffice_is_unverifiable_not_crash(tmp_path, monkeypatch):
    monkeypatch.setattr(soffice_tools, "find_soffice", lambda: None)
    f = _long_doc(tmp_path, 1)
    r = inspect_docx(f, pages=True)
    assert r["pages"] is None and any("LibreOffice" in w for w in r["warnings"])
    with pytest.raises(InspectError):
        from inspect_docx import check_rules
        check_rules(r, {"pages": {"max": 2}})
    assert inspect_docx(f)["pages"] is None  # không --pages thì không đòi LibreOffice


def test_diacritic_split_heuristic():
    from inspect_docx import _diacritic_split_warning
    broken = "Đo n s 1 ch Vi t Nam\nạ\nố\nữ\nệ\n"  # kiểu LibreOffice thay font Cambria
    assert "thiếu font" in _diacritic_split_warning([broken])
    assert _diacritic_split_warning(["Đoạn số 1 chữ Việt Nam\nA\nb\n"]) is None  # ASCII lẻ không tính
    assert _diacritic_split_warning(["ạ\nố\n"]) is None  # dưới ngưỡng 3


def test_require_page_numbers_rule(tmp_path):
    d = Document()
    d.add_paragraph("không số trang")
    f = tmp_path / "np.docx"
    d.save(f)
    res = _run_rules(f, tmp_path, {"require_page_numbers": True})
    assert res.returncode == 1 and "RULE require_page_numbers:" in res.stdout
