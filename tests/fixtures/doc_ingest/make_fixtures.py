#!/usr/bin/env python3
"""Sinh lại bộ fixture nhỏ cho tests/test_doc_ingest.py (đã commit sẵn; chỉ chạy khi cần tạo lại).

Cần: python-docx, openpyxl, python-pptx, pymupdf, Pillow (repo .venv) + LibreOffice (soffice) cho doc/odt/rtf/pdf/xlsx.
    .venv/bin/python tests/fixtures/doc_ingest/make_fixtures.py
Tổng dung lượng giữ < 1 MB.
"""
import io
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[2] / "scripts"))
from doc_ingest import find_soffice  # noqa: E402

TITLE = "Báo cáo thử nghiệm định dạng"
P1 = "Nghị định 30/2020/NĐ-CP quy định thể thức văn bản hành chính tại Thành phố Hồ Chí Minh."
P2 = "Kiểm tra dấu: ấ ầ ẩ ẫ ậ ắ ằ ẳ ẵ ặ ế ề ể ễ ệ ố ồ ổ ỗ ộ ớ ờ ở ỡ ợ ứ ừ ử ữ ự ỳ ỷ ỹ ỵ đ Đ."
P3 = "Mọi nghĩa vụ thuế thu nhập cá nhân phải được thực hiện đúng hạn."
HDR = ["Hạng mục", "Số lượng", "Đơn giá (đồng)"]
ROWS = [["Bút bi xanh", 120, 3500], ["Giấy A4 định lượng 80", 15, 72000], ["Mực in Canon", 4, 450000]]
HEADER_TXT = "Công ty Thử Nghiệm — tài liệu nội bộ"
FONT_CANDIDATES = ["/System/Library/Fonts/Supplemental/Arial.ttf", "/Library/Fonts/Arial Unicode.ttf",
                   "C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]


def font(size):
    from PIL import ImageFont
    for f in FONT_CANDIDATES:
        if Path(f).exists():
            return ImageFont.truetype(f, size)
    return ImageFont.load_default()


def small_png() -> bytes:
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (240, 64), (230, 240, 255))
    d = ImageDraw.Draw(img)
    d.rectangle([2, 2, 237, 61], outline=(0, 60, 160), width=2)
    d.text((10, 20), "HÌNH MINH HỌA", fill=(0, 0, 0), font=font(20))
    buf = io.BytesIO()
    img.save(buf, "PNG", optimize=True)
    return buf.getvalue()


def soffice(src: Path, fmt: str, outdir: Path):
    exe = find_soffice()
    prof = Path(tempfile.mkdtemp()) / "lo"
    subprocess.run([exe, f"-env:UserInstallation={prof.as_uri()}", "--headless", "--convert-to", fmt,
                    "--outdir", str(outdir), str(src)], check=True, capture_output=True, timeout=240)


def make_docx(png: bytes):
    import docx
    from docx.shared import Inches
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = HEADER_TXT
    d.add_heading(TITLE, 0)
    d.add_heading("Chương 1: Mở đầu", 1)
    d.add_paragraph(P1)
    d.add_paragraph(P2)
    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(HDR):
        t.rows[0].cells[i].text = h
    for r in ROWS:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = f"{v:,}".replace(",", ".") if isinstance(v, int) else v
    d.add_picture(io.BytesIO(png), width=Inches(2))
    d.add_heading("Chương 2: Kết luận", 1)
    d.add_paragraph(P3)
    d.save(HERE / "sample.docx")


def make_tcvn3_docx():
    import docx
    from docx.oxml.ns import qn
    d = docx.Document()
    for text, fnt in (("CéNG HOµ X· HéI CHñ NGHÜA VIÖT NAM", ".VnTimeH"),
                      ("§éc lËp - Tù do - H¹nh phóc", ".VnTime"),
                      ("Céng hoµ x· héi chñ nghÜa ViÖt Nam", ".VnTime"),
                      ("Ng\u00adêi lao ®éng ph¶i ®\u00adîc tr¶ l\u00ad¬ng ®óng h¹n.", ".VnTime")):
        run = d.add_paragraph().add_run(text)
        rpr = run._element.get_or_add_rPr()
        rf = rpr.makeelement(qn("w:rFonts"), {qn("w:ascii"): fnt, qn("w:hAnsi"): fnt})
        rpr.append(rf)
    d.save(HERE / "tcvn3.docx")


def make_xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chi phí"
    ws.append(HDR + ["Thành tiền"])
    for i, r in enumerate(ROWS, start=2):
        ws.append(r + [f"=B{i}*C{i}"])
    ws.append(["Tổng cộng", "=SUM(B2:B4)", None, "=SUM(D2:D4)"])
    ws2 = wb.create_sheet("Tổng hợp")
    ws2.append(["Chỉ tiêu", "Giá trị"])
    ws2.append(["Tổng chi phí quý III", "='Chi phí'!D5"])
    ws2.append(["Ghi chú", P3])
    wb.save(HERE / "sample_nocache.xlsx")
    tmp = Path(tempfile.mkdtemp())
    shutil.copy(HERE / "sample_nocache.xlsx", tmp / "sample.xlsx")
    out = tmp / "out"
    soffice(tmp / "sample.xlsx", "xlsx:Calc MS Excel 2007 XML", out)
    shutil.copy(out / "sample.xlsx", HERE / "sample.xlsx")


def make_pptx(png: bytes):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = TITLE
    s.placeholders[1].text = P1 + "\n" + P2
    s.notes_slide.notes_text_frame.text = "Ghi chú diễn giả: nhấn mạnh tiến độ quý ba."
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Bảng chi phí"
    tb = s.shapes.add_table(4, 3, Inches(0.5), Inches(1.8), Inches(9), Inches(2)).table
    for i, h in enumerate(HDR):
        tb.cell(0, i).text = h
    for ri, r in enumerate(ROWS, 1):
        for ci, v in enumerate(r):
            tb.cell(ri, ci).text = str(v)
    s.notes_slide.notes_text_frame.text = "Ghi chú bảng: đơn giá đã gồm VAT."
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Chương 2: Kết luận"
    s.shapes.add_picture(io.BytesIO(png), Inches(0.5), Inches(1.6), width=Inches(3))
    cd = CategoryChartData()
    cd.categories = ["Quý I", "Quý II", "Quý III"]
    cd.add_series("Doanh thu (tỷ đồng)", (12.5, 15.0, 18.2))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4), Inches(1.6), Inches(5), Inches(3), cd)
    tx = s.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(8), Inches(1))
    tx.text_frame.text = P3
    s.notes_slide.notes_text_frame.text = "Ghi chú kết luận: hạn chót ngày 31 tháng 12."
    prs.save(HERE / "sample.pptx")


def make_epub(png: bytes):
    """Tên tệp chương đảo ngược thứ tự spine (z_ là chương 1, a_ là chương 2) để kiểm tra đọc theo spine."""
    xhtml = ('<?xml version="1.0" encoding="utf-8"?>\n<html xmlns="http://www.w3.org/1999/xhtml"><head><title>{t}</title>'
             '</head><body>{b}</body></html>')
    tbl = "<table><tr>" + "".join(f"<td>{h}</td>" for h in HDR) + "</tr>" + "".join(
        "<tr>" + "".join(f"<td>{v}</td>" for v in r) + "</tr>" for r in ROWS) + "</table>"
    ch1 = xhtml.format(t="Chương 1", b=f"<h1>Chương 1: Mở đầu</h1><p>{P1}</p><p>{P2}</p>{tbl}"
                                        '<p><img src="../images/hinh.png" alt="Hình minh họa"/></p>')
    ch2 = xhtml.format(t="Chương 2", b=f"<h1>Chương 2: Kết luận</h1><p>{P3}</p>")
    opf = f"""<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">urn:uuid:aiwf-test</dc:identifier>
<dc:title>{TITLE}</dc:title><dc:creator>AIWF</dc:creator><dc:language>vi</dc:language></metadata>
<manifest>
<item id="c1" href="text/z_chuong1.xhtml" media-type="application/xhtml+xml"/>
<item id="c2" href="text/a_chuong2.xhtml" media-type="application/xhtml+xml"/>
<item id="img" href="images/hinh.png" media-type="image/png"/>
</manifest>
<spine><itemref idref="c1"/><itemref idref="c2"/></spine>
</package>"""
    container = ('<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                 '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                 '</rootfiles></container>')
    with zipfile.ZipFile(HERE / "sample.epub", "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", container, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("OEBPS/content.opf", opf, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("OEBPS/text/a_chuong2.xhtml", ch2, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("OEBPS/text/z_chuong1.xhtml", ch1, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("OEBPS/images/hinh.png", png)


def to_cp1258_bytes(s: str) -> bytes:
    """CP1258: chữ có mũ/móc/trăng dựng sẵn + dấu thanh dạng tổ hợp."""
    tones = {"\u0300", "\u0301", "\u0303", "\u0309", "\u0323"}
    out = []
    for ch in s:
        d = unicodedata.normalize("NFD", ch)
        base = unicodedata.normalize("NFC", "".join(c for c in d if c not in tones))
        out.append(base + "".join(c for c in d if c in tones))
    return "".join(out).encode("cp1258")


def make_texts():
    import csv
    with open(HERE / "sample.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(HDR + ["Thành tiền"])
        for r in ROWS:
            w.writerow(r + [r[1] * r[2]])
    (HERE / "cp1258.txt").write_bytes(to_cp1258_bytes(P1 + "\n" + P3 + "\n"))
    tcvn3 = "Céng hoµ x· héi chñ nghÜa ViÖt Nam\n§éc lËp - Tù do - H¹nh phóc\nNg\u00adêi lao ®éng ph¶i ®\u00adîc tr¶ l\u00ad¬ng ®óng h¹n.\n"
    (HERE / "tcvn3.txt").write_bytes(tcvn3.encode("latin-1"))


def make_pdfs():
    import pymupdf
    from PIL import Image
    tmp = Path(tempfile.mkdtemp())
    shutil.copy(HERE / "sample.docx", tmp / "sample.docx")
    soffice(tmp / "sample.docx", "pdf", tmp)
    shutil.copy(tmp / "sample.pdf", HERE / "sample.pdf")
    src = pymupdf.open(HERE / "sample.pdf")
    pix = src[0].get_pixmap(dpi=100, colorspace=pymupdf.csGRAY)
    img = Image.frombytes("L", (pix.width, pix.height), pix.samples)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=55)
    (HERE / "scan.jpg").write_bytes(buf.getvalue())
    out = pymupdf.open()
    pg = out.new_page(width=src[0].rect.width, height=src[0].rect.height)
    pg.insert_image(pg.rect, stream=buf.getvalue())
    out.save(HERE / "scanned.pdf", deflate=True)
    src.save(HERE / "password.pdf", encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="matkhau-thu", owner_pw="chu-so-huu")


def make_legacy():
    tmp = Path(tempfile.mkdtemp())
    shutil.copy(HERE / "sample.docx", tmp / "sample.docx")
    for fmt in ("doc:MS Word 97", "odt", "rtf"):
        soffice(tmp / "sample.docx", fmt, tmp)
    for ext in ("doc", "odt", "rtf"):
        shutil.copy(tmp / f"sample.{ext}", HERE / f"sample.{ext}")


if __name__ == "__main__":
    if not find_soffice():
        sys.exit("Cần LibreOffice (soffice) để sinh doc/odt/rtf/pdf/xlsx.")
    png = small_png()
    make_docx(png)
    make_tcvn3_docx()
    make_xlsx()
    make_pptx(png)
    make_epub(png)
    make_texts()
    make_pdfs()
    make_legacy()
    total = sum(f.stat().st_size for f in HERE.iterdir() if f.is_file())
    for f in sorted(HERE.iterdir()):
        print(f"{f.stat().st_size:>8}  {f.name}")
    print(f"Tổng: {total / 1024:.0f} KB")
